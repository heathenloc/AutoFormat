"""Run with base Python -I -S -B: verify only bundled dependencies and real V1.

Example: python -I -S -B scripts/smoke_windows_bundle.py --bundle <directory>
    --config <external config> --work <new project runtime directory>
Uses the explicit offline parser; the DOCX editor and PDF renderer are real.
"""
import argparse
from io import BytesIO
import json
from pathlib import Path
import runpy
import sys


def smoke(bundle, config_path, work):
    if not sys.flags.isolated or not sys.flags.no_site or not sys.dont_write_bytecode:
        raise ValueError("Use Python -I -S -B to isolate runtime dependencies")
    bundle = bundle.resolve(strict=True)
    verify = runpy.run_path(str(bundle / "launch.py"))["verify_bundle"]
    verify(bundle)
    work = work.resolve()
    if not work.is_relative_to(Path(__file__).resolve().parents[1] / "runtime") or work.exists():
        raise ValueError("Use a new work directory under the project's runtime")
    work.mkdir(parents=True)
    sys.path[:0] = [str(bundle / "apps/api/src"), str(bundle / "modules/brief_parsing/src"), str(bundle / "python")]
    from fastapi.testclient import TestClient
    from docx import Document
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from autoformat.desktop_host import create_application, load_config
    from autoformat.selection_indexes import SelectionIndexes
    from autoformat.preview import _sha256
    from autoformat.version_store import VersionStore, new_id

    config = load_config(bundle, config_path)
    host, diagnostics = create_application(bundle, config, work / "store", demo=True)
    document = Document()
    document.styles["Normal"].font.name = "Arial"
    document.styles["Normal"].font.size = Pt(12)
    document.add_paragraph("Bundle smoke paragraph", style="Body Text")
    paper = BytesIO(); document.save(paper)
    original = paper.getvalue()
    try:
        with TestClient(host.app, base_url=host.url) as client:
            headers = {"X-AutoFormat-Session": host.session_token, "Origin": host.url}
            assert client.get("/api/tasks").status_code == 401
            assert client.get("/", headers=headers).status_code == 200
            assert client.get("/api/application", headers=headers).json()["mode"] == "demo"
            form = {"request_id": new_id(), "model_processing_acknowledged": "true"}
            files = {"brief": ("school.txt", "正文：14pt，居中。".encode("utf-8"), "text/plain"),
                     "paper": ("paper.docx", original, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
            response = client.post("/api/tasks", headers=headers, data=form, files=files)
            assert response.status_code == 200, response.text
            task = response.json(); assert task["phase"] == "initial_draft_ready", task
            replay = client.post("/api/tasks", headers=headers, data=form, files=files)
            assert replay.json()["initial_version_id"] == task["initial_version_id"]
            base = f'/api/tasks/{task["task_id"]}/versions'
            v0 = client.get(f'{base}/{task["base_version_id"]}/artifacts/docx', headers=headers)
            assert v0.content == original
            v1 = client.get(f'{base}/{task["initial_version_id"]}/artifacts/docx', headers=headers)
            edited = Document(BytesIO(v1.content)); paragraph = edited.paragraphs[0]
            assert paragraph.text == "Bundle smoke paragraph"
            assert paragraph.alignment == WD_ALIGN_PARAGRAPH.CENTER
            assert all(run.font.size.pt == 14 for run in paragraph.runs)
            report = client.get(f'{base}/{task["initial_version_id"]}/artifacts/report_json', headers=headers).json()
            assert {r["property_id"] for r in report["rules"] if r["status"] == "modified_verified"} >= {"font.size", "paragraph.alignment"}
            assert len(client.get(base, headers=headers).json()["versions"]) == 2
            store = VersionStore(work / "store")
            pdf = store.artifact_path(task["task_id"], task["initial_version_id"], "preview")
            indexes = SelectionIndexes(node=config.node, dotnet=config.dotnet, script=bundle / "scripts/pdf_text_index.mjs", editor_dll=config.editor_dll)
            assert "Bundle smoke paragraph" in "".join(item.text for item in indexes.pdf(pdf, [1]).pages[0].items)
            assert indexes.docx(store.artifact_path(task["task_id"], task["initial_version_id"], "docx")).paragraphs
            published_hashes = {_name: _sha256(store.artifact_path(task["task_id"], task["initial_version_id"], _name))
                                for _name in ("docx", "report_json", "preview")}
    finally: host.close()
    recovered, _ = create_application(bundle, config, work / "store", demo=True)
    try:
        assert recovered.session_token != host.session_token
        with TestClient(recovered.app, base_url=recovered.url) as client:
            assert client.get("/api/tasks", headers={"X-AutoFormat-Session": host.session_token}).status_code == 401
            catalog = client.get("/api/tasks", headers={"X-AutoFormat-Session": recovered.session_token}).json()
            assert catalog["tasks"][0]["working_version_id"] == task["initial_version_id"]
            assert catalog["tasks"][0]["version_count"] == 2
            for name, digest in published_hashes.items():
                assert _sha256(store.artifact_path(task["task_id"], task["initial_version_id"], name)) == digest
    finally: recovered.close()
    verify(bundle)  # Running must not write caches, profiles or tasks into the bundle.
    modules = ("autoformat", "brief_parser", "jsonschema", "fastapi", "pydantic", "docx", "lxml", "pypdf", "httpx2")
    for name in modules: assert Path(sys.modules[name].__file__).resolve().is_relative_to(bundle), name
    print(json.dumps({"status": "passed", "isolation": "-I -S -B", "mode": diagnostics["mode"],
        "real_editor": True, "real_pdf_renderer": True, "pdfjs_index": True, "upload_replay": True,
        "restart_recovery": True, "original_and_history_unchanged": True, "bundled_modules": list(modules)}))


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--bundle", type=Path, required=True)
    parser.add_argument("--config", type=Path, required=True)
    parser.add_argument("--work", type=Path, required=True)
    args = parser.parse_args()
    smoke(args.bundle, args.config, args.work)
