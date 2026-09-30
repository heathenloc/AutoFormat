"""DOCX adapter preserving body order, tables, headers, footers, and footnotes."""

from __future__ import annotations

import re
from io import BytesIO
from zipfile import ZipFile

from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from lxml import etree  # type: ignore[import-untyped]

from brief_parser.adapters.parsers.base import BlockDraft, ParserResult
from brief_parser.domain.document import (
    BlockKind,
    BlockStyle,
    DocumentLocator,
    InlineStyleSpan,
)
from brief_parser.errors import BriefParserError, ErrorCode
from brief_parser.ports.document_ingestion import SecuredSource, SourceKind

_HEADING_STYLE = re.compile(r"(?:heading|标题)\s*([1-9])", re.IGNORECASE)
_WORD_NAMESPACE = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


class DocxParser:
    supported_kinds: frozenset[SourceKind] = frozenset({"docx"})
    parser_id = "docx-python-docx"
    parser_version = "1.0.0"

    def parse(self, source: SecuredSource) -> ParserResult:
        if source.kind != "docx":
            raise ValueError(f"unsupported source kind for DocxParser: {source.kind}")
        try:
            document = Document(BytesIO(source.content))
        except (PackageNotFoundError, ValueError, KeyError, OSError) as exc:
            raise BriefParserError(
                ErrorCode.SOURCE_EXTRACTION_FAILED,
                "DOCX 正文无法稳定提取。",
            ) from exc

        drafts: list[BlockDraft] = []
        section_titles: dict[int, str] = {}
        paragraph_index = 0
        table_index = 0
        for item in document.iter_inner_content():
            if isinstance(item, Paragraph):
                paragraph_index += 1
                drafts.extend(
                    self._paragraph_drafts(
                        item,
                        paragraph_index=paragraph_index,
                        section_titles=section_titles,
                    )
                )
            elif isinstance(item, Table):
                table_index += 1
                drafts.extend(
                    self._table_drafts(
                        item,
                        table_index=table_index,
                        section_titles=section_titles,
                    )
                )

        duplicate_header_count = 0
        duplicate_footer_count = 0
        seen_headers: set[str] = set()
        seen_footers: set[str] = set()
        for section_index, section in enumerate(document.sections, start=1):
            for paragraph_index, paragraph in enumerate(section.header.paragraphs, start=1):
                text = paragraph.text
                if not text.strip():
                    continue
                if text in seen_headers:
                    duplicate_header_count += 1
                    continue
                seen_headers.add(text)
                drafts.append(
                    BlockDraft(
                        kind=BlockKind.HEADER,
                        raw_text=text,
                        locator=DocumentLocator(
                            paragraph=paragraph_index,
                            section=section_index,
                            section_path=["header"],
                        ),
                        style=self._paragraph_style(paragraph),
                    )
                )
            for paragraph_index, paragraph in enumerate(section.footer.paragraphs, start=1):
                text = paragraph.text
                if not text.strip():
                    continue
                if text in seen_footers:
                    duplicate_footer_count += 1
                    continue
                seen_footers.add(text)
                drafts.append(
                    BlockDraft(
                        kind=BlockKind.FOOTER,
                        raw_text=text,
                        locator=DocumentLocator(
                            paragraph=paragraph_index,
                            section=section_index,
                            section_path=["footer"],
                        ),
                        style=self._paragraph_style(paragraph),
                    )
                )

        drafts.extend(self._footnote_drafts(source.content))
        warnings: list[str] = []
        if duplicate_header_count:
            warnings.append("DUPLICATE_HEADERS_REMOVED")
        if duplicate_footer_count:
            warnings.append("DUPLICATE_FOOTERS_REMOVED")
        return ParserResult(
            parser_id=self.parser_id,
            parser_version=self.parser_version,
            drafts=tuple(drafts),
            duplicate_header_count=duplicate_header_count,
            duplicate_footer_count=duplicate_footer_count,
            warnings=tuple(warnings),
        )

    def _paragraph_drafts(
        self,
        paragraph: Paragraph,
        *,
        paragraph_index: int,
        section_titles: dict[int, str],
    ) -> list[BlockDraft]:
        drafts: list[BlockDraft] = []
        text = paragraph.text
        heading_level = self._heading_level(paragraph)
        list_level = self._list_level(paragraph)
        if text.strip():
            if heading_level is not None:
                for level in tuple(section_titles):
                    if level >= heading_level:
                        del section_titles[level]
                section_titles[heading_level] = text.strip()
                kind = BlockKind.HEADING
            elif list_level is not None:
                kind = BlockKind.LIST_ITEM
            else:
                kind = BlockKind.PARAGRAPH
            drafts.append(
                BlockDraft(
                    kind=kind,
                    raw_text=text,
                    locator=DocumentLocator(
                        paragraph=paragraph_index,
                        section_path=[
                            section_titles[level] for level in sorted(section_titles)
                        ],
                    ),
                    style=BlockStyle(
                        heading_level=heading_level,
                        list_level=list_level,
                        bold_ratio=self._bold_ratio(paragraph),
                        inline_spans=self._inline_spans([paragraph]),
                    ),
                )
            )
        if self._has_page_break(paragraph):
            drafts.append(
                BlockDraft(
                    kind=BlockKind.PAGE_BREAK,
                    raw_text="\f",
                    locator=DocumentLocator(
                        paragraph=paragraph_index,
                        section_path=[
                            section_titles[level] for level in sorted(section_titles)
                        ],
                    ),
                )
            )
        return drafts

    def _table_drafts(
        self,
        table: Table,
        *,
        table_index: int,
        section_titles: dict[int, str],
    ) -> list[BlockDraft]:
        drafts: list[BlockDraft] = []
        seen_cells: set[int] = set()
        for row_index, row in enumerate(table.rows, start=1):
            for column_index, cell in enumerate(row.cells, start=1):
                cell_identity = id(cell._tc)
                if cell_identity in seen_cells:
                    continue
                seen_cells.add(cell_identity)
                paragraphs = [paragraph for paragraph in cell.paragraphs if paragraph.text.strip()]
                text = "\n".join(paragraph.text for paragraph in paragraphs)
                if not text.strip():
                    continue
                total_length = sum(len(paragraph.text) for paragraph in paragraphs)
                bold_length = sum(
                    len(run.text)
                    for paragraph in paragraphs
                    for run in paragraph.runs
                    if run.bold is True
                )
                bold_ratio = bold_length / total_length if total_length else 0.0
                drafts.append(
                    BlockDraft(
                        kind=BlockKind.TABLE_CELL,
                        raw_text=text,
                        locator=DocumentLocator(
                            table=table_index,
                            row=row_index,
                            column=column_index,
                            section_path=[
                                section_titles[level] for level in sorted(section_titles)
                            ],
                        ),
                        style=BlockStyle(
                            bold_ratio=bold_ratio,
                            inline_spans=self._inline_spans(paragraphs),
                        ),
                    )
                )
        return drafts

    def _footnote_drafts(self, content: bytes) -> list[BlockDraft]:
        with ZipFile(BytesIO(content)) as archive:
            if "word/footnotes.xml" not in archive.namelist():
                return []
            xml = archive.read("word/footnotes.xml")
        parser = etree.XMLParser(
            resolve_entities=False,
            no_network=True,
            load_dtd=False,
            recover=False,
            huge_tree=False,
        )
        try:
            root = etree.fromstring(xml, parser=parser)
        except etree.XMLSyntaxError as exc:
            raise BriefParserError(
                ErrorCode.SOURCE_EXTRACTION_FAILED,
                "DOCX 脚注 XML 无法安全解析。",
            ) from exc

        drafts: list[BlockDraft] = []
        namespaces = {"w": _WORD_NAMESPACE}
        for footnote in root.findall("w:footnote", namespaces):
            raw_id = footnote.get(qn("w:id"))
            if raw_id is None or raw_id.startswith("-"):
                continue
            text_parts = footnote.xpath(".//w:t/text()", namespaces=namespaces)
            text = "".join(str(part) for part in text_parts)
            if not text.strip():
                continue
            footnote_index = (
                int(raw_id)
                if raw_id.isdigit() and int(raw_id) > 0
                else len(drafts) + 1
            )
            drafts.append(
                BlockDraft(
                    kind=BlockKind.FOOTNOTE,
                    raw_text=text,
                    locator=DocumentLocator(
                        footnote=footnote_index,
                        section_path=["footnotes"],
                    ),
                )
            )
        return drafts

    def _heading_level(self, paragraph: Paragraph) -> int | None:
        style = paragraph.style
        style_names = [style.name or "", style.style_id or ""] if style is not None else []
        for style_name in style_names:
            match = _HEADING_STYLE.search(style_name)
            if match is not None:
                return int(match.group(1))
        return None

    def _list_level(self, paragraph: Paragraph) -> int | None:
        paragraph_properties = paragraph._p.pPr
        if paragraph_properties is not None and paragraph_properties.numPr is not None:
            level = paragraph_properties.numPr.ilvl  # type: ignore[attr-defined]
            return 0 if level is None or level.val is None else int(level.val)
        style = paragraph.style
        style_name = "" if style is None else f"{style.name} {style.style_id}".lower()
        if "list" in style_name or "列表" in style_name:
            return 0
        return None

    def _bold_ratio(self, paragraph: Paragraph) -> float:
        total_length = sum(len(run.text) for run in paragraph.runs)
        if total_length == 0:
            return 0.0
        bold_length = sum(len(run.text) for run in paragraph.runs if run.bold is True)
        return bold_length / total_length

    def _paragraph_style(self, paragraph: Paragraph) -> BlockStyle:
        return BlockStyle(
            heading_level=self._heading_level(paragraph),
            list_level=self._list_level(paragraph),
            bold_ratio=self._bold_ratio(paragraph),
            inline_spans=self._inline_spans([paragraph]),
        )

    def _inline_spans(self, paragraphs: list[Paragraph]) -> list[InlineStyleSpan]:
        spans: list[InlineStyleSpan] = []
        offset = 0
        for paragraph_index, paragraph in enumerate(paragraphs):
            for run in paragraph.runs:
                if not run.text:
                    continue
                start = offset
                end = start + len(run.text)
                style_name = None if run.style is None else run.style.name
                spans.append(
                    InlineStyleSpan(
                        start=start,
                        end=end,
                        bold=run.bold is True,
                        italic=run.italic is True,
                        underline=run.underline is True,
                        style_name=style_name,
                    )
                )
                offset = end
            if paragraph_index < len(paragraphs) - 1:
                offset += 1
        return spans

    def _has_page_break(self, paragraph: Paragraph) -> bool:
        page_break_tag = qn("w:lastRenderedPageBreak")
        break_tag = qn("w:br")
        break_type = qn("w:type")
        for element in paragraph._p.iter():
            if element.tag == page_break_tag:
                return True
            if element.tag == break_tag and element.get(break_type) == "page":
                return True
        return False
