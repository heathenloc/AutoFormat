"""Real Brief requirements are extracted; the category checklist is only a guide."""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document

from brief_parser.adapters.clause_segmenter import ClauseSegmenter
from brief_parser.adapters.document_ingestion import DocumentIngestor
from brief_parser.adapters.evidence_index import EvidenceIndexBuilder
from brief_parser.adapters.input_security import FileSecurityValidator
from brief_parser.application.service import ParseServiceRequest
from brief_parser.bootstrap import build_service
from brief_parser.compilers.academic_common import AcademicCommonFormatExtractor
from brief_parser.config import Settings
from brief_parser.domain.requirement_ir import RequirementIR
from tests.factories import requirement_ir_payload


CATEGORY_CHECKLIST = (
    "摘要：字体，字号",
    "目录：一号标题字体，二级标题字体，三级标题字体",
    "正文：",
    "一级标题：位置（居中），字体，是否加粗，字号",
    "二级标题：位置（居左），字体，是否加粗，字号",
    "正文：横排，通栏，字体，字号",
    "图示：字体，字号",
    "表格：边框样式，表头字体字号，内容字体字号",
    "参考文献位置，字体，字号，是否加粗",
    "英文字母字体，字号",
)
PROJECT_ROOT = Path(__file__).resolve().parents[4]


def _ir_from_docx(tmp_path: Path, lines: tuple[str, ...]) -> RequirementIR:
    path = tmp_path / "brief.docx"
    document = Document()
    for line in lines:
        document.add_paragraph(line)
    document.save(path)
    source = FileSecurityValidator().secure_file(
        path, source_id="src_001", client_source_id="school-brief",
    )
    extracted = DocumentIngestor().ingest(source)
    segments = ClauseSegmenter().segment(extracted)
    evidence = EvidenceIndexBuilder().build(extracted, segments)
    payload = requirement_ir_payload()
    payload["domain_pack"] = {"id": "academic_paper", "version": "1.0.0"}
    payload["source_manifest"] = [extracted.source_manifest.model_dump(mode="json")]
    payload["evidence_index"] = {
        key: value.model_dump(mode="json") for key, value in evidence.items()
    }
    payload["requirements"] = []
    payload["facts"] = []
    payload["objectives"] = []
    payload["references"] = []
    payload["conflicts"] = []
    payload["unresolved_items"] = []
    payload["unclassified_content"] = []
    return RequirementIR.model_validate(payload)


def test_category_checklist_does_not_invent_unspecified_values(tmp_path: Path) -> None:
    ir = _ir_from_docx(tmp_path, CATEGORY_CHECKLIST)
    ruleset = AcademicCommonFormatExtractor().extract(ir)
    assert not any(rule.property_id in {"font.family", "font.size", "font.bold",
                                        "table.border_style"} for rule in ruleset.rules)
    assert ruleset.issues
    assert any("标题级别" in issue.reason for issue in ruleset.issues)
    assert {issue.property_id for issue in ruleset.issues} >= {
        "font.family", "font.size", "font.bold", "table.border_style",
    }
    assert all(issue.target is not None and issue.original_text
               for issue in ruleset.issues)
    assert all(issue.evidence_ids[0] in ir.evidence_index for issue in ruleset.issues)
    assert all(rule.evidence_refs[0] == ir.evidence_index[rule.evidence_refs[0].evidence_id]
               for rule in ruleset.rules)


@pytest.mark.parametrize(("brief_line", "expected"), (
    ("摘要的中文字体为宋体，字号为小四号。", {
        ("abstract", "font.family", "chinese"): "宋体",
        ("abstract", "font.size", "all"): 12,
    }),
    ("目录中的一级标题字体设为黑体，二级标题采用宋体，三级标题使用楷体。", {
        ("toc_heading_1", "font.family", "all"): "黑体",
        ("toc_heading_2", "font.family", "all"): "宋体",
        ("toc_heading_3", "font.family", "all"): "楷体",
    }),
    ("一级标题应居中排列，字体为黑体，字号小三号，加粗。", {
        ("heading_1", "paragraph.alignment", "all"): "center",
        ("heading_1", "font.family", "all"): "黑体",
        ("heading_1", "font.size", "all"): 15,
        ("heading_1", "font.bold", "all"): True,
    }),
    ("二级标题左对齐，采用宋体四号字，不加粗。", {
        ("heading_2", "paragraph.alignment", "all"): "left",
        ("heading_2", "font.family", "all"): "宋体",
        ("heading_2", "font.size", "all"): 14,
        ("heading_2", "font.bold", "all"): False,
    }),
    ("正文须横排通栏，字体宋体，小四号。", {
        ("body", "section.text_direction", "all"): "horizontal",
        ("body", "section.column_count", "all"): 1,
        ("body", "font.family", "all"): "宋体",
        ("body", "font.size", "all"): 12,
    }),
    ("图示说明文字统一使用宋体五号。", {
        ("figure_text", "font.family", "all"): "宋体",
        ("figure_text", "font.size", "all"): 10.5,
    }),
    ("表格边框采用三线表，表头字体为黑体五号，表内文字使用宋体小五号。", {
        ("table", "table.border_style", "all"): "three_line",
        ("table_header", "font.family", "all"): "黑体",
        ("table_header", "font.size", "all"): 10.5,
        ("table_body", "font.family", "all"): "宋体",
        ("table_body", "font.size", "all"): 9,
    }),
    ("参考文献标题居中，字体为黑体，字号三号，加粗。", {
        ("bibliography_heading", "paragraph.alignment", "all"): "center",
        ("bibliography_heading", "font.family", "all"): "黑体",
        ("bibliography_heading", "font.size", "all"): 16,
        ("bibliography_heading", "font.bold", "all"): True,
    }),
    ("英文字母统一使用 Times New Roman 字体，字号为 12pt。", {
        ("latin_text", "font.family", "latin"): "Times New Roman",
        ("latin_text", "font.size", "latin"): 12,
    }),
    ("论文摘要部分要求字体为宋体，字号为小四号。", {
        ("abstract", "font.family", "all"): "宋体",
        ("abstract", "font.size", "all"): 12,
    }),
    ("参考文献应置于正文之后。", {
        ("bibliography", "document.position", "all"): "after_body",
    }),
    ("图题采用黑体五号，置于图下方。", {
        ("figure_caption", "font.family", "all"): "黑体",
        ("figure_caption", "font.size", "all"): 10.5,
        ("figure_caption", "caption.position", "all"): "below",
    }),
    ("英文及数字采用 Times New Roman 字体，字号为 12 磅。", {
        ("latin_digits", "font.family", "latin_digits"): "Times New Roman",
        ("latin_digits", "font.size", "latin_digits"): 12,
    }),
    ("表头文字使用黑体五号。", {
        ("table_header", "font.family", "all"): "黑体",
        ("table_header", "font.size", "all"): 10.5,
    }),
    ("表格：表头黑体五号，内容宋体五号，边框三线表。", {
        ("table_header", "font.family", "all"): "黑体",
        ("table_header", "font.size", "all"): 10.5,
        ("table_body", "font.family", "all"): "宋体",
        ("table_body", "font.size", "all"): 10.5,
        ("table", "table.border_style", "all"): "three_line",
    }),
    ("摘要字体：宋体；字号：小四号。", {
        ("abstract", "font.family", "all"): "宋体",
        ("abstract", "font.size", "all"): 12,
    }),
    ("表格边框样式：三线表。", {
        ("table", "table.border_style", "all"): "three_line",
    }),
    ("参考文献位置：正文之后。", {
        ("bibliography", "document.position", "all"): "after_body",
    }),
))
def test_real_brief_phrasings_cover_common_categories(
    tmp_path: Path, brief_line: str, expected: dict[tuple[str, str, str], object],
) -> None:
    ruleset = AcademicCommonFormatExtractor().extract(_ir_from_docx(tmp_path, (brief_line,)))
    actual = {(rule.target.semantic_role, rule.property_id, rule.character_scope):
              rule.value.normalized for rule in ruleset.rules if rule.status == "verified"}
    assert actual == expected
    assert not ruleset.issues


def test_followup_clauses_keep_target_and_both_evidence_refs(tmp_path: Path) -> None:
    ir = _ir_from_docx(tmp_path, ("摘要字体：宋体；字号：小四号。",))
    ruleset = AcademicCommonFormatExtractor().extract(ir)
    size = next(rule for rule in ruleset.rules if rule.property_id == "font.size")
    assert size.target.semantic_role == "abstract"
    assert len(size.evidence_refs) == 2
    assert all(ref == ir.evidence_index[ref.evidence_id] for ref in size.evidence_refs)


def test_followup_context_does_not_cross_paragraphs(tmp_path: Path) -> None:
    ir = _ir_from_docx(tmp_path, ("摘要：宋体。", "字号：五号。"))
    ruleset = AcademicCommonFormatExtractor().extract(ir)
    assert not any(rule.property_id == "font.size" for rule in ruleset.rules)


def test_toc_heading_clauses_keep_toc_context(tmp_path: Path) -> None:
    ir = _ir_from_docx(tmp_path, (
        "目录：一级标题黑体；二级标题宋体；三级标题楷体。",
    ))
    ruleset = AcademicCommonFormatExtractor().extract(ir)
    assert {(rule.target.semantic_role, rule.value.normalized) for rule in ruleset.rules} == {
        ("toc_heading_1", "黑体"), ("toc_heading_2", "宋体"),
        ("toc_heading_3", "楷体"),
    }
    assert all(len(rule.evidence_refs) == 2 for rule in ruleset.rules[1:])


def test_filled_values_are_atomic_and_conflicts_are_blocked(tmp_path: Path) -> None:
    ir = _ir_from_docx(tmp_path, (
        "一级标题：居中、黑体、小三号、加粗",
        "一级标题：左对齐",
        "正文：宋体、小四号、1.5倍行距、首行缩进2字符",
        "表格：三线表，表头黑体小四号",
        "页面：页边距2.5cm",
        "纸张：A4纸",
        "页码：阿拉伯数字",
        "示例：一级标题居右，仅供参考",
    ))
    ruleset = AcademicCommonFormatExtractor().extract(ir)
    heading_alignment = [rule for rule in ruleset.rules
                         if rule.target.semantic_role == "heading_1"
                         and rule.property_id == "paragraph.alignment"]
    assert len(heading_alignment) == 2
    assert {rule.status for rule in heading_alignment} == {"pending_conflict"}
    values = {(r.target.semantic_role, r.property_id): r.value.normalized
              for r in ruleset.rules if r.status == "verified"}
    assert values[("heading_1", "font.family")] == "黑体"
    assert values[("heading_1", "font.size")] == 15
    assert values[("heading_1", "font.bold")] is True
    assert values[("body", "font.family")] == "宋体"
    assert values[("body", "font.size")] == 12
    assert values[("body", "paragraph.line_spacing")] == 1.5
    assert values[("body", "paragraph.first_line_indent")] == 2
    assert values[("table", "table.border_style")] == "three_line"
    assert values[("table_header", "font.family")] == "黑体"
    assert values[("table_header", "font.size")] == 12
    assert values[("page", "page.margin.all")] == 2.5
    assert values[("page", "page.paper_size")] == "A4"
    assert values[("page_number", "page_number.number_format")] == "arabic"


def test_category_targets_and_character_scope_remain_distinct(tmp_path: Path) -> None:
    ir = _ir_from_docx(tmp_path, (
        "摘要：宋体、小四号",
        "目录：一级标题黑体，二级标题宋体",
        "参考文献：居左、宋体、五号、不加粗",
        "英文字母：Times New Roman、12pt",
        "表格：内容宋体五号",
    ))
    ruleset = AcademicCommonFormatExtractor().extract(ir)
    values = {(r.target.semantic_role, r.property_id): r.value.normalized
              for r in ruleset.rules if any(f.rule_id == r.rule_id and f.status == "verified"
                                           for f in ruleset.fragments)}
    assert values[("abstract", "font.size")] == 12
    assert values[("toc_heading_1", "font.family")] == "黑体"
    assert values[("toc_heading_2", "font.family")] == "宋体"
    assert values[("bibliography", "font.bold")] is False
    assert values[("latin_text", "font.family")] == "Times New Roman"
    assert values[("latin_text", "font.size")] == 12
    assert values[("table_body", "font.size")] == 10.5
    assert {rule.character_scope for rule in ruleset.rules
            if rule.target.semantic_role == "latin_text"} == {"latin"}


def test_conditional_values_do_not_become_global_rules(tmp_path: Path) -> None:
    ir = _ir_from_docx(tmp_path, (
        "正文：纸质版使用宋体",
        "正文：电子版使用黑体",
    ))
    ruleset = AcademicCommonFormatExtractor().extract(ir)
    assert len(ruleset.rules) == 2
    assert {rule.status for rule in ruleset.rules} == {"pending_condition"}
    assert {rule.value.normalized for rule in ruleset.rules} == {"宋体", "黑体"}
    assert not ruleset.conflicts


def test_alternatives_are_not_reduced_to_first_recognized_value(tmp_path: Path) -> None:
    ir = _ir_from_docx(tmp_path, (
        "正文：宋体或其他字体",
        "正文：行距1.5倍或2倍",
        "页面：A4或A3",
        "页码：小写罗马数字或阿拉伯数字",
    ))
    ruleset = AcademicCommonFormatExtractor().extract(ir)
    assert ruleset.rules == []
    assert {issue.property_id for issue in ruleset.issues} >= {
        "font.family", "paragraph.line_spacing", "page.paper_size",
        "page_number.number_format",
    }


async def test_docx_parse_service_includes_editor_handoff_without_model_rules(
    tmp_path: Path,
) -> None:
    path = tmp_path / "common-brief.docx"
    document = Document()
    for line in ("摘要：宋体、小四号", "一级标题：居中、黑体、小三号、加粗"):
        document.add_paragraph(line)
    document.save(path)
    service = build_service(Settings(
        environment="test", provider="fake", parser_isolation_mode="in_process",
    ))
    bundle = await service.parse_academic(ParseServiceRequest(
        request_id="req_common_format", domain_pack_id="academic_paper",
        requested_task_types=("format_academic_paper",), source_paths=(path,),
    ))
    assert bundle.bundle_version == "academic-parse-bundle/1.6"
    assert bundle.format_ruleset is not None
    assert {rule.target.semantic_role for rule in bundle.format_ruleset.rules
            if rule.status == "verified"} == {"abstract", "heading_1"}
    assert bundle.requirement_ir.evidence_index


async def test_project_template_distinguishes_toc_from_body_headings() -> None:
    service = build_service(Settings(
        environment="test", provider="fake", parser_isolation_mode="in_process",
    ))
    bundle = await service.parse_academic(ParseServiceRequest(
        request_id="req_template_format", domain_pack_id="academic_paper",
        requested_task_types=("format_academic_paper",),
        source_paths=(PROJECT_ROOT / "TestDocuments" /
                      "材料学院毕业论文模板2026.docx",),
    ))
    assert bundle.format_ruleset is not None
    rules = {(rule.target.semantic_role, rule.property_id): rule
             for rule in bundle.format_ruleset.rules}
    assert rules[("toc", "toc.max_depth")].value.normalized == 3
    assert rules[("toc_heading_3", "font.size")].value.normalized == 12
    assert rules[("body", "font.size")].value.normalized == 12
    assert rules[("heading_1", "paragraph.alignment")].value.normalized == "center"
    assert rules[("latin_digits", "font.family")].value.normalized == "Times New Roman"
    assert rules[("latin_digits", "font.family")].character_scope == "latin_digits"
    assert rules[("latin_digits", "font.family")].status == "pending_conflict"
    assert rules[("table", "table.border_style")].value.normalized == "three_line"
    assert rules[("figure_caption", "caption.position")].value.normalized == "below"
    assert rules[("figure_caption", "font.size")].value.normalized == 10.5
    assert not any(rule.property_id.startswith("page.margin")
                   for rule in bundle.format_ruleset.rules)
