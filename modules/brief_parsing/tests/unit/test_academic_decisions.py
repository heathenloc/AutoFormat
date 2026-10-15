"""User decisions change typed projections without mutating source parsing."""

import asyncio
import json
from uuid import uuid4

import pytest
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from pydantic import ValidationError

from brief_parser.application.service import ParseServiceRequest, TextSourceInput
from brief_parser.bootstrap import build_service
from brief_parser.config import Settings
from brief_parser.application.academic_decisions import AcademicDecisionEngine, AcademicDecisionError, render_decision_snapshot
from brief_parser.domain.academic_decisions import AcademicDecisionRequest, AcademicDecisionSnapshot
from brief_parser.compilers.academic_scope import intersection, base_scope, axis


def engine(text=None, path=None):
    service = build_service(Settings(environment="test", provider="fake", parser_isolation_mode="in_process"))
    bundle = asyncio.run(service.parse_academic(ParseServiceRequest(request_id="decision-test",
        domain_pack_id="academic_paper", requested_task_types=("format_academic_paper",),
        text_sources=(TextSourceInput(text=text),) if text else (), source_paths=(path,) if path else ())))
    return AcademicDecisionEngine(bundle, task_id="a" * 32, initial_version_id="b" * 32)


def answer(parser, snapshot, question, data, supersedes=None):
    return parser.submit(snapshot, AcademicDecisionRequest(request_id=uuid4().hex, base_version_id="b" * 32,
        expected_snapshot_id=snapshot.snapshot_id, question_id=question.question_id, answer=data,
        supersedes_decision_id=supersedes), decision_id=uuid4().hex, created_at="2026-10-03T01:00:00+00:00")


def question(snapshot, kind):
    return next(item for item in snapshot.questions if item.kind == kind and item.status in {"pending", "deferred"})


def test_conflict_selection_preserves_nonintersecting_scope_and_source():
    parser = engine("正文：宋体。\n英文字母：Times New Roman。")
    source = parser.bundle.model_dump_json()
    initial = parser.compile()
    conflict = question(initial, "conflict")
    winner = next(rule for rule in initial.rules if rule.value.normalized == "Times New Roman")
    resolved = answer(parser, initial, conflict, {"action": "select_rule", "selected_id": winner.rule_id})
    assert not resolved.conflicts and all(rule.status == "verified" for rule in resolved.rules)
    loser = next(rule for rule in resolved.rules if rule.value.normalized == "宋体")
    assert intersection(loser.semantics.scope, conflict.scope) is None
    assert loser.source_rule_ids and loser.decision_ids
    assert parser.bundle.model_dump_json() == source
    assert parser.compile(initial.events) == initial
    assert parser.compile(resolved.events) == resolved


@pytest.mark.parametrize("medium,channel,expected", [("print", None, "宋体"), ("electronic", "other", "黑体")])
def test_scenario_activates_only_matching_branch(medium, channel, expected):
    parser = engine("纸质版正文：宋体。\n电子版正文：黑体。")
    initial = parser.compile()
    result = answer(parser, initial, question(initial, "submission"), {"action": "set_submission", "submission": {"medium": medium, "channel": channel}})
    assert [rule.value.normalized for rule in result.rules] == [expected]
    assert result.rules[0].status == "verified" and result.rules[0].semantics.condition is None
    assert result.excluded_source_rule_ids


def test_upload_channel_still_requires_overlapping_conflict_choice():
    parser = engine("电子版正文：宋体。\n系统上传版正文：黑体。")
    initial = parser.compile()
    other = answer(parser, initial, question(initial, "submission"), {"action": "set_submission", "submission": {"medium": "electronic", "channel": "other"}})
    assert len(other.rules) == 1 and not other.conflicts
    uploaded = answer(parser, initial, question(initial, "submission"), {"action": "set_submission", "submission": {"medium": "electronic", "channel": "system_upload"}})
    assert uploaded.conflicts and question(uploaded, "conflict")


@pytest.mark.parametrize("applies", [True, False])
def test_unknown_condition_requires_explicit_applicability(applies):
    parser = engine("若需要，正文：黑体。")
    initial = parser.compile()
    resolved = answer(parser, initial, question(initial, "condition"), {"action": "confirm_condition", "applies": applies})
    assert bool(resolved.rules) == applies
    assert all(rule.status == "verified" for rule in resolved.rules)


@pytest.mark.parametrize("value,unit,valid", [(13, "pt", True), (16, "pt", False), (13, "cm", False), (0, "pt", False), (True, "pt", False)])
def test_interval_answer_must_be_typed_and_within_constraint(value, unit, valid):
    parser = engine("第一章正文：字号12-14pt。")
    initial = parser.compile()
    data = {"action": "clarify_format", "format": {"property_id": "font.size", "normalized": value, "unit": unit}}
    if valid:
        resolved = answer(parser, initial, question(initial, "value"), data)
        assert resolved.rules[0].semantics.scope.axes["chapter"].include == ["1"]
        assert resolved.rules[0].value.normalized == 13 and resolved.rules[0].status == "verified"
    else:
        with pytest.raises(AcademicDecisionError):
            answer(parser, initial, question(initial, "value"), data)


@pytest.mark.parametrize("text,axis_name,expected", [("第一章正文：字号。", "chapter", ["1"]), ("英文摘要：字号。", "object_language", ["english"]), ("正文：中文字体。", "characters", ["chinese"])])
def test_missing_value_preserves_narrow_object_scope(text, axis_name, expected):
    parser = engine(text)
    initial = parser.compile()
    pending = question(initial, "clarification")
    data = {"action": "clarify_format", "format": {"property_id": pending.property_id, "normalized": "宋体" if pending.property_id == "font.family" else 12, "unit": None if pending.property_id == "font.family" else "pt"}}
    resolved = answer(parser, initial, pending, data)
    assert resolved.rules[0].semantics.scope.axes[axis_name].include == expected
    assert resolved.rules[0].target.selector == "semantic_role_scope" or axis_name == "characters"


def test_defer_skip_and_ignore_never_invent_rule():
    parser = engine("正文：字号。")
    initial = parser.compile()
    deferred = answer(parser, initial, question(initial, "clarification"), {"action": "defer"})
    assert not deferred.rules and question(deferred, "clarification").status == "deferred"
    skipped = answer(parser, deferred, question(deferred, "clarification"), {"action": "skip"})
    assert not skipped.rules and skipped.issues
    ignored = answer(parser, initial, question(initial, "clarification"), {"action": "ignore"})
    assert not ignored.rules and not ignored.issues


def test_content_limits_never_become_questions():
    snapshot = engine("摘要不少于300字。\n论文应论证观点并说明创新点。").compile()
    assert not snapshot.rules and not snapshot.questions


def test_revise_scenario_invalidates_dependent_answer_and_keeps_history():
    parser = engine("电子版正文：宋体。\n系统上传版正文：黑体。\n纸质版正文：楷体。")
    initial = parser.compile()
    electronic = answer(parser, initial, question(initial, "submission"), {"action": "set_submission", "submission": {"medium": "electronic", "channel": "system_upload"}})
    old_context = electronic.events[-1].decision_id
    conflict = question(electronic, "conflict")
    resolved = answer(parser, electronic, conflict, {"action": "select_rule", "selected_id": conflict.rule_ids[0]})
    old_bytes = resolved.model_dump_json()
    reopened = next(item for item in parser.reopen_questions(resolved, old_context) if item.kind == "submission")
    revised = answer(parser, resolved, reopened, {"action": "set_submission", "submission": {"medium": "print"}}, old_context)
    assert [item.status for item in revised.decision_effects] == ["superseded", "invalidated", "applied"]
    assert [rule.value.normalized for rule in revised.rules] == ["楷体"]
    assert resolved.model_dump_json() == old_bytes


def test_cannot_supersede_unrelated_question():
    parser = engine("正文：字号。\n一级标题：字号。")
    initial = parser.compile()
    pending = [item for item in initial.questions if item.kind == "clarification"]
    skipped = answer(parser, initial, pending[0], {"action": "skip"})
    with pytest.raises(AcademicDecisionError, match="原决定"):
        answer(parser, skipped, pending[1], {"action": "ignore"}, skipped.events[-1].decision_id)


def template(tmp_path, mismatch=False):
    path = tmp_path / "template.docx"
    doc = Document()
    doc.styles["Normal"].font.name = "黑体"
    doc.styles["Normal"].font.size = Pt(13)
    doc.add_heading("毕业论文模板", 0)
    if mismatch:
        doc.add_paragraph("正文：宋体。")
    doc.add_heading("第一章 绪论", 1)
    p = doc.add_paragraph("在此填写论文正文。")
    p.add_run("模板正文样例").font.size = Pt(13)
    p.runs[-1].font.name = "黑体"
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            run.font.name = "黑体"
            run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "黑体")
    doc.save(path)
    return engine(path=path)


def test_template_requires_explicit_confirmation_and_has_xml_proof(tmp_path):
    parser = template(tmp_path)
    initial = parser.compile()
    # Resolve diagnostic questions first using explicit skip, without any default.
    while any(item.phase == "resolve_requirements" and item.status == "pending" for item in initial.questions):
        q = next(item for item in initial.questions if item.phase == "resolve_requirements" and item.status == "pending")
        initial = answer(parser, initial, q, {"action": "skip"})
    category = question(initial, "template_category")
    assert not any(rule.source_kind == "template_observation_confirmed" for rule in initial.rules)
    resolved = answer(parser, initial, category, {"action": "accept_template"})
    confirmed = [rule for rule in resolved.rules if rule.source_kind == "template_observation_confirmed"]
    assert confirmed and all(rule.observation_ids and rule.decision_ids and not rule.evidence_refs for rule in confirmed)
    assert all(rule.execution_support == "typed" or rule.status == "check_only" for rule in confirmed)
    assert "Brief 解析决策报告" in render_decision_snapshot(resolved)


def test_template_phase_blocked_until_conflict_handled(tmp_path):
    parser = template(tmp_path, mismatch=True)
    initial = parser.compile()
    mismatch = question(initial, "source_mismatch")
    blocked = next(item for item in initial.questions if item.kind == "template_category")
    assert blocked.status == "blocked"
    with pytest.raises(AcademicDecisionError):
        answer(parser, initial, blocked, {"action": "accept_template"})
    kept = answer(parser, initial, mismatch, {"action": "keep_brief"})
    assert not any(rule.source_kind == "template_observation_confirmed" for rule in kept.rules)
    selected = answer(parser, initial, mismatch, {"action": "select_template", "selected_id": mismatch.observations[0].observation_id})
    assert any(rule.source_kind == "template_observation_confirmed" for rule in selected.rules)


def test_snapshot_and_answer_extra_fields_cannot_bypass_validation():
    initial = engine("正文：宋体。").compile()
    payload = initial.model_dump(mode="json")
    payload["rules"][0]["value"]["normalized"] = "黑体"
    with pytest.raises(ValidationError):
        AcademicDecisionSnapshot.model_validate(payload)
    with pytest.raises(ValidationError):
        AcademicDecisionRequest(request_id="a" * 32, base_version_id="b" * 32, expected_snapshot_id=initial.snapshot_id,
            question_id="bogus", answer={"action": "skip", "selected_id": "bogus"})


def test_registered_schemas_accept_actual_snapshot_and_queue():
    from jsonschema import Draft202012Validator
    initial = engine("正文：字号12-14pt。").compile()
    schema = AcademicDecisionSnapshot.model_json_schema()
    Draft202012Validator.check_schema(schema)
    Draft202012Validator(schema).validate(initial.model_dump(mode="json"))


@pytest.mark.parametrize("text", ["表格：无边框。", "正文：段前6pt。", "正文：段后12pt。", "正文：行距固定值20磅。", "页面：页边距2.5cm。"])
def test_current_brief_values_remain_valid_in_effective_contract(text):
    parser = engine(text)
    snapshot = parser.compile()
    assert snapshot.rules
    assert AcademicDecisionSnapshot.model_validate_json(snapshot.model_dump_json()) == snapshot


def test_resource_limit_fails_without_publishing_subset(monkeypatch):
    parser = engine("正文：宋体。\n正文：黑体。")
    import brief_parser.compilers.academic_scope as scopes
    monkeypatch.setattr(scopes, "MAX_RULES", 1)
    with pytest.raises(AcademicDecisionError, match="分析上限"):
        parser.compile()


def test_deferred_template_preconditions_are_not_default_acceptance(tmp_path):
    parser = template(tmp_path, mismatch=True)
    initial = parser.compile()
    pending = question(initial, "source_mismatch")
    deferred = answer(parser, initial, pending, {"action": "defer"})
    assert all(q.status == "blocked" for q in deferred.questions if q.kind == "template_category")
    assert not any(rule.source_kind == "template_observation_confirmed" for rule in deferred.rules)
    with pytest.raises(AcademicDecisionError):
        answer(parser, initial, pending, {"action": "select_template", "selected_id": "tplobs_unknown"})


def test_unknown_target_requires_explicit_typed_scope_and_cannot_expand():
    parser = engine("字体应为宋体。")
    initial = parser.compile()
    pending = question(initial, "source_review")
    data = {"action": "clarify_format", "format": {"property_id": "font.family", "normalized": "宋体", "semantic_role": "body"}}
    with pytest.raises(AcademicDecisionError, match="类型化范围"):
        answer(parser, initial, pending, data)
    data["format"]["scope"] = base_scope("body", "font.family").model_dump(mode="json")
    resolved = answer(parser, initial, pending, data)
    assert resolved.rules[0].status == "verified" and resolved.rules[0].target.semantic_role == "body"
    data["format"]["scope"]["axes"]["roles"] = axis(["heading_1"]).model_dump()
    with pytest.raises(AcademicDecisionError):
        answer(parser, initial, pending, data)
    data["format"]["scope"] = base_scope("body", "font.family").model_dump(mode="json")
    data["format"]["scope"]["axes"]["unknown_axis"] = axis(["unknown"]).model_dump()
    with pytest.raises(AcademicDecisionError):
        answer(parser, initial, pending, data)


def test_clarified_unknown_target_retains_submission_condition():
    parser = engine("电子版字体应为宋体。")
    initial = parser.compile()
    pending = question(initial, "source_review")
    data = {"action": "clarify_format", "format": {"property_id": "font.family", "normalized": "宋体", "semantic_role": "body",
            "scope": base_scope("body", "font.family").model_dump(mode="json")}}
    clarified = answer(parser, initial, pending, data)
    assert clarified.rules[0].semantics.scope.axes["medium"].include == ["electronic"]
    resolved = answer(parser, clarified, question(clarified, "submission"), {"action": "set_submission", "submission": {"medium": "print"}})
    assert not resolved.rules


def test_confirmed_invalid_template_value_remains_report_only(tmp_path):
    from brief_parser.domain.template_layout import TemplatePropertyGroup
    parser = template(tmp_path)
    observed = next(item for item in parser.bundle.template_layout.observations if item.property_id == "font.size" and item.resolution == "resolved")
    invalid = observed.model_copy(update={"value": 0})
    group = TemplatePropertyGroup(group_id="test_group", category="body", semantic_role=observed.semantic_role,
        property_id=observed.property_id, character_scope=observed.character_scope, status="needs_template_confirmation",
        observation_ids=[observed.observation_id], explicit_rule_ids=[], values=[0], unit="pt", reason="Test unsupported value", next_action="Report only")
    rule = parser._template_rule(group, invalid, "c" * 32)
    assert rule.execution_support == "report_only" and rule.status == "check_only"


@pytest.mark.parametrize("prefix", ["本章", "某章", "指定章节"])
def test_unknown_chapter_condition_requires_explicit_chapter(prefix):
    parser = engine(prefix + "正文：宋体。")
    initial = parser.compile()
    pending = question(initial, "condition")
    assert pending.requires_scope
    with pytest.raises(AcademicDecisionError, match="类型化范围"):
        answer(parser, initial, pending, {"action": "confirm_condition", "applies": True})
    with pytest.raises(AcademicDecisionError, match="明确章号"):
        answer(parser, initial, pending, {"action": "confirm_condition", "applies": True, "scope": pending.scope.model_dump(mode="json")})
    scope = pending.scope.model_dump(mode="json")
    scope["axes"]["chapter"] = axis(["2"]).model_dump()
    resolved = answer(parser, initial, pending, {"action": "confirm_condition", "applies": True, "scope": scope})
    assert resolved.rules[0].status == "verified"
    assert resolved.rules[0].semantics.scope.axes["chapter"].include == ["2"]
    assert resolved.rules[0].target.selector == "semantic_role_scope"


def test_unknown_exception_cannot_confirm_whole_original_scope():
    parser = engine("全文：除封面外使用宋体。")
    initial = parser.compile()
    pending = question(initial, "condition")
    assert pending.requires_scope
    with pytest.raises(AcademicDecisionError, match="排除范围"):
        answer(parser, initial, pending, {"action": "confirm_condition", "applies": True, "scope": pending.scope.model_dump(mode="json")})
