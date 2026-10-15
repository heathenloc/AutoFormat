"""Expression coverage and adversarial model-candidate proof checks."""

from __future__ import annotations

import pytest
from docx import Document
from pydantic import ValidationError

from brief_parser.compilers.academic_common import AcademicCommonFormatExtractor
from brief_parser.compilers.academic_semantic_extractor_v1_4 import AcademicCommonFormatExtractor as Frozen14
from brief_parser.domain.requirement_ir import RequirementIR, Requirement
from brief_parser.domain.academic_format import AcademicFormatRuleSet
from tests.unit.test_academic_common import _ir_from_docx


@pytest.mark.parametrize("text,role,prop,value,unit", [
    ("Body text: font size 12pt, justified.", "body", "font.size", 12, "pt"),
    ("Body text: font size 12pt, justified.", "body", "paragraph.alignment", "both", None),
    ('Body: font family "Source Serif 4".', "body", "font.family", "Source Serif 4", None),
    ("正文：字体为Garamond。", "body", "font.family", "Garamond", None),
    ("正文：字体为SimSun。", "body", "font.family", "宋体", None),
    ("Heading 5: bold, 14pt.", "heading_5", "font.bold", True, None),
    ("Level 9 heading: centered.", "heading_9", "paragraph.alignment", "center", None),
    ("第八级标题：斜体。", "heading_8", "font.italic", True, None),
    ("正文：不使用斜体。", "body", "font.italic", False, None),
    ("Footnotes: double-spaced.", "footnote", "paragraph.line_spacing", 2, "multiple"),
    ("Endnotes: single-spaced.", "endnote", "paragraph.line_spacing", 1, "multiple"),
    ("正文：双倍行距。", "body", "paragraph.line_spacing", 2, "multiple"),
    ("正文：段前0.5行。", "body", "paragraph.spacing_before", .5, "lines"),
    ("Body: spacing after 6pt.", "body", "paragraph.spacing_after", 6, "pt"),
    ("Body: hanging indent 2 characters.", "body", "paragraph.hanging_indent", 2, "characters"),
    ("页面：A3纸。", "page", "page.paper_size", "A3", None),
    ("Paper: Letter, landscape.", "page", "page.paper_size", "Letter", None),
    ("Paper: Letter, landscape.", "page", "page.orientation", "landscape", None),
    ("页面：上下边距25mm。", "page", "page.margin.top", 2.5, "cm"),
    ("页面：上下边距25mm。", "page", "page.margin.bottom", 2.5, "cm"),
    ("Page: left margin 3cm.", "page", "page.margin.left", 3, "cm"),
    ("页面：装订线0.5厘米。", "page", "page.margin.gutter", .5, "cm"),
    ("页面：页眉距顶端1.5厘米。", "page", "page.margin.header_distance", 1.5, "cm"),
    ("页码：大写罗马数字。", "page_number", "page_number.number_format", "upper_roman", None),
    ("页码：从1开始。", "page_number", "page_number.start", 1, None),
    ("目录：最多五级。", "toc", "toc.max_depth", 5, None),
    ("表题：置于表格上方。", "table_caption", "caption.position", "above", None),
    ("图题：位于图下方。", "figure_caption", "caption.position", "below", None),
    ("摘要标题：居中。", "abstract_heading", "paragraph.alignment", "center", None),
    ("目录标题：加粗。", "toc_heading", "font.bold", True, None),
    ("正文：与下段同页。", "body", "paragraph.keep_with_next", True, None),
    ("正文：取消段前分页。", "body", "paragraph.page_break_before", False, None),
    ("正文：禁止加粗。", "body", "font.bold", False, None),
    ("正文：不得使用斜体。", "body", "font.italic", False, None),
    ("Body: do not use bold.", "body", "font.bold", False, None),
    ("参 考 文 献（居中、小三号字、黑体加粗）", "bibliography", "font.bold", True, None),
    ("致    谢（居中、小三号字、黑体加粗）", "acknowledgements", "font.size", 15, "pt"),
])
def test_new_expressions_preserve_actual_evidence(tmp_path, text, role, prop, value, unit):
    ir = _ir_from_docx(tmp_path, (text,))
    ruleset = AcademicCommonFormatExtractor().extract(ir)
    matches = [rule for rule in ruleset.rules if rule.target.semantic_role == role and rule.property_id == prop]
    assert len(matches) == 1, (ruleset.rules, ruleset.issues)
    rule = matches[0]
    assert (rule.value.normalized, rule.value.unit, rule.status) == (value, unit, "verified")
    assert all(ref == ir.evidence_index[ref.evidence_id] for ref in rule.evidence_refs)
    assert rule.value.original_text in text


def candidate_ir(tmp_path, text="第一章正文：字号12pt。", **changes):
    ir = _ir_from_docx(tmp_path, (text,))
    ref = next(iter(ir.evidence_index.values()))
    payload = dict(id="req_model_size", semantic_key="body.font.size", category="body_text", label="正文的字号",
                   subject="body", modality="must", operator="eq", value={"type": "font_size", "amount": 12, "unit": "pt"},
                   scope={"applies_to": ["body"], "excludes": []}, priority="normal", machine_checkability="deterministic",
                   verification_status="verified", verification_basis=["exact_evidence"], evidence_refs=[ref.evidence_id],
                   origin="source", downstream_policy="enforce")
    payload.update(changes)
    return ir.model_copy(update={"requirements": [Requirement.model_validate(payload)]})


def test_model_projection_preserves_chapter_and_actual_requirement_id(tmp_path):
    ir = candidate_ir(tmp_path)
    rules = AcademicCommonFormatExtractor().extract(ir)
    assert len(rules.rules) == 1
    assert rules.candidate_reviews[0].status == "accepted"
    assert rules.rules[0].source_requirement_id == ir.requirements[0].id
    assert rules.rules[0].semantics.scope.axes["chapter"].include == ["1"]
    assert rules.rules[0].target.selector == "semantic_role_scope"
    assert rules.candidate_reviews[0].rule_ids == [rules.rules[0].rule_id]


def test_english_chapter_keeps_proven_restriction(tmp_path):
    ir = candidate_ir(tmp_path, "Chapter 2 body text: font size 12pt.")
    rules = AcademicCommonFormatExtractor().extract(ir)
    assert rules.candidate_reviews[0].status == "accepted"
    assert rules.rules[0].semantics.scope.axes["chapter"].include == ["2"]


@pytest.mark.parametrize("text,changes,reason", [
    ("正文：字号12pt。", {"value": {"type": "font_size", "amount": 14, "unit": "pt"}}, "value_scope_evidence_mismatch"),
    ("正文：字号12pt。", {"value": {"type": "font_size", "amount": True, "unit": "pt"}}, "invalid_typed_value"),
    ("正文：字号12pt。", {"value": {"type": "font_size", "amount": 12, "unit": None}}, "invalid_typed_value"),
    ("正文：字号12pt。", {"subject": "heading_1", "scope": {"applies_to": ["heading_1"], "excludes": []}}, "value_scope_evidence_mismatch"),
    ("正文：字号12pt。", {"scope": {"applies_to": ["body", "heading_1"], "excludes": []}}, "unsupported_model_scope"),
    ("正文：字号12pt。", {"scope": {"applies_to": ["body"], "excludes": ["abstract"]}}, "unsupported_model_scope"),
    ("正文：字号12pt。", {"origin": "user_revision"}, "source_not_explicit"),
    ("正文：字号12pt。", {"verification_status": "unresolved"}, "evidence_not_verified"),
    ("正文：字号12pt。", {"modality": "should", "downstream_policy": "prefer"}, "non_mandatory_candidate"),
    ("正文：字号12pt。", {"value": {"type": "font_size", "amount": 12, "unit": "pt"}, "semantic_key": "body.font.bold"}, "unsupported_property"),
    ("示例：正文：字号12pt。", {}, "source_not_explicit"),
    ("模板显示正文：字号12pt。", {}, "source_not_explicit"),
    ("正文：行距12pt。", {}, "value_scope_evidence_mismatch"),
])
def test_bad_model_values_and_authority_never_promote(tmp_path, text, changes, reason):
    ir = candidate_ir(tmp_path, text, **changes)
    result = AcademicCommonFormatExtractor().extract(ir)
    review = result.candidate_reviews[0]
    assert review.status in {"rejected", "pending"}
    assert review.reason_code == reason and not review.rule_ids
    assert not any(rule.source_requirement_id == "req_model_size" for rule in result.rules)


def test_duplicate_model_proposals_share_one_rule(tmp_path):
    ir = candidate_ir(tmp_path)
    duplicate = ir.requirements[0].model_copy(update={"id": "req_model_size_2"})
    ir = ir.model_copy(update={"requirements": [duplicate, ir.requirements[0]]})
    result = AcademicCommonFormatExtractor().extract(ir)
    assert len(result.rules) == 1
    assert [item.status for item in result.candidate_reviews] == ["accepted", "covered"]
    assert result.candidate_reviews[0].rule_ids == result.candidate_reviews[1].rule_ids


def test_model_does_not_remove_proven_submission_condition(tmp_path):
    ir = candidate_ir(tmp_path, "电子版正文：字号12pt。")
    result = AcademicCommonFormatExtractor().extract(ir)
    assert result.candidate_reviews[0].status == "accepted"
    assert result.rules[0].status == "pending_condition"
    assert result.rules[0].semantics.scope.axes["medium"].include == ["electronic"]


@pytest.mark.parametrize("text", ["正文：建议使用12pt。", "正文：字号12pt或14pt。", "正文：行距1.5。", "正文：字体。"])
def test_ambiguous_or_optional_expression_does_not_become_verified(tmp_path, text):
    result = AcademicCommonFormatExtractor().extract(_ir_from_docx(tmp_path, (text,)))
    assert not result.rules and result.issues


def test_count_candidate_is_excluded_even_when_it_points_at_format(tmp_path):
    ir = candidate_ir(tmp_path, semantic_key="abstract.word_count", label="摘要字数", value={"type": "word_count", "amount": 300, "unit": "字"})
    result = AcademicCommonFormatExtractor().extract(ir)
    assert not result.candidate_reviews
    assert len(result.rules) == 1


def test_millimeters_convert_and_bounds_remain_check_only(tmp_path):
    ir = candidate_ir(tmp_path, "正文：悬挂缩进25mm。", semantic_key="body.paragraph.hanging_indent",
                      value={"type": "hanging_indent", "amount": 25, "unit": "mm"})
    result = AcademicCommonFormatExtractor().extract(ir)
    assert result.candidate_reviews[0].status == "accepted"
    assert result.rules[0].value.normalized == 2.5 and result.rules[0].value.unit == "cm"
    ir = candidate_ir(tmp_path, "正文：字号不小于12pt。", operator="gte", value={"type": "font_size", "amount": {"min": 12}, "unit": "pt"})
    result = AcademicCommonFormatExtractor().extract(ir)
    assert result.candidate_reviews[0].status == "accepted" and result.rules[0].status == "check_only"
    invalid = candidate_ir(tmp_path, "正文：字号不小于1pt。", operator="gte", value={"type": "font_size", "amount": {"min": True}, "unit": "pt"})
    assert AcademicCommonFormatExtractor().extract(invalid).candidate_reviews[0].reason_code == "invalid_typed_value"


def test_candidate_limit_never_returns_a_subset(tmp_path):
    from brief_parser.errors import BriefParserError, ErrorCode
    ir = candidate_ir(tmp_path)
    ir = ir.model_copy(update={"requirements": [ir.requirements[0].model_copy(update={"id": f"req_{i}"}) for i in range(1025)]})
    with pytest.raises(BriefParserError) as error:
        AcademicCommonFormatExtractor().extract(ir)
    assert error.value.code == ErrorCode.SOURCE_PARSER_RESOURCE_LIMIT


@pytest.mark.parametrize("text", ["页面：禁止双栏。", "正文：不得不加粗。", "正文：字体为Arial或Calibri。"])
def test_negative_and_disjunctive_values_are_not_guessed(tmp_path, text):
    result = AcademicCommonFormatExtractor().extract(_ir_from_docx(tmp_path, (text,)))
    assert not result.rules and result.issues


def test_candidate_review_tampering_and_frozen_history(tmp_path):
    ir = candidate_ir(tmp_path, "正文：字号12pt。")
    result = AcademicCommonFormatExtractor().extract(ir)
    damaged = result.model_dump(mode="json")
    damaged["candidate_reviews"][0]["rule_ids"] = ["fmt_forged"]
    with pytest.raises(ValidationError, match="unknown rule"):
        AcademicFormatRuleSet.model_validate(damaged)
    old = Frozen14().extract(ir).model_dump(mode="json")
    assert old["schema_version"] == "academic-format-rule-set/1.4"
    assert "candidate_reviews" not in old
    assert all(rule["value"]["dictionary_version"] == "academic-common-format/1.4" for rule in old["rules"])
    assert not Frozen14().extract(_ir_from_docx(tmp_path, ("Body: font size 12pt.",))).rules


def test_explicit_section_heading_can_supply_object_and_evidence(tmp_path):
    from brief_parser.adapters.document_ingestion import DocumentIngestor
    from brief_parser.adapters.input_security import FileSecurityValidator
    from brief_parser.adapters.clause_segmenter import ClauseSegmenter
    from brief_parser.adapters.evidence_index import EvidenceIndexBuilder
    ir = _ir_from_docx(tmp_path, ("placeholder",))
    document = Document()
    document.add_heading("正文", level=1)
    document.add_paragraph("使用小四号字，行距为。")
    document.add_heading("章标题", level=1)
    document.add_paragraph("使用三号字，黑体，靠左对齐。")
    path = tmp_path / "section-brief.docx"
    document.save(path)
    source = DocumentIngestor().ingest(FileSecurityValidator().secure_file(path, source_id="src_001", client_source_id="brief"))
    index = EvidenceIndexBuilder().build(source, ClauseSegmenter().segment(source))
    ir = ir.model_copy(update={"source_manifest": [source.source_manifest], "evidence_index": dict(index)})
    result = AcademicCommonFormatExtractor().extract(ir)
    assert len(result.rules) == 1
    rule = result.rules[0]
    assert (rule.target.semantic_role, rule.property_id, rule.value.normalized) == ("body", "font.size", 12)
    assert "正文" in {ref.quote for ref in rule.evidence_refs}
    assert all(ref == index[ref.evidence_id] for ref in rule.evidence_refs)
    assert any(issue.property_id == "paragraph.line_spacing" for issue in result.issues)
    assert not Frozen14().extract(ir).rules
