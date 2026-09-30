from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK

from brief_parser.adapters.document_ingestion import DocumentIngestor
from brief_parser.adapters.input_security import FileSecurityValidator
from brief_parser.domain.document import BlockKind


def _write_docx(path: Path) -> None:
    document = Document()
    document.sections[0].header.paragraphs[0].text = "Synthetic Header"
    document.sections[0].footer.paragraphs[0].text = "Synthetic Footer"
    document.add_heading("Delivery Requirements", level=1)
    paragraph = document.add_paragraph()
    paragraph.add_run("Brand name must appear.").bold = True
    document.add_paragraph("Do not mention competitors.", style="List Bullet")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Duration"
    table.cell(0, 1).text = "60 seconds maximum"
    page_break = document.add_paragraph()
    page_break.add_run().add_break(WD_BREAK.PAGE)
    document.add_paragraph("Second page content.")
    document.save(path)


def test_docx_preserves_structure_and_locations(tmp_path: Path) -> None:
    path = tmp_path / "synthetic.docx"
    _write_docx(path)
    source = FileSecurityValidator().secure_file(
        path,
        source_id="src_docx",
        client_source_id="brief",
        language_hints=("en-US",),
    )

    result = DocumentIngestor().ingest(source)
    kinds = [block.kind for block in result.blocks]

    assert BlockKind.HEADING in kinds
    assert BlockKind.LIST_ITEM in kinds
    assert kinds.count(BlockKind.TABLE_CELL) == 2
    assert BlockKind.PAGE_BREAK in kinds
    assert BlockKind.HEADER in kinds
    assert BlockKind.FOOTER in kinds
    duration = next(block for block in result.blocks if block.raw_text == "Duration")
    assert (duration.locator.table, duration.locator.row, duration.locator.column) == (
        1,
        1,
        1,
    )
    bold = next(block for block in result.blocks if "Brand name" in block.raw_text)
    assert bold.style.bold_ratio == 1.0
    assert bold.style.inline_spans[0].bold is True
    assert (bold.style.inline_spans[0].start, bold.style.inline_spans[0].end) == (
        0,
        len(bold.raw_text),
    )
    assert result.source_manifest.parser_id == "docx-python-docx"
