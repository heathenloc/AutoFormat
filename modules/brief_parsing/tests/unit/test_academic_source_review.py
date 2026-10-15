"""Authority, scope exclusions and omissions are checked against source text."""

from __future__ import annotations

import json
from dataclasses import replace
from io import BytesIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document
from pydantic import ValidationError

from brief_parser.adapters.document_ingestion import DocumentIngestor
from brief_parser.adapters.fake_provider import FakeModelProvider
from brief_parser.adapters.input_security import FileSecurityValidator
from brief_parser.application.pipeline import BriefParsingPipeline
from brief_parser.application.service import ParseServiceRequest, TextSourceInput
from brief_parser.bootstrap import build_service
from brief_parser.compilers.academic_common import AcademicCommonFormatExtractor
from brief_parser.compilers.academic_common_v1_2 import AcademicCommonFormatExtractor as Legacy12
from brief_parser.compilers.academic_source_review import build_source_review
from brief_parser.config import Settings
from brief_parser.domain.academic_bundle import AcademicParseBundle
from brief_parser.domain.candidates import ExtractorOutput
from brief_parser.domain.enums import ModelRole
from brief_parser.domain.requirement_ir import DomainPackRef
from tests.e2e.test_pipeline import _duration_extractor, _pipeline, _request, _source
from tests.unit.test_academic_common import _ir_from_docx


def _review(tmp_path: Path, lines: tuple[str, ...]):
    ir = _ir_from_docx(tmp_path, lines)
    rules = AcademicCommonFormatExtractor().extract(ir)
    return ir, rules, build_source_review(ir, rules)


def test_example_region_does_not_authorize_following_sample_values(tmp_path: Path) -> None:
    ir, rules, review = _review(tmp_path, (
        "示例：", "正文：黑体，小三号。", "论文格式要求：", "正文：宋体，小四号。",
    ))
    assert {(rule.property_id, rule.value.normalized) for rule in rules.rules} == {
        ("font.family", "宋体"), ("font.size", 12),
    }
    assert review.classification_counts["example"] == 2
    assert not review.items
    # JSON object key ordering cannot change region boundaries or decisions.
    reordered = ir.model_copy(update={"evidence_index": dict(reversed(list(ir.evidence_index.items())))})
    assert build_source_review(reordered, AcademicCommonFormatExtractor().extract(reordered)) == review


def test_template_observation_and_uncertain_value_are_not_executable(tmp_path: Path) -> None:
    _, rules, review = _review(tmp_path, (
        "模板观察：正文默认宋体，小四号。", "宋体，小四号。",
    ))
    assert not rules.rules
    assert review.classification_counts == {"template_observation": 1, "uncertain": 1}
    assert {item.kind for item in review.items} == {"source_uncertain"}


def test_partial_property_extraction_is_reported_even_with_successful_rules(tmp_path: Path) -> None:
    ir, rules, review = _review(tmp_path, ("正文：宋体，段前自动间距。",))
    assert any(rule.property_id == "font.family" for rule in rules.rules)
    assert not any(rule.property_id == "font.size" and rule.value.normalized == 6 for rule in rules.rules)
    gaps = [item for item in review.items if item.kind == "suspected_omission"]
    assert {item.property_id for item in gaps} == {"paragraph.spacing_before"}
    assert ir.evidence_index[gaps[0].evidence_ids[0]].locator.paragraph == 1
    assert "段前" in ir.evidence_index[gaps[0].evidence_ids[0]].quote
    assert review.represented_format_property_count < review.detected_format_property_count


def test_conflict_and_missing_value_are_pending_rather_than_omissions(tmp_path: Path) -> None:
    _, rules, review = _review(tmp_path, ("正文：宋体。", "正文：黑体。", "摘要：字号。"))
    assert {rule.status for rule in rules.rules} == {"pending_conflict"}
    assert review.items and {item.kind for item in review.items} == {"recognized_pending"}


def test_explicitly_unsupported_property_is_not_reported_as_unnoticed(tmp_path: Path) -> None:
    _, rules, review = _review(tmp_path, ("正文：分栏。",))
    assert rules.issues[0].status == "unsupported"
    assert len(review.items) == 1 and review.items[0].kind == "unsupported_format"
    assert review.items[0].property_id == "section.column_count"


@pytest.mark.parametrize("text", (
    "摘要字数不少于300字。", "正文至少10000字。", "关键词3～5个。",
    "Abstract must contain 200-300 words.", "论文页数不得超过30页。",
    "摘要应包括研究目的、方法和结论。", "参考文献至少20篇，近5年文献占一半。",
))
def test_excluded_requirements_generate_no_rules_gaps_or_questions(tmp_path: Path, text: str) -> None:
    _, rules, review = _review(tmp_path, (text,))
    assert not rules.rules and not rules.issues and not review.items
    assert review.detected_format_property_count == 0
    assert review.classification_counts == {"out_of_scope": 1}


class CapturingProvider(FakeModelProvider):
    def __init__(self):
        super().__init__()
        self.inputs = []

    async def structured_generate(self, **kwargs):
        role = kwargs["role"]
        payload = json.loads(kwargs["messages"][-1].content)
        self.inputs.append((role, payload))
        if role == ModelRole.EXTRACTOR:
            response = dict(facts=[], objectives=[], requirements=[], references=[], candidate_conflicts=[], unresolved_items=[])
        else:
            response = {"clauses": [{"clause_id": clause["clause_id"], "classification": "context",
                                      "evidence_refs": [clause["evidence_id"]], "reason_code": "format_context"}
                                     for clause in payload["clauses"]]}
        self._responses[role].append(response)
        return await super().structured_generate(**kwargs)


@pytest.mark.parametrize("text", (
    "摘要字数不少于300字，字体为宋体，字号为小四。正文应包括研究方法。",
    "摘要字数不少于300字字体为宋体，字号为小四。正文应包括研究方法。",
    "摘要：字体为宋体，字号为小四，字数不少于300字。正文应包括研究方法。",
))
async def test_mixed_count_and_format_keeps_values_and_never_sends_count_to_model(text: str) -> None:
    source = _source(text)
    provider = CapturingProvider()
    result = await _pipeline(provider).run(replace(
        _request(source), domain_pack=DomainPackRef(id="academic_paper", version="1.1.0"),
        requested_task_types=("format_academic_paper",),
    ))
    assert provider.inputs
    for _, payload in provider.inputs:
        texts = " ".join(clause["text"] for clause in payload["clauses"])
        assert "300" not in texts and "研究方法" not in texts
        assert "字体为宋体" in texts
    rules = AcademicCommonFormatExtractor().extract(result.requirement_ir)
    assert {(rule.target.semantic_role, rule.property_id, rule.value.normalized) for rule in rules.rules} == {
        ("abstract", "font.family", "宋体"), ("abstract", "font.size", 12),
    }
    assert not build_source_review(result.requirement_ir, rules).items


async def test_all_excluded_source_never_calls_model_or_creates_pending_items() -> None:
    source = _source("论文至少10000字。摘要应包括研究目的和结论。")
    provider = CapturingProvider()
    result = await _pipeline(provider).run(replace(
        _request(source), domain_pack=DomainPackRef(id="academic_paper", version="1.1.0"),
        requested_task_types=("format_academic_paper",),
    ))
    assert not provider.inputs
    assert not result.requirement_ir.requirements and not result.requirement_ir.unresolved_items
    assert not result.requirement_ir.unclassified_content


def test_model_count_hallucination_is_rejected_at_output_boundary() -> None:
    output = _duration_extractor("ev_format")
    output["requirements"][0].update(semantic_key="abstract.word_count", category="word_count", label="摘要字数至少300字")
    assert not BriefParsingPipeline._filter_academic_output(ExtractorOutput.model_validate(output)).requirements


@pytest.mark.parametrize("text", ("正文：首行缩进2字。", "正文：首行缩进 2 字符。"))
def test_character_indentation_remains_formatting(tmp_path: Path, text: str) -> None:
    _, rules, review = _review(tmp_path, (text,))
    assert len(rules.rules) == 1 and rules.rules[0].property_id == "paragraph.first_line_indent"
    assert rules.rules[0].value.normalized == 2 and not review.items


def test_docx_reading_blindspots_are_reported_without_executing_embedded_data(tmp_path: Path) -> None:
    path = tmp_path / "images.docx"
    document = Document()
    document.add_paragraph("正文：宋体。")
    document.sections[0].header.paragraphs[0].text = "大学名称"
    document.save(path)
    buffer = BytesIO()
    with ZipFile(path) as original, ZipFile(buffer, "w", ZIP_DEFLATED) as output:
        for entry in original.infolist():
            output.writestr(entry, original.read(entry.filename))
        output.writestr("word/media/brief.png", b"unread image")
        output.writestr("word/embeddings/object.bin", b"unread object")
    path.write_bytes(buffer.getvalue())
    source = FileSecurityValidator().secure_file(path, source_id="src_001", client_source_id="school-brief")
    document = DocumentIngestor().ingest(source)
    assert {"DOCX_MEDIA_NOT_READ", "DOCX_EMBEDDED_OBJECTS_NOT_READ", "DOCX_HEADERS_FOOTERS_NOT_AUDITED"} <= set(document.source_manifest.warnings)
    ir = _ir_from_docx(tmp_path, ("正文：宋体。",))
    ir = ir.model_copy(update={"source_manifest": [document.source_manifest]})
    review = build_source_review(ir, AcademicCommonFormatExtractor().extract(ir))
    assert len([item for item in review.items if item.kind == "unread_source"]) == 3


async def test_review_tampering_rejected_and_historical_12_retains_wire_shape() -> None:
    service = build_service(Settings(environment="test", provider="fake", parser_isolation_mode="in_process"))
    bundle = await service.parse_academic(ParseServiceRequest(
        request_id="req_source_review", domain_pack_id="academic_paper", requested_task_types=("format_academic_paper",),
        text_sources=(TextSourceInput(text="正文：宋体，段前自动间距。"),),
    ))
    modified = bundle.model_dump(mode="python")
    modified["source_review"]["items"] = []
    with pytest.raises(ValidationError, match="source review differs"):
        AcademicParseBundle.model_validate(modified)
    previous = bundle.model_dump(mode="python")
    previous["bundle_version"] = "academic-parse-bundle/1.2"
    previous.pop("source_review")
    previous.pop("template_layout")
    previous["format_ruleset"] = Legacy12().extract(bundle.requirement_ir).model_dump(mode="python")
    historical = AcademicParseBundle.model_validate(previous)
    assert "source_review" not in historical.model_dump(mode="json")
    assert AcademicParseBundle.model_validate_json(historical.model_dump_json()) == historical
