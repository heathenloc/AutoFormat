"""Format scope, interval and conditional semantics, with protected history."""

from __future__ import annotations

import json
from pathlib import Path

import pytest
from pydantic import ValidationError
from jsonschema import Draft202012Validator

from brief_parser.compilers.academic_common import AcademicCommonFormatExtractor
from brief_parser.compilers.academic_common_v1_3 import AcademicCommonFormatExtractor as Legacy13
from brief_parser.compilers.academic_scope import axis, intersection, subtract_scope
from brief_parser.domain.academic_format import AcademicFormatRuleSet
from brief_parser.domain.academic_semantics import AcademicRuleScope
from brief_parser.domain.academic_bundle import AcademicParseBundle
from brief_parser.application.service import ParseServiceRequest, TextSourceInput
from brief_parser.bootstrap import build_service
from brief_parser.config import Settings
from tests.unit.test_academic_common import _ir_from_docx


def extract(tmp_path, *lines):
    return AcademicCommonFormatExtractor().extract(_ir_from_docx(tmp_path, lines))


def test_print_and_electronic_are_branches_with_preserved_values(tmp_path):
    rules = extract(tmp_path, "正文：纸质版使用宋体。", "正文：电子版使用黑体。")
    assert {rule.value.normalized for rule in rules.rules} == {"宋体", "黑体"}
    assert {rule.status for rule in rules.rules} == {"pending_condition"}
    assert not rules.conflicts and rules.relations[0].kind == "conditional_branch"
    assert all(rule.semantics.condition.resolution == "known" for rule in rules.rules)


def test_system_upload_is_subset_of_electronic_not_an_exclusive_branch(tmp_path):
    rules = extract(tmp_path, "电子版正文：宋体。", "系统上传版正文：黑体。")
    assert len(rules.conflicts) == 1
    assert rules.conflicts[0].intersection.axes["medium"].include == ["electronic"]
    assert rules.conflicts[0].intersection.axes["channel"].include == ["system_upload"]
    clear = [item for item in rules.fragments if item.rule_id == rules.rules[0].rule_id and not item.conflict_ids]
    assert clear and clear[0].scope.axes["channel"].exclude == ["system_upload"]


def test_unknown_condition_retains_value_and_conservatively_blocks_overlap(tmp_path):
    rules = extract(tmp_path, "正文：宋体。", "若需要，正文：黑体。")
    assert {rule.value.normalized for rule in rules.rules} == {"宋体", "黑体"}
    assert rules.conflicts[0].certainty == "potential"
    assert all(item.status == "pending_condition" for item in rules.fragments)


def test_characters_chapters_and_unaffected_properties_are_distinct(tmp_path):
    rules = extract(tmp_path, "正文：中文宋体，英文Times New Roman，数字Arial。", "正文：居中。", "第一章正文：居左。")
    fonts = [item for item in rules.rules if item.property_id == "font.family"]
    assert {item.character_scope for item in fonts} == {"chinese", "latin", "digits"}
    assert all(item.status == "verified" for item in fonts)
    assert len(rules.conflicts) == 1
    conflict = rules.conflicts[0]
    assert conflict.intersection.axes["chapter"].include == ["1"]
    body = next(item for item in rules.rules if item.value.normalized == "center")
    residual = next(item for item in rules.fragments if item.rule_id == body.rule_id and item.status == "verified")
    assert residual.scope.axes["chapter"].exclude == ["1"]
    assert next(item for item in rules.rules if item.value.normalized == "left").target.selector == "semantic_role_scope"


def test_global_font_and_script_font_need_explicit_conflict_handling(tmp_path):
    rules = extract(tmp_path, "正文：宋体。", "英文字母：Times New Roman。")
    assert rules.conflicts[0].intersection.axes["characters"].include == ["latin"]
    body = next(item for item in rules.rules if item.target.semantic_role == "body")
    clear = [item for item in rules.fragments if item.rule_id == body.rule_id and item.status == "verified"]
    assert clear and "latin" not in clear[0].scope.axes["characters"].include
    assert "chinese" in clear[0].scope.axes["characters"].include


@pytest.mark.parametrize("expression", ("字号小四至四号", "字号四号至小四", "字号12-14pt", "字号12磅～14磅", "字号小4号至4号"))
def test_size_intervals_never_pick_an_endpoint(tmp_path, expression):
    rules = extract(tmp_path, "正文：" + expression + "。")
    rule = rules.rules[0]
    assert rule.value.normalized == {"min": 12, "max": 14}
    assert rule.semantics.constraint.operator == "between" and rule.status == "check_only"


def test_compatible_interval_and_scalar_vs_disjoint_intervals(tmp_path):
    compatible = extract(tmp_path, "正文：字号12-14pt。", "正文：小四号。")
    assert not compatible.conflicts and compatible.relations[0].kind == "compatible"
    incompatible = extract(tmp_path, "正文：字号12-14pt。", "正文：字号15-16pt。")
    assert len(incompatible.conflicts) == 1
    assert all(rule.status == "pending_conflict" for rule in incompatible.rules)


def test_lower_upper_bounds_are_format_not_text_limits(tmp_path):
    rules = extract(tmp_path, "正文：字号不小于12磅。", "正文：字号不超过14磅。", "摘要不少于300字。")
    assert len(rules.rules) == 2 and not rules.conflicts
    assert {item.semantics.constraint.operator for item in rules.rules} == {"gte", "lte"}
    assert not rules.issues


@pytest.mark.parametrize("expression,operator", [("字号不得小于小四", "gte"), ("字号12pt及以上", "gte"), ("字号12pt以下", "lte")])
def test_bound_phrasings_do_not_fall_back_to_equality(tmp_path, expression, operator):
    rule = extract(tmp_path, "正文：" + expression + "。").rules[0]
    assert rule.semantics.constraint.operator == operator and rule.status == "check_only"


def test_compound_font_range_retains_other_format_properties(tmp_path):
    rules = extract(tmp_path, "正文：宋体字号12-14pt加粗。")
    assert {rule.property_id for rule in rules.rules} == {"font.family", "font.size", "font.bold"}
    assert next(rule for rule in rules.rules if rule.property_id == "font.size").status == "check_only"
    assert not rules.issues


def test_unsupported_strict_or_physical_bounds_are_not_fixed_values(tmp_path):
    rules = extract(tmp_path, "正文：字号大于12pt。", "页面：页边距不少于2.5cm。", "正文：段前12pt以上。")
    assert not rules.rules and rules.issues


def test_fixed_line_spacing_and_paragraph_points_are_not_font_sizes(tmp_path):
    rules = extract(tmp_path, "正文：行距固定值20磅，段前6磅，段后3磅，悬挂缩进5毫米。")
    values = {item.property_id: item.value.normalized for item in rules.rules}
    assert values == {"paragraph.line_spacing": {"mode": "exact", "amount": 20}, "paragraph.spacing_before": 6,
                      "paragraph.spacing_after": 3, "paragraph.hanging_indent": .5}


def test_regional_page_numbers_and_specific_table_ordinals(tmp_path):
    rules = extract(tmp_path, "正文页码：阿拉伯数字。", "前置部分页码：小写罗马数字。", "第2个表格：全框线。", "第3个表格：三线表。")
    assert len(rules.rules) == 4 and not rules.conflicts
    page_rules = [item for item in rules.rules if item.target.semantic_role == "page_number"]
    assert {item.semantics.scope.axes["region_zone"].include[0] for item in page_rules} == {"body", "front_matter"}
    tables = [item for item in rules.rules if item.target.semantic_role == "table"]
    assert {item.semantics.scope.axes["object_ordinal"].include[0] for item in tables} == {"2", "3"}


def test_front_matter_can_overlap_an_abstract_without_assuming_paper_position(tmp_path):
    rules = extract(tmp_path, "前置部分页码：小写罗马数字。", "摘要页码：阿拉伯数字。")
    assert len(rules.conflicts) == 1
    common = rules.conflicts[0].intersection.axes
    assert common["region_zone"].include == ["front_matter"] and common["regions"].include == ["abstract"]


def test_explicit_heading_exclusion_and_deduplication_keep_evidence(tmp_path):
    rules = extract(tmp_path, "除一级标题外，各级标题：宋体。", "一级标题：黑体。", "正文：宋体。", "正文：宋体。")
    assert not rules.conflicts
    broad = next(item for item in rules.rules if item.target.semantic_role == "heading")
    assert "heading_1" not in broad.semantics.scope.axes["roles"].include
    body = next(item for item in rules.rules if item.target.semantic_role == "body")
    assert len(body.evidence_refs) == 2


def test_uniform_and_specific_margins_conflict_only_on_that_side(tmp_path):
    rules = extract(tmp_path, "页面：页边距2.5cm。", "页面：上页边距3cm。")
    assert len(rules.conflicts) == 1
    assert rules.conflicts[0].intersection.axes["property_component"].include == ["top"]
    assert any(item.status == "verified" and item.scope.axes["property_component"].include == ["bottom", "left", "right"] for item in rules.fragments)


def test_exclusion_constraint_does_not_invent_a_replacement_font(tmp_path):
    rules = extract(tmp_path, "正文不得使用宋体。", "正文：黑体。")
    assert not rules.conflicts
    forbidden = next(item for item in rules.rules if item.semantics.constraint.operator == "ne")
    assert forbidden.value.normalized == "宋体" and forbidden.status == "check_only"


def test_branch_continues_to_following_size_and_changes_only_when_named(tmp_path):
    rules = extract(tmp_path, "正文：纸质版宋体，小四号；电子版黑体，四号。")
    assert len(rules.rules) == 4 and not rules.conflicts
    assert all(rule.status == "pending_condition" for rule in rules.rules)
    assert {(rule.property_id, rule.value.normalized, rule.semantics.scope.axes["medium"].include[0]) for rule in rules.rules} == {
        ("font.family", "宋体", "print"), ("font.size", 12, "print"),
        ("font.family", "黑体", "electronic"), ("font.size", 14, "electronic"),
    }


async def test_docx_submission_heading_constrains_following_requirements(tmp_path):
    from docx import Document
    path = tmp_path / "conditional.docx"
    document = Document()
    document.add_heading("纸质版格式要求", level=1)
    document.add_paragraph("正文：宋体，小四号。")
    document.add_heading("电子版格式要求", level=1)
    document.add_paragraph("正文：黑体，四号。")
    document.save(path)
    service = build_service(Settings(environment="test", provider="fake", parser_isolation_mode="in_process"))
    bundle = await service.parse_academic(ParseServiceRequest(request_id="req_condition_heading", domain_pack_id="academic_paper",
        requested_task_types=("format_academic_paper",), source_paths=(path,)))
    rules = bundle.format_ruleset
    assert len(rules.rules) == 4 and not rules.conflicts
    assert {rule.status for rule in rules.rules} == {"pending_condition"}
    assert {rule.semantics.scope.axes["medium"].include[0] for rule in rules.rules} == {"print", "electronic"}


@pytest.mark.parametrize("text", ("正文：字号不小于12。", "正文：字号不超过12号。", "正文：行距固定值0磅。", "正文：段前6-12pt。"))
def test_missing_units_or_invalid_lengths_do_not_become_font_values(tmp_path, text):
    rules = extract(tmp_path, text)
    assert not rules.rules
    assert rules.issues


@pytest.mark.parametrize("text", ("正文：字号0pt。", "正文：行距固定值" + "9" * 400 + "磅。"))
def test_invalid_numeric_values_are_reported_without_executable_rules(tmp_path, text):
    rules = extract(tmp_path, text)
    assert not rules.rules and rules.issues


def test_ambiguous_or_submission_condition_does_not_create_exclusive_branch(tmp_path):
    rules = extract(tmp_path, "纸质版或电子版正文：宋体。", "正文：黑体。")
    assert rules.rules[0].semantics.condition.resolution == "unknown"
    assert rules.conflicts[0].certainty == "potential"


def test_whole_table_and_toc_rules_overlap_their_component_text(tmp_path):
    rules = extract(tmp_path, "表格：宋体。", "表头文字：黑体。", "目录：小四号。", "目录：一级标题五号。")
    assert len(rules.conflicts) == 2
    assert {tuple(item.intersection.axes["roles"].include) for item in rules.conflicts} == {("table_header",), ("toc_heading_1",)}
    assert any(part.status == "verified" and "table_body" in part.scope.axes["roles"].include for part in rules.fragments)


def test_abstract_language_is_object_scope_and_heading_sublevels_are_preserved(tmp_path):
    rules = extract(tmp_path, "中文摘要：宋体，居中。", "英文摘要：Times New Roman，居左。",
                    "各级标题：一级标题黑体，二级标题宋体。", "英数字：Arial。")
    abstracts = [rule for rule in rules.rules if rule.target.semantic_role == "abstract"]
    assert len(abstracts) == 4
    assert {rule.semantics.scope.axes["object_language"].include[0] for rule in abstracts} == {"chinese", "english"}
    assert all(rule.target.selector == "semantic_role_scope" for rule in abstracts)
    assert any(relation.kind == "disjoint" for relation in rules.relations if set(relation.rule_ids) <= {rule.rule_id for rule in abstracts})
    assert {rule.target.semantic_role for rule in rules.rules if rule.target.semantic_role.startswith("heading_")} == {"heading_1", "heading_2"}
    digits = next(rule for rule in rules.rules if rule.target.semantic_role == "latin_digits")
    assert digits.semantics.scope.axes["characters"].include == ["digits", "latin"]


def test_negated_submission_label_is_not_assumed_to_be_positive(tmp_path):
    rules = extract(tmp_path, "非纸质版正文：宋体。")
    assert rules.rules[0].semantics.condition.resolution == "unknown"
    assert "medium" not in rules.rules[0].semantics.scope.axes


def test_fixed_line_spacing_and_minimum_share_physical_units(tmp_path):
    rules = extract(tmp_path, "正文：行距固定值20磅。", "正文：行距最小值18磅。")
    assert not rules.conflicts
    assert rules.relations[0].kind == "compatible"


def test_changed_constraint_or_unscoped_selector_is_rejected(tmp_path):
    rules = extract(tmp_path, "第一章正文：字号12-14pt。")
    changed = rules.model_dump(mode="python")
    changed["rules"][0]["semantics"]["constraint"]["upper"] = 16
    with pytest.raises(ValidationError, match="normalized value differs"):
        AcademicFormatRuleSet.model_validate(changed)
    changed = rules.model_dump(mode="python")
    changed["rules"][0]["target"]["selector"] = "semantic_role"
    with pytest.raises(ValidationError, match="scoped selector"):
        AcademicFormatRuleSet.model_validate(changed)


def test_scope_analysis_limits_fail_without_discarding_conflicts(tmp_path, monkeypatch):
    import brief_parser.compilers.academic_scope as scope
    from brief_parser.errors import BriefParserError, ErrorCode
    monkeypatch.setattr(scope, "MAX_ANALYSIS_ITEMS", 1)
    with pytest.raises(BriefParserError) as error:
        extract(tmp_path, "正文：宋体。", "正文：黑体。")
    assert error.value.code == ErrorCode.SOURCE_PARSER_RESOURCE_LIMIT


def test_scope_subtraction_is_complete_and_never_contains_blocked_points():
    whole = AcademicRuleScope(axes={"roles": axis(["body"]), "characters": axis(["chinese", "latin", "digits"])})
    blocked = AcademicRuleScope(axes={"roles": axis(["body"]), "characters": axis(["latin"]), "chapter": axis(["1"])})
    parts = subtract_scope(whole, blocked)
    assert parts and all(intersection(part, blocked) is None for part in parts)
    for chapter in ("1", "2", "unknown"):
        for char in ("chinese", "latin", "digits"):
            point = AcademicRuleScope(axes={"roles": axis(["body"]), "characters": axis([char]), "chapter": axis([chapter])})
            count = sum(intersection(point, part) is not None for part in parts)
            assert count == (0 if chapter == "1" and char == "latin" else 1)


def test_forged_fragments_or_verified_conflicts_fail_model_validation(tmp_path):
    rules = extract(tmp_path, "正文：宋体。", "正文：黑体。")
    changed = rules.model_dump(mode="python")
    changed["fragments"] = []
    with pytest.raises(ValidationError, match="scope analysis differs"):
        AcademicFormatRuleSet.model_validate(changed)
    changed = rules.model_dump(mode="python")
    changed["rules"][0]["status"] = "verified"
    with pytest.raises(ValidationError, match="status differs"):
        AcademicFormatRuleSet.model_validate(changed)


async def test_current_bundle_and_legacy_14_use_original_semantics(tmp_path):
    service = build_service(Settings(environment="test", provider="fake", parser_isolation_mode="in_process"))
    bundle = await service.parse_academic(ParseServiceRequest(request_id="req_semantics", domain_pack_id="academic_paper",
        requested_task_types=("format_academic_paper",), text_sources=(TextSourceInput(text="正文：宋体。"),)))
    assert bundle.bundle_version == "academic-parse-bundle/1.6"
    assert AcademicParseBundle.model_validate_json(bundle.model_dump_json()) == bundle
    from brief_parser.adapters.template_layout_v1_0 import DocxTemplateLayoutReader
    from brief_parser.adapters.input_security import FileSecurityValidator
    from brief_parser.compilers.template_layout_v1_0 import build_template_report
    from brief_parser.compilers.academic_source_review_v1_0 import build_source_review
    source = FileSecurityValidator().secure_text("正文：宋体。", source_id="src_001", client_source_id="brief")
    legacy_rules = Legacy13().extract(bundle.requirement_ir)
    old = bundle.model_dump(mode="python")
    old.update(bundle_version="academic-parse-bundle/1.4", format_ruleset=legacy_rules.model_dump(mode="python"),
               source_review=build_source_review(bundle.requirement_ir, legacy_rules).model_dump(mode="python"),
               template_layout=build_template_report(DocxTemplateLayoutReader().read((source,)), bundle.requirement_ir, legacy_rules).model_dump(mode="python"))
    historical = AcademicParseBundle.model_validate(old)
    assert "semantics" not in historical.model_dump(mode="json")["format_ruleset"]["rules"][0]
    assert AcademicParseBundle.model_validate_json(historical.model_dump_json()) == historical
    root = Path(__file__).resolve().parents[4]
    Draft202012Validator(json.loads((root / "contracts/academic-parse-bundle/academic-parse-bundle-1.4.schema.json").read_text(encoding="utf-8"))).validate(historical.model_dump(mode="json"))
    from brief_parser.compilers.academic_semantic_extractor_v1_4 import AcademicCommonFormatExtractor as Frozen14
    from brief_parser.compilers.academic_source_review_v1_1 import build_source_review as Review11
    old_rules = Frozen14().extract(bundle.requirement_ir)
    old15 = bundle.model_dump(mode="python")
    old15.update(bundle_version="academic-parse-bundle/1.5", format_ruleset=old_rules.model_dump(mode="python"),
                 source_review=Review11(bundle.requirement_ir, old_rules).model_dump(mode="python"))
    previous = AcademicParseBundle.model_validate(old15)
    assert AcademicParseBundle.model_validate_json(previous.model_dump_json()) == previous
    Draft202012Validator(json.loads((root / "contracts/academic-parse-bundle/academic-parse-bundle-1.5.schema.json").read_text(encoding="utf-8"))).validate(previous.model_dump(mode="json"))
