"""Local HTTP boundary: trusted rules, scope approval, execution and downloads."""

from __future__ import annotations

import hashlib
import json
import os
import shutil
import socket
import subprocess
import threading
import time
from pathlib import Path

import httpx
import pytest
from docx import Document
from fastapi.testclient import TestClient
from jsonschema import Draft202012Validator

from brief_parser.domain.academic_format import (
    AcademicFormatRule, AcademicFormatRuleSet, AcademicTarget, AcademicValue,
)
from brief_parser.domain.requirement_ir import (
    DomainPackRef, EvidenceReference, SourceLocator,
)

from easyformat.local_api import bind_local_api, create_local_api
from easyformat.preview import LibreOfficeRenderer
from easyformat.rule_snapshots import RuleSnapshotError, RuleSnapshotStore
from easyformat.target_planning import DotnetTargetIndexer
from easyformat.version_store import VersionStore, new_id
from easyformat.workflow import DotnetEditor


ROOT = Path(__file__).resolve().parents[2]
ORIGIN = "http://127.0.0.1:57843"
TOKEN = "local-api-test-session-token-with-32-characters"


@pytest.fixture(scope="module")
def engines() -> tuple[DotnetTargetIndexer, DotnetEditor]:
    dotnet = shutil.which("dotnet")
    if dotnet is None:
        pytest.skip(".NET 10 SDK is required for local API integration.")
    index_project = ROOT / "services" / "docx_engine" / "EasyFormat.DocxIndex"
    editor_project = ROOT / "services" / "docx_engine" / "EasyFormat.DocxEditor"
    for project in (index_project, editor_project):
        completed = subprocess.run(
            [dotnet, "build", str(project), "--no-restore", "-c", "Debug"],
            cwd=ROOT, capture_output=True, text=True, timeout=120, check=False,
            encoding="utf-8", errors="replace",
        )
        assert completed.returncode == 0, completed.stderr[-2000:]
    index_dll = index_project / "bin" / "Debug" / "net10.0" / "EasyFormat.DocxIndex.dll"
    editor_dll = editor_project / "bin" / "Debug" / "net10.0" / "EasyFormat.DocxEditor.dll"
    return (
        DotnetTargetIndexer(
            dotnet_path=Path(dotnet), index_dll=index_dll,
            expected_index_sha256=hashlib.sha256(index_dll.read_bytes()).hexdigest(),
        ),
        DotnetEditor(
            dotnet_path=Path(dotnet), editor_dll=editor_dll,
            expected_editor_sha256=hashlib.sha256(editor_dll.read_bytes()).hexdigest(),
        ),
    )


def _ruleset() -> AcademicFormatRuleSet:
    quote = "正文段落居中"
    evidence = EvidenceReference(
        evidence_id="evidence-body", source_id="brief-source", block_id="block-body",
        start=0, end=len(quote), quote=quote,
        quote_sha256=hashlib.sha256(quote.encode("utf-8")).hexdigest(),
        locator=SourceLocator(paragraph=1), match_status="exact",
    )
    rule = AcademicFormatRule(
        rule_id="fmt-body-center", source_requirement_id="req-body",
        source_kind="brief_explicit", source_artifact_id="brief-artifact",
        source_version="brief-artifact:v1", evidence_refs=[evidence],
        target=AcademicTarget(
            object_kind="paragraph", semantic_role="body", selector="semantic_role"),
        property_id="paragraph.alignment",
        value=AcademicValue(
            original_text="居中", normalized="center",
            dictionary_version="academic-alignment/1.0"),
        requested_scope="same_kind", effective_scope="same_kind",
        status="verified", reason="Explicit Brief evidence verified.",
    )
    return AcademicFormatRuleSet(
        source_artifact_id="brief-artifact", source_artifact_version=1,
        domain_pack=DomainPackRef(id="academic_paper", version="1.0.0"),
        rules=[rule], issues=[],
    )


def _setup(tmp_path: Path, engines):
    indexer, editor = engines
    source = tmp_path / "paper.docx"
    document = Document()
    document.add_paragraph("First body paragraph.")
    document.add_paragraph("Second body paragraph.")
    document.add_heading("Heading is excluded.", level=1)
    document.save(source)
    store = VersionStore(tmp_path / "runtime")
    v0 = store.create_task(source)
    snapshots = RuleSnapshotStore(store)
    snapshot = snapshots.register(
        task_id=v0.task_id, base_version_id=v0.version_id, ruleset=_ruleset())

    def renderer(work_dir: Path):
        executable = os.environ.get("EASYFORMAT_TEST_SOFFICE")
        expected_hash = os.environ.get("EASYFORMAT_TEST_SOFFICE_SHA256")
        if not executable or not expected_hash:
            raise AssertionError("Unexpected renderer call without integration configuration.")
        return LibreOfficeRenderer(
            soffice_path=Path(executable), expected_soffice_sha256=expected_hash,
            data_root=work_dir, timeout_seconds=60,
        )

    app = create_local_api(
        store=store, snapshots=snapshots, indexer=indexer, editor=editor,
        renderer_factory=renderer, session_token=TOKEN,
        allowed_origin=ORIGIN, font_fingerprint="fixture-fonts",
    )
    client = TestClient(app, base_url=ORIGIN)
    headers = {"X-EasyFormat-Session": TOKEN, "Origin": ORIGIN}
    prefix = f"/api/tasks/{v0.task_id}/versions/{v0.version_id}"
    return client, headers, prefix, store, v0, snapshots, snapshot


def test_rule_snapshot_is_immutable_and_detects_tampering(
    tmp_path: Path, engines,
) -> None:
    client, _, _, store, v0, snapshots, snapshot = _setup(tmp_path, engines)
    client.close()
    assert snapshots.load(v0.task_id).snapshot_id == snapshot.snapshot_id
    assert RuleSnapshotStore(store).load(v0.task_id).ruleset == _ruleset()
    altered_rule = _ruleset().rules[0].model_copy(update={"reason": "Changed reason."})
    altered = _ruleset().model_copy(update={"rules": [altered_rule]})
    with pytest.raises(RuleSnapshotError) as duplicate:
        snapshots.register(task_id=v0.task_id, base_version_id=v0.version_id,
                           ruleset=altered)
    assert duplicate.value.status == "blocked_conflict"
    path = store.root / "rule_snapshots" / f"{v0.task_id}.json"
    envelope = json.loads(path.read_text(encoding="utf-8"))
    envelope["ruleset"]["rules"][0]["value"]["normalized"] = "right"
    path.write_text(json.dumps(envelope, ensure_ascii=False), encoding="utf-8")
    with pytest.raises(RuleSnapshotError) as tampered:
        snapshots.load(v0.task_id)
    assert tampered.value.status == "validation_failed"


def test_local_server_binding_reserves_loopback_random_port(
    tmp_path: Path, engines,
) -> None:
    client, _, _, store, _, snapshots, _ = _setup(tmp_path, engines)
    client.close()
    indexer, editor = engines
    binding = bind_local_api(
        store=store, snapshots=snapshots, indexer=indexer, editor=editor,
        renderer_factory=lambda _: None, font_fingerprint="fixture-fonts",
    )
    try:
        address, port = binding.listener.getsockname()
        assert address == "127.0.0.1" and port > 0
        assert binding.url == f"http://127.0.0.1:{port}"
        assert len(binding.session_token) >= 32
        assert binding.listener.family == socket.AF_INET
    finally:
        binding.close()


def test_local_server_serves_reserved_socket_and_stops(
    tmp_path: Path, engines,
) -> None:
    client, _, _, store, v0, snapshots, _ = _setup(tmp_path, engines)
    client.close()
    indexer, editor = engines
    binding = bind_local_api(
        store=store, snapshots=snapshots, indexer=indexer, editor=editor,
        renderer_factory=lambda _: None, font_fingerprint="fixture-fonts",
    )
    thread = threading.Thread(target=binding.serve, daemon=True)
    thread.start()
    try:
        for _ in range(100):
            if binding.server is not None and binding.server.started:
                break
            time.sleep(0.05)
        assert binding.server is not None and binding.server.started
        with httpx.Client(timeout=3, trust_env=False) as network_client:
            response = network_client.get(
                binding.url + f"/api/tasks/{v0.task_id}",
                headers={"X-EasyFormat-Session": binding.session_token,
                         "Origin": binding.url},
            )
        assert response.status_code == 200
        assert response.json()["working_version_id"] == v0.version_id
    finally:
        binding.close()
        thread.join(timeout=5)
    assert not thread.is_alive()


def test_local_review_and_confirmation_require_session_and_fresh_review(
    tmp_path: Path, engines,
) -> None:
    client, headers, prefix, store, v0, _, snapshot = _setup(tmp_path, engines)
    review_url = prefix + "/body-alignment/review"
    assert client.get(review_url).status_code == 401
    assert client.get(review_url, headers={**headers, "Origin": "http://evil.test"}).status_code == 403
    assert client.get(review_url, headers={**headers, "Host": "evil.test"}).status_code == 403
    review_response = client.get(review_url, headers=headers)
    assert review_response.status_code == 200
    review = review_response.json()
    schema = json.loads((ROOT / "contracts" / "target-review-plan" /
                         "target-review-plan-1.0.schema.json").read_text(encoding="utf-8"))
    Draft202012Validator(schema).validate(review)
    assert review["rules_snapshot_id"] == snapshot.snapshot_id
    assert review["rule_reviews"][0]["estimated_affected_count"] is None
    ids = [item["object_id"] for item in review["rule_reviews"][0]["targets"]]
    assert len(ids) == 3
    payload = {
        "schema_version": "body-scope-confirmation/1.0",
        "request_id": new_id(),
        "expected_base_docx_sha256": v0.docx_sha256,
        "review_sha256": review["review_sha256"],
        "body_object_ids": ids[:2],
        "excluded_object_ids": ids[2:],
        "raw_answer": "前两段是正文，标题排除。",
    }
    contract = json.loads((ROOT / "contracts" / "body-scope-confirmation" /
                           "body-scope-confirmation-1.0.schema.json").read_text(encoding="utf-8"))
    Draft202012Validator(contract).validate(payload)
    confirm_url = prefix + "/body-alignment/confirm"
    assert client.post(confirm_url, headers=headers,
                       json={**payload, "value": "right"}).status_code == 422
    assert client.post(confirm_url, headers=headers,
                       json={**payload, "review_sha256": "0" * 64}).status_code == 409
    assert client.post(confirm_url, headers=headers,
                       json={**payload, "excluded_object_ids": []}).status_code == 422
    assert store.list_decisions(v0.task_id) == []
    response = client.post(confirm_url, headers=headers, json=payload)
    assert response.status_code == 200, response.text
    confirmed = response.json()
    plan_schema = json.loads((ROOT / "contracts" / "edit-plan" /
                              "edit-plan-1.0.schema.json").read_text(encoding="utf-8"))
    Draft202012Validator(plan_schema).validate(confirmed["plan"])
    assert confirmed["plan"]["estimated_affected_count"] == 2
    assert [item["object_id"] for item in confirmed["plan"]["operations"]] == ids[:2]
    assert len(store.list_decisions(v0.task_id)) == 1
    assert client.post(confirm_url, headers=headers, json=payload).json() == confirmed
    conflicting = {**payload, "body_object_ids": ids[:1],
                   "excluded_object_ids": ids[1:]}
    assert client.post(confirm_url, headers=headers, json=conflicting).status_code == 409
    assert len(store.list_decisions(v0.task_id)) == 1
    client.close()


def test_local_execute_rechecks_plan_and_serves_registered_artifacts(
    tmp_path: Path, engines,
) -> None:
    if not os.environ.get("EASYFORMAT_TEST_SOFFICE"):
        pytest.skip("Set verified LibreOffice path for API execution integration.")
    client, headers, prefix, store, v0, _, _ = _setup(tmp_path, engines)
    review = client.get(prefix + "/body-alignment/review", headers=headers).json()
    ids = [item["object_id"] for item in review["rule_reviews"][0]["targets"]]
    confirmation = client.post(prefix + "/body-alignment/confirm", headers=headers,
                               json={
                                   "schema_version": "body-scope-confirmation/1.0",
                                   "request_id": new_id(),
                                   "expected_base_docx_sha256": v0.docx_sha256,
                                   "review_sha256": review["review_sha256"],
                                   "body_object_ids": ids[:2],
                                   "excluded_object_ids": ids[2:],
                                   "raw_answer": "前两段是正文。",
                               }).json()
    execute_url = prefix + "/body-alignment/execute"
    request = {
        "schema_version": "body-plan-execution/1.0",
        "request_id": new_id(),
        "expected_base_docx_sha256": v0.docx_sha256,
        "decision_id": confirmation["decision_id"],
        "plan_id": confirmation["plan"]["plan_id"],
    }
    contract = json.loads((ROOT / "contracts" / "body-plan-execution" /
                           "body-plan-execution-1.0.schema.json").read_text(encoding="utf-8"))
    Draft202012Validator(contract).validate(request)
    assert client.post(execute_url, headers=headers,
                       json={**request, "plan_id": "0" * 64}).status_code == 409
    assert client.post(execute_url, headers=headers,
                       json={**request, "value": "right"}).status_code == 422
    assert len(store.list_versions(v0.task_id)) == 1
    response = client.post(execute_url, headers=headers, json=request)
    assert response.status_code == 200, response.text
    result = response.json()
    assert result["status"] == "ready" and result["version_id"]
    assert len(store.list_versions(v0.task_id)) == 2
    assert client.post(execute_url, headers=headers, json=request).json() == result
    task_response = client.get(f"/api/tasks/{v0.task_id}", headers=headers).json()
    assert task_response["working_version_id"] == result["version_id"]
    versions = client.get(f"/api/tasks/{v0.task_id}/versions", headers=headers).json()
    assert len(versions["versions"]) == 2
    assert "docx_relative_path" not in versions["versions"][1]
    artifacts = f"/api/tasks/{v0.task_id}/versions/{result['version_id']}/artifacts"
    assert client.get(artifacts + "/docx").status_code == 401
    assert client.get(artifacts + "/preview", headers=headers).content.startswith(b"%PDF")
    report = client.get(artifacts + "/report_json", headers=headers).json()
    assert report["summary"]["this_round_change_count"] == 2
    assert client.get(artifacts + "/unknown", headers=headers).status_code == 422
    client.close()
