"""Task recovery, trusted bootstrap and fail-closed launch configuration."""
import hashlib
import json
from pathlib import Path
import runpy

import pytest
from docx import Document
from fastapi.testclient import TestClient

from autoformat.desktop_host import DesktopConfig, font_inventory, load_config, parser_settings
from autoformat.feedback_contracts import contract_schema
from autoformat.task_catalog import TASK_CONTRACTS
from autoformat.local_api import bind_local_api, create_local_api
from autoformat.rule_snapshots import RuleSnapshotStore
from autoformat.task_catalog import list_tasks
from autoformat.version_store import VersionStore

ROOT = Path(__file__).resolve().parents[2]
TOKEN = "desktop-test-token-with-32-characters"
ORIGIN = "http://127.0.0.1:57843"


def test_home_schemas_match_models_and_bundled_copies():
    from jsonschema import Draft202012Validator
    for version, model in TASK_CONTRACTS.items():
        kind, number = version.split("/")
        name = f"{kind}-{number}.schema.json"
        schema = json.loads((ROOT / "contracts" / kind / name).read_text(encoding="utf-8"))
        assert schema == contract_schema(model)
        assert schema == json.loads((ROOT / "apps/api/src/autoformat/schemas" / name).read_text(encoding="utf-8"))
        Draft202012Validator.check_schema(schema)


@pytest.mark.parametrize("error_type", [ValueError, TimeoutError])
def test_startup_failures_do_not_print_dependency_secrets(tmp_path, monkeypatch, capsys, error_type):
    from autoformat import desktop_host
    monkeypatch.setattr("sys.argv", ["autoformat", "--config", "test.json", "--check"])
    monkeypatch.setattr(desktop_host, "load_config", lambda *_: None)
    def fail(*args, **kwargs): raise error_type("provider-secret-and-private-path")
    monkeypatch.setattr(desktop_host, "preflight", fail)
    with pytest.raises(SystemExit) as caught: desktop_host.main()
    assert caught.value.code == 2
    message = capsys.readouterr().err
    assert "provider-secret-and-private-path" not in message and "运行组件检查失败" in message


def test_catalog_is_session_protected_and_restores_origins(tmp_path):
    doc = Document(); doc.add_paragraph("Private paper text")
    path = tmp_path / "paper.docx"; doc.save(path)
    store = VersionStore(tmp_path / "store"); original = store.create_task(path)
    app = create_local_api(store=store, snapshots=RuleSnapshotStore(store), indexer=None, editor=None,
        renderer_factory=lambda _: None, font_fingerprint="test", session_token=TOKEN, allowed_origin=ORIGIN)
    with TestClient(app, base_url=ORIGIN) as client:
        for route in ["/api/tasks", "/api/application"]: assert client.get(route).status_code == 401
        headers = {"X-AutoFormat-Session": TOKEN}
        result = client.get("/api/tasks", headers=headers)
        assert result.json()["tasks"][0] == dict(task_id=original.task_id, created_at=original.created_at,
            working_version_id=original.version_id, working_sequence=0, version_count=1, has_initial_version=False)
        assert str(tmp_path) not in result.text and "Private paper text" not in result.text
        assert client.get("/api/tasks?before=0", headers=headers).status_code == 422
        assert client.get("/api/application", headers=headers).json()["upload_available"] is False
        assert client.get("/api/tasks", headers={**headers, "Origin": "https://evil.test"}).status_code == 403
    restored = VersionStore(store.root)
    assert list_tasks(restored).tasks[0].working_version_id == original.version_id


def test_catalog_cursor_has_no_duplicates_when_new_tasks_arrive(tmp_path):
    doc = Document(); doc.add_paragraph("Catalog fixture")
    path = tmp_path / "paper.docx"; doc.save(path)
    store = VersionStore(tmp_path / "store")
    for _ in range(52): store.create_task(path)
    first = list_tasks(store); assert len(first.tasks) == 50 and first.next_before
    store.create_task(path)
    second = list_tasks(store, first.next_before)
    assert len(second.tasks) == 2 and second.next_before is None
    assert not {t.task_id for t in first.tasks} & {t.task_id for t in second.tasks}


def test_home_session_url_contains_no_task_and_is_not_in_repr(tmp_path):
    store = VersionStore(tmp_path / "store")
    host = bind_local_api(store=store, snapshots=RuleSnapshotStore(store), indexer=None, editor=None,
        renderer_factory=lambda _: None, font_fingerprint="test")
    try:
        assert "#session=" in host.ui_url() and "&task=" not in host.ui_url()
        assert host.session_token not in repr(host)
        with pytest.raises(ValueError): host.ui_url("not-a-task")
    finally: host.close()


def test_real_parser_never_implicitly_uses_demo(monkeypatch):
    monkeypatch.delenv("BRIEF_PARSER_PROVIDER", raising=False)
    with pytest.raises(ValueError, match="不会自动降级"): parser_settings(False)
    assert parser_settings(True).provider == "fake"
    monkeypatch.setenv("BRIEF_PARSER_PROVIDER", "openai"); monkeypatch.setenv("OPENAI_API_KEY", "test-secret-never-logged")
    monkeypatch.setenv("BRIEF_PARSER_LOG_CONTENT", "true")
    result = parser_settings(False)
    assert not result.log_content and result.provider == "openai"
    assert "test-secret-never-logged" not in repr(result)


def config_data():
    return dict(schema_version="desktop-config/1.0", node="node.exe", dotnet="dotnet.exe", editor_dll="engine/editor.dll",
        editor_sha256="0"*64, soffice="soffice.com", soffice_sha256="1"*64, font_files=[])


def test_config_rejects_secret_fields_and_resolves_relative_paths(tmp_path):
    data = config_data()
    with pytest.raises(ValueError): DesktopConfig.model_validate_json(json.dumps({**data, "api_key": "must-not-be-here"}))
    path = tmp_path / "config.json"; path.write_text(json.dumps(data), encoding="utf-8")
    config = load_config(tmp_path, path)
    assert config.editor_dll == tmp_path / "engine/editor.dll"


def test_font_fingerprint_tracks_content_not_just_names(tmp_path):
    path = tmp_path / "test.ttf"; path.write_bytes(b"\x00\x01\x00\x00font-one")
    config = DesktopConfig.model_validate_json(json.dumps({**config_data(), "font_files": [str(path)]}))
    first = font_inventory(config)
    path.write_bytes(b"\x00\x01\x00\x00font-two")
    assert font_inventory(config)[0] != first[0] and first[1] == 1
    path.write_bytes(b"not a font")
    with pytest.raises(ValueError): font_inventory(config)


def test_builder_uses_current_dependency_declarations_instead_of_stale_editable_metadata(tmp_path, monkeypatch):
    from types import SimpleNamespace
    from importlib import metadata
    vendor = runpy.run_path(str(ROOT / "scripts/build_windows_bundle.py"))["vendor_python"]
    monkeypatch.setitem(vendor.__globals__, "ROOT", tmp_path)
    for relative in ("apps/api", "modules/brief_parsing"):
        directory = tmp_path / relative; directory.mkdir(parents=True)
        (directory / "pyproject.toml").write_text('[project]\ndependencies=["jsonschema==4.26.0"]\n')
    # The installed editable metadata deliberately lacks the newly added dependency.
    monkeypatch.setattr(metadata, "distribution", lambda name: SimpleNamespace(
        version="4.26.0", requires=[], files=[], locate_file=lambda _: tmp_path))
    assert "jsonschema" in {item["name"] for item in vendor(tmp_path / "python")}


@pytest.mark.parametrize("corruption", ["changed", "extra", "cache", "traversal", "missing"])
def test_bundle_verification_rejects_incomplete_or_changed_inputs(tmp_path, corruption):
    verify = runpy.run_path(str(ROOT / "desktop/launch.py"))["verify_bundle"]
    payload = tmp_path / "app.py"; payload.write_bytes(b"pass\n")
    manifest = dict(schema_version="windows-test-bundle/1.0", files={"app.py": hashlib.sha256(payload.read_bytes()).hexdigest()})
    manifest_path = tmp_path / "bundle-manifest.json"
    manifest_path.write_text(json.dumps(manifest))
    assert verify(tmp_path) == manifest
    if corruption == "changed": payload.write_bytes(b"changed")
    elif corruption == "extra": (tmp_path / "injected.py").write_text("pass")
    elif corruption == "cache":
        (tmp_path / "__pycache__").mkdir(); (tmp_path / "__pycache__/app.cpython-312.pyc").write_bytes(b"stale cache")
    elif corruption == "missing": payload.unlink()
    else:
        manifest["files"] = {"../outside.py": "0"*64}; manifest_path.write_text(json.dumps(manifest))
    with pytest.raises((ValueError, OSError)): verify(tmp_path)


@pytest.mark.parametrize("legacy_database,current_directory,expected", [
    (False, False, "AutoFormat"),
    (True, False, "EasyFormat"),
    (True, True, "AutoFormat"),
])
def test_default_store_reopens_legacy_tasks_without_moving_them(
    tmp_path, monkeypatch, legacy_database, current_directory, expected,
):
    from autoformat.desktop_host import default_store_path
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path))
    legacy = tmp_path / "EasyFormat/tasks"
    if legacy_database:
        legacy.mkdir(parents=True)
        (legacy / "versions.sqlite3").write_bytes(b"legacy-database")
    if current_directory:
        (tmp_path / "AutoFormat/tasks").mkdir(parents=True)
    assert default_store_path() == tmp_path / expected / "tasks"
    if legacy_database:
        assert (legacy / "versions.sqlite3").read_bytes() == b"legacy-database"
    if not current_directory:
        assert not (tmp_path / "AutoFormat").exists()


@pytest.mark.parametrize("headers,status", [
    ({"X-AutoFormat-Session": TOKEN}, 200),
    ({"X-EasyFormat-Session": TOKEN}, 200),
    ({"X-AutoFormat-Session": TOKEN, "X-EasyFormat-Session": TOKEN}, 200),
    ({"X-AutoFormat-Session": TOKEN, "X-EasyFormat-Session": "wrong"}, 401),
    ({"X-AutoFormat-Session": "wrong", "X-EasyFormat-Session": TOKEN}, 401),
    ({"X-EasyFormat-Session": "wrong"}, 401),
    ({}, 401),
])
def test_renamed_session_header_preserves_authentication(tmp_path, headers, status):
    store = VersionStore(tmp_path / "store")
    app = create_local_api(store=store, snapshots=RuleSnapshotStore(store), indexer=None, editor=None,
        renderer_factory=lambda _: None, font_fingerprint="test", session_token=TOKEN, allowed_origin=ORIGIN)
    with TestClient(app, base_url=ORIGIN) as client:
        assert client.get("/api/tasks", headers=headers).status_code == status
        assert client.get("/api/tasks", headers={**headers, "Origin": "https://evil.test"}).status_code == 403
