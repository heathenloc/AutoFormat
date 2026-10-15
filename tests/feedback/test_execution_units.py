"""Feedback respects table/section units and preserves unrelated objects."""
import json

import pytest
from docx import Document

from autoformat.feedback_store import FeedbackStore
from autoformat.version_store import VersionStore
from conftest import publish_fixture
from test_execution import request_for
from test_execution_integrity import executor_for
from test_processing import saved, confirm_for
from test_rules import rules_store
from test_selection import indexes, answer_for


def test_local_table_edit_does_not_change_other_table_or_text(tmp_path, indexes):
    document = Document()
    document.add_paragraph("Outside both tables")
    first = document.add_table(rows=2, cols=2)
    for cell, value in zip(first._cells, ["Selectable preview text", "Header", "A", "B"]):
        cell.text = value
    second = document.add_table(rows=2, cols=2)
    for cell, value in zip(second._cells, ["Other table", "Header", "C", "D"]):
        cell.text = value
    source = tmp_path / "paper.docx"
    document.save(source)
    versions = VersionStore(tmp_path / "store")
    v0 = versions.create_task(source)
    v1 = publish_fixture(versions, v0, tmp_path / "output")
    before = versions.artifact_path(v1.task_id, v1.version_id, "docx").read_bytes()
    feedback = FeedbackStore(versions)
    store = rules_store(feedback, indexes)
    receipt, context, _ = saved(feedback, v1, "三线表")
    analysis = store.processor.get(v1.task_id, v1.version_id, receipt.feedback_id)
    if analysis.selection.status != "ready":
        store.processor.selections.confirm(v1.task_id, v1.version_id, receipt.feedback_id,
            answer_for(receipt, context, analysis.selection))
        analysis = store.processor.get(v1.task_id, v1.version_id, receipt.feedback_id)
    scope = analysis.interpretation.scope_reviews[0]
    assert scope.minimum_unit == "table" and scope.estimated_affected_count == 1
    store.processor.confirm(v1.task_id, v1.version_id, receipt.feedback_id, confirm_for(analysis, context,
        dict(action="confirm_scope", accepted_review_ids=[scope.review_id])))
    review = store.review(v1.task_id, v1.version_id, receipt.feedback_id)
    result = executor_for(store).execute(request_for(review.snapshot, context))
    output = Document(versions.artifact_path(v1.task_id, result.version_id, "docx"))
    w = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    assert output.tables[0]._tbl.find(w + "tblPr/" + w + "tblBorders") is not None
    assert output.tables[1]._tbl.xml == second._tbl.xml
    assert [cell.text for table in output.tables for cell in table._cells] == [cell.text for table in document.tables for cell in table._cells]
    assert [p.text for p in output.paragraphs] == [p.text for p in document.paragraphs]
    assert versions.artifact_path(v1.task_id, v1.version_id, "docx").read_bytes() == before
    report = json.loads(versions.artifact_path(v1.task_id, result.version_id, "report_json").read_text(encoding="utf-8"))
    fact = report["rules"][-1]
    assert fact["property_id"] == "table.border_style" and fact["effective_scope"] == "same_kind"


def test_page_geometry_and_margin_components_publish_together(feedback_env, indexes):
    versions, feedback, _, v1, _ = feedback_env
    before = versions.artifact_path(v1.task_id, v1.version_id, "docx").read_bytes()
    store = rules_store(feedback, indexes)
    receipt, context, _ = saved(feedback, v1, "横向，页边距25mm", requested_scope="document", selection=None)
    analysis = store.processor.get(v1.task_id, v1.version_id, receipt.feedback_id)
    assert all(r.minimum_unit == "section" for r in analysis.interpretation.scope_reviews)
    store.processor.confirm(v1.task_id, v1.version_id, receipt.feedback_id, confirm_for(analysis, context,
        dict(action="confirm_scope", accepted_review_ids=[r.review_id for r in analysis.interpretation.scope_reviews])))
    review = store.review(v1.task_id, v1.version_id, receipt.feedback_id)
    result = executor_for(store).execute(request_for(review.snapshot, context))
    output = Document(versions.artifact_path(v1.task_id, result.version_id, "docx"))
    section = output.sections[0]
    assert section.page_width > section.page_height
    assert [section.top_margin.mm, section.bottom_margin.mm, section.left_margin.mm, section.right_margin.mm] == pytest.approx([25] * 4, abs=0.02)
    assert output.paragraphs[0].text == "Selectable preview text"
    assert versions.artifact_path(v1.task_id, v1.version_id, "docx").read_bytes() == before
    report = json.loads(versions.artifact_path(v1.task_id, result.version_id, "report_json").read_text(encoding="utf-8"))
    assert {c["property_id"] for c in report["changes_this_round"]} == {"page.orientation", "page.margin.top", "page.margin.bottom", "page.margin.left", "page.margin.right"}
