"""Real .NET feedback edits, publication rollback and multi-round inheritance."""
import json
import os
import shutil
from pathlib import Path

import pytest
from docx import Document
from pypdf import PdfReader

from autoformat.feedback_editor import FeedbackEditor
from autoformat.feedback_execution import FeedbackExecutor
from autoformat.feedback_execution_contracts import FeedbackExecutionRequest
from autoformat.feedback_store import FeedbackError
from autoformat.preview import PreviewArtifact, LibreOfficeRenderer, PreviewError, _sha256
from autoformat.version_store import VersionStore, new_id
from conftest import submission_for, text_pdf
from test_selection import indexes
from test_processing import saved
from test_rules import rules_store, academic_rule, decision_for


class FixtureRenderer:
    def __init__(self, root): self.root = root
    def render(self, source, artifact_id):
        relative = f"artifacts/previews/{artifact_id}.pdf"
        pdf = self.root / relative
        pdf.parent.mkdir(parents=True)
        text_pdf(pdf)
        return PreviewArtifact("preview-artifact/1.0", artifact_id, _sha256(source), _sha256(pdf),
            relative, 1, len(PdfReader(pdf).pages[0].extract_text()), "fixture", "fixture/1.0")


def editor():
    dll = Path(__file__).resolve().parents[2] / "services/docx_engine/AutoFormat.DocxEditor/bin/Debug/net10.0/AutoFormat.DocxEditor.dll"
    return FeedbackEditor(dotnet_path=Path(shutil.which("dotnet")), editor_dll=dll, expected_editor_sha256=_sha256(dll))


def request_for(snapshot, context):
    return FeedbackExecutionRequest(schema_version="feedback-execution-request/1.0", request_id=new_id(),
        task_id=context.task_id, base_version_id=context.base_version_id,
        expected_base_docx_sha256=context.base_docx_sha256, expected_context_id=context.context_id,
        expected_snapshot_id=snapshot.snapshot_id)


def character_formats(paragraph):
    return [(char, run.font.size.pt if run.font.size else None, run.bold)
            for run in paragraph.runs for char in run.text]


def prepared(feedback_env, indexes, text="字号14pt", rules=(), length=10):
    versions, feedback, _, v1, _ = feedback_env
    store = rules_store(feedback, indexes, rules)
    context = feedback.context(v1.task_id, v1.version_id)
    hint = submission_for(context).selection.model_dump(mode="json")
    hint["selected_text"] = "Selectable preview text"[:length]
    hint["spans"][0]["end_offset_utf16"] = length
    receipt, _, _ = saved(feedback, v1, text, selection=hint)
    review = store.review(v1.task_id, v1.version_id, receipt.feedback_id)
    return versions, feedback, v1, store, context, receipt, review


def test_partial_run_edit_publishes_exact_scope_and_preserves_history(feedback_env, indexes):
    versions, feedback, v1, store, context, _, review = prepared(feedback_env, indexes)
    before = {k: versions.artifact_path(v1.task_id, v1.version_id, k).read_bytes() for k in ["docx", "preview", "report_json"]}
    executor = FeedbackExecutor(store, editor=editor(), renderer_factory=FixtureRenderer, font_fingerprint="test-fonts")
    request = request_for(review.snapshot, context)
    result = executor.execute(request)
    assert result.phase == "ready"
    version = versions.get_version(v1.task_id, result.version_id)
    assert version.parent_version_id == v1.version_id
    assert version.references.feedback_rules_snapshot_id == review.snapshot.snapshot_id
    output = Document(versions.artifact_path(v1.task_id, version.version_id, "docx"))
    assert output.paragraphs[0].text == "Selectable preview text"
    assert [(r.text, r.font.size.pt if r.font.size else None) for r in output.paragraphs[0].runs] == [("Selectable", 14), (" preview text", None)]
    assert {k: versions.artifact_path(v1.task_id, v1.version_id, k).read_bytes() for k in before} == before
    report = json.loads(versions.artifact_path(v1.task_id, version.version_id, "report_json").read_text(encoding="utf-8"))
    assert report["changes_this_round"] and report["rules"][-1]["source_kind"] == "user_feedback"
    assert report["rules"][-1]["after_value"]["span"]["length_utf16"] == 10
    assert executor.execute(request) == result
    assert len(versions.list_versions(v1.task_id)) == 3
    assert not list((versions.root / ".operations").iterdir())
    with pytest.raises(FeedbackError): executor.execute(request.model_copy(update={"expected_snapshot_id": "a"*64}))


def test_local_exception_survives_reanchor_and_next_round(feedback_env, indexes):
    versions, feedback, v1, store, context, receipt, review = prepared(feedback_env, indexes, rules=[academic_rule()])
    review = store.decide(v1.task_id, v1.version_id, receipt.feedback_id, decision_for(review, context, "local_exception"))
    executor = FeedbackExecutor(store, editor=editor(), renderer_factory=FixtureRenderer, font_fingerprint="test-fonts")
    result = executor.execute(request_for(review.snapshot, context))
    v2 = versions.get_version(v1.task_id, result.version_id)
    document = Document(versions.artifact_path(v1.task_id, v2.version_id, "docx"))
    assert [(char, size) for char, size, _ in character_formats(document.paragraphs[0])] == [(char, 14 if i < 10 else 12) for i, char in enumerate("Selectable preview text")]
    second, context2, _ = saved(feedback, v2, "加粗")
    next_review = store.review(v1.task_id, v2.version_id, second.feedback_id)
    assert next_review.snapshot.status == "ready" and next_review.snapshot.parent_snapshot_id == review.snapshot.snapshot_id
    assert not any(f.status == "pending_target" for f in next_review.snapshot.fragments)
    next_result = executor.execute(request_for(next_review.snapshot, context2))
    output = Document(versions.artifact_path(v1.task_id, next_result.version_id, "docx"))
    assert character_formats(output.paragraphs[0]) == [(char, 14 if i < 10 else 12, True) for i, char in enumerate("Selectable preview text")]
    old = json.loads(versions.artifact_path(v1.task_id, v2.version_id, "report_json").read_text(encoding="utf-8"))
    new = json.loads(versions.artifact_path(v1.task_id, next_result.version_id, "report_json").read_text(encoding="utf-8"))
    assert new["changes_cumulative"][:len(old["changes_cumulative"])] == old["changes_cumulative"]
    assert new["decisions"][:len(old["decisions"])] == old["decisions"]
    restarted = rules_store(feedback, indexes, [academic_rule()])
    assert restarted.load(v1.task_id, v2.version_id) == next_review.snapshot


def test_pending_and_stale_snapshots_cannot_execute(feedback_env, indexes):
    versions, _, v1, store, context, _, review = prepared(feedback_env, indexes, rules=[academic_rule()])
    executor = FeedbackExecutor(store, editor=editor(), renderer_factory=FixtureRenderer, font_fingerprint="test-fonts")
    with pytest.raises(FeedbackError): executor.execute(request_for(review.snapshot, context))
    assert len(versions.list_versions(v1.task_id)) == 2


def test_render_failure_rolls_back_and_same_request_can_retry(feedback_env, indexes):
    versions, _, v1, store, context, _, review = prepared(feedback_env, indexes)
    class BrokenRenderer:
        def __init__(self, root): pass
        def render(self, source, artifact_id): raise PreviewError("render_failed", "test renderer unavailable")
    executor = FeedbackExecutor(store, editor=editor(), renderer_factory=BrokenRenderer, font_fingerprint="test-fonts")
    request = request_for(review.snapshot, context)
    with pytest.raises(PreviewError): executor.execute(request)
    assert executor.get(v1.task_id, request.request_id).phase == "failed"
    assert versions.working_version_id(v1.task_id) == v1.version_id and len(versions.list_versions(v1.task_id)) == 2
    assert not list((versions.root / ".operations").iterdir())
    executor.renderer_factory = FixtureRenderer
    assert executor.execute(request).phase == "ready"
    assert len(versions.list_versions(v1.task_id)) == 3


def test_already_satisfied_has_a_durable_receipt_without_new_version(feedback_env, indexes):
    versions, feedback, v1, store, context, _, review = prepared(feedback_env, indexes, "取消加粗")
    executor = FeedbackExecutor(store, editor=editor(), renderer_factory=FixtureRenderer, font_fingerprint="test-fonts")
    request = request_for(review.snapshot, context)
    result = executor.execute(request)
    assert result.phase == "no_document_change" and result.version_id is None
    restarted = FeedbackExecutor(rules_store(feedback, indexes), editor=editor(), renderer_factory=FixtureRenderer, font_fingerprint="test-fonts")
    assert restarted.execute(request) == result and len(versions.list_versions(v1.task_id)) == 2


@pytest.mark.skipif(not os.environ.get("AUTOFORMAT_TEST_SOFFICE"), reason="Pinned LibreOffice integration is opt-in")
def test_real_feedback_edit_and_libreoffice_preview(feedback_env, indexes):
    versions, _, v1, store, context, _, review = prepared(feedback_env, indexes)
    def renderer(root):
        return LibreOfficeRenderer(soffice_path=Path(os.environ["AUTOFORMAT_TEST_SOFFICE"]),
            expected_soffice_sha256=os.environ["AUTOFORMAT_TEST_SOFFICE_SHA256"], data_root=root)
    executor = FeedbackExecutor(store, editor=editor(), renderer_factory=renderer, font_fingerprint="integration-fonts")
    result = executor.execute(request_for(review.snapshot, context))
    context2 = store.feedback.context(v1.task_id, result.version_id)
    assert context2.renderer_id == "libreoffice" and context2.base_docx_sha256 != context.base_docx_sha256
    assert "Selectable preview text" in PdfReader(versions.artifact_path(v1.task_id, result.version_id, "preview")).pages[0].extract_text()
