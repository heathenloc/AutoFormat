"""Real immutable version artifacts; no .NET, rendering or model service needed."""

from __future__ import annotations

import hashlib
import json
import shutil
from pathlib import Path

import pytest
from docx import Document
from pypdf import PdfReader

from autoformat.feedback_contracts import FeedbackSubmission
from autoformat.feedback_store import FeedbackStore
from autoformat.preview import PreviewArtifact
from autoformat.reports import ReportContext, build_report, write_report
from autoformat.version_store import VersionReferences, FeedbackVersionReferences, VersionStore, new_id


def hash_file(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def text_pdf(path):
    content = b"BT /F1 14 Tf 72 720 Td (Selectable preview text) Tj ET"
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream",
    ]
    output = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for number, body in enumerate(objects, 1):
        offsets.append(len(output))
        output.extend(f"{number} 0 obj\n".encode() + body + b"\nendobj\n")
    cross_reference = len(output)
    output.extend(f"xref\n0 {len(offsets)}\n0000000000 65535 f \n".encode())
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode())
    output.extend(f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{cross_reference}\n%%EOF\n".encode())
    path.write_bytes(output)


def publish_fixture(store, parent, directory, *, rules_id="rules-v1", feedback_snapshot_id=None):
    directory.mkdir()
    candidate = directory / "candidate.docx"
    source = store.artifact_path(parent.task_id, parent.version_id, "docx")
    changes = []
    rule_facts = []
    if parent.kind == "output":
        document = Document(source)
        before = document.paragraphs[0].runs[0].bold
        after = not before
        document.paragraphs[0].runs[0].bold = after
        document.save(candidate)
        assert Document(candidate).paragraphs[0].runs[0].bold == after
        rule_id = f"fixture-bold-{new_id()}"
        rule_facts = [{"rule_id": rule_id, "source_kind": "brief_explicit",
            "source_artifact_id": "fixture-brief", "source_version": "brief:v1",
            "original_requirement": "测试样本首段加粗", "evidence_refs": [{
                "source_artifact_id": "fixture-brief", "source_version": "brief:v1",
                "locator": "fixture:1", "quote": "测试样本首段加粗"}],
            "object_id": "p:000001", "semantic_role": "body", "property_id": "font.bold",
            "requested_scope": "paragraph", "effective_scope": "paragraph",
            "before_value": before, "after_value": after, "status": "modified_verified",
            "reason": "Fixture format checked after reopening DOCX.", "next_action": None,
            "decision_id": None, "verification_ref": rule_id}]
        changes = [{"change_id": "change-" + rule_id, "rule_id": rule_id, "object_id": "p:000001",
            "property_id": "font.bold", "before_value": before, "after_value": after,
            "effective_scope": "paragraph", "validation_status": "verified", "verification_ref": rule_id}]
    else:
        shutil.copyfile(source, candidate)
    pdf = directory / "preview.pdf"
    text_pdf(pdf)
    version_id, artifact_id = new_id(), new_id()
    text_chars = len(PdfReader(pdf).pages[0].extract_text())
    artifact = PreviewArtifact(
        schema_version="preview-artifact/1.0", artifact_id=artifact_id,
        source_sha256=hash_file(candidate), pdf_sha256=hash_file(pdf),
        pdf_relative_path=f"artifacts/previews/{artifact_id}.pdf", page_count=1,
        selectable_text_chars=text_chars, renderer_id="fixture",
        renderer_version="LibreOffice 26.8.0.3 fixture",
    )
    previous = (json.loads(store.artifact_path(parent.task_id, parent.version_id, "report_json").read_text(encoding="utf-8"))
                if parent.kind == "output" else None)
    report = build_report(
        ReportContext(task_id=parent.task_id, version_id=version_id, parent_version_id=parent.version_id,
            source_docx_sha256=parent.docx_sha256, output_docx_sha256=hash_file(candidate),
            brief_version_id="brief:v1", rules_snapshot_id=rules_id, edit_plan_id="fixture-plan",
            preview_artifact_id=artifact_id, renderer_id=artifact.renderer_id,
            renderer_version=artifact.renderer_version, font_fingerprint="fixture-fonts"),
        rules=(previous["rules"] if previous else []) + rule_facts, changes_this_round=changes,
        changes_cumulative=(previous["changes_cumulative"] if previous else []) + changes,
        decisions=[], issues=[], note="Feedback storage fixture; no format compliance claim.",
    )
    report_path, markdown_path = directory / "report.json", directory / "report.md"
    write_report(report_path, markdown_path, report)
    return store.publish_version(task_id=parent.task_id, version_id=version_id,
        parent_version_id=parent.version_id, expected_parent_sha256=parent.docx_sha256,
        request_id=new_id(), candidate_docx=candidate, report_json=report_path,
        report_markdown=markdown_path, preview_pdf=pdf, preview_artifact=artifact,
        references=(FeedbackVersionReferences(rules_id, "brief:v1", "fixture-plan", "fixture-fonts", feedback_rules_snapshot_id=feedback_snapshot_id)
            if feedback_snapshot_id else VersionReferences(rules_id, "brief:v1", "fixture-plan", "fixture-fonts")),
        expected_feedback_snapshot_id=feedback_snapshot_id)


def submission_for(context, **updates):
    payload = {
        "schema_version": "feedback-submission/1.0", "request_id": new_id(),
        "task_id": context.task_id, "base_version_id": context.base_version_id,
        "expected_base_docx_sha256": context.base_docx_sha256,
        "expected_context_id": context.context_id,
        "preview_artifact_id": context.preview_artifact_id,
        "expected_preview_pdf_sha256": context.preview_pdf_sha256,
        "raw_text": "  这一段改为居中\n", "requested_scope": "selection",
        "selection": {"selected_text": "Selectable preview text", "context_before": "", "context_after": "",
            "spans": [{"page": 1, "start_item": 0, "start_offset_utf16": 0,
                       "end_item": 0, "end_offset_utf16": 23}]},
    }
    payload.update(updates)
    return FeedbackSubmission.model_validate(payload)


@pytest.fixture
def feedback_env(tmp_path):
    source = tmp_path / "original.docx"
    document = Document()
    document.add_paragraph("Selectable preview text")
    document.save(source)
    store = VersionStore(tmp_path / "store")
    v0 = store.create_task(source)
    v1 = publish_fixture(store, v0, tmp_path / "v1")
    feedback = FeedbackStore(store)
    return store, feedback, v0, v1, source
