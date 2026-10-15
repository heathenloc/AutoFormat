"""Real PDF.js + academic .NET IDs + immutable feedback/target decision storage."""
from __future__ import annotations

import hashlib
import json
import shutil
import subprocess
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient

from autoformat.feedback_contracts import SelectionHint, SelectionTargetConfirmation
from autoformat.feedback_store import FeedbackError, FeedbackStore
from autoformat.local_api import create_local_api
from autoformat.rule_snapshots import RuleSnapshotStore
from autoformat.selection_indexes import SelectionIndexes
from autoformat.selection_resolver import find_matches, verify_hint
from autoformat.selection_store import SelectionStore
from autoformat.version_store import VersionStore, new_id

from conftest import publish_fixture, submission_for

ROOT = Path(__file__).resolve().parents[2]
TOKEN = "selection-test-session-token-1234567890"
ORIGIN = "http://127.0.0.1:57843"
HEADERS = {"X-AutoFormat-Session": TOKEN, "Origin": ORIGIN}


@pytest.fixture(scope="module")
def indexes():
    subprocess.run([shutil.which("dotnet"), "build", str(ROOT / "services/docx_engine/AutoFormat.DocxEditor"), "--no-restore"],
        cwd=ROOT, capture_output=True, check=True)
    return SelectionIndexes.from_project(ROOT)


def custom_task(tmp_path, paragraphs):
    source = tmp_path / "paper.docx"
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    doc.save(source)
    versions = VersionStore(tmp_path / "store")
    v0 = versions.create_task(source)
    v1 = publish_fixture(versions, v0, tmp_path / "output")
    return versions, FeedbackStore(versions), v1


def submitted(feedback, version, **updates):
    context = feedback.context(version.task_id, version.version_id)
    request = submission_for(context, **updates)
    receipt = feedback.submit(version.task_id, version.version_id, request)
    return receipt, context


def answer_for(receipt, context, result, **updates):
    payload = dict(schema_version="selection-target-confirmation/1.0", request_id=new_id(), task_id=context.task_id,
        base_version_id=context.base_version_id, feedback_id=receipt.feedback_id,
        expected_base_docx_sha256=context.base_docx_sha256, expected_context_id=context.context_id,
        expected_resolution_id=result.resolution_id, candidate_id=result.candidates[-1].candidate_id,
        raw_answer="确认第二个位置")
    payload.update(updates)
    return SelectionTargetConfirmation(**payload)


def test_real_indexes_produce_editor_ids_and_preserve_all_artifacts(feedback_env, indexes):
    versions, feedback, _, v1, _ = feedback_env
    receipt, context = submitted(feedback, v1)
    before = {kind: versions.artifact_path(v1.task_id, v1.version_id, kind).read_bytes()
              for kind in ["docx", "preview", "report_json", "report_markdown"]}
    selections = SelectionStore(feedback, indexes)
    first = selections.get(v1.task_id, v1.version_id, receipt.feedback_id)
    assert first.status == "ready" and first.anchor is not None
    span = first.anchor.ranges[0]
    assert span.matched_text == "Selectable preview text" and span.length_utf16 == 23
    assert span.object_id.startswith("o:")
    academic = json.loads(subprocess.run([str(indexes.dotnet), str(indexes.editor_dll), "inspect-academic",
        str(versions.artifact_path(v1.task_id, v1.version_id, "docx"))], capture_output=True, encoding="utf-8", check=True).stdout)
    assert span.object_id in {item["object_id"] for item in academic["objects"]}
    assert selections.get(v1.task_id, v1.version_id, receipt.feedback_id) == first
    restarted = SelectionStore(FeedbackStore(VersionStore(versions.root)), SelectionIndexes.from_project(ROOT))
    assert restarted.get(v1.task_id, v1.version_id, receipt.feedback_id) == first
    assert context.base_docx_sha256 == first.anchor.base_docx_sha256
    assert len(versions.list_versions(v1.task_id)) == 2
    assert {kind: versions.artifact_path(v1.task_id, v1.version_id, kind).read_bytes() for kind in before} == before


def test_duplicate_text_requires_one_bound_target_and_replays_after_working_version_changes(tmp_path, indexes):
    versions, feedback, v1 = custom_task(tmp_path, ["Selectable preview text", "Selectable preview text"])
    receipt, context = submitted(feedback, v1)
    selections = SelectionStore(feedback, indexes)
    result = selections.get(v1.task_id, v1.version_id, receipt.feedback_id)
    assert result.status == "ambiguous" and len(result.candidates) == 2 and result.anchor is None
    answer = answer_for(receipt, context, result)
    ready = selections.confirm(v1.task_id, v1.version_id, receipt.feedback_id, answer)
    assert ready.status == "ready" and ready.anchor.ranges[0].object_id == result.candidates[1].object_id
    assert "用户选择" in ready.reason
    assert selections.get(v1.task_id, v1.version_id, receipt.feedback_id) == ready
    publish_fixture(versions, v1, tmp_path / "v2")
    assert selections.confirm(v1.task_id, v1.version_id, receipt.feedback_id, answer) == ready
    changed = answer.model_copy(update={"candidate_id": result.candidates[0].candidate_id})
    with pytest.raises(FeedbackError, match="请求 ID"):
        selections.confirm(v1.task_id, v1.version_id, receipt.feedback_id, changed)
    assert len(versions.list_versions(v1.task_id)) == 3


def test_trusted_pdf_context_disambiguates_repeated_substring(tmp_path, indexes):
    versions, feedback, v1 = custom_task(tmp_path, ["Selectable preview text", "Another preview sample"])
    hint = {"selected_text": "preview", "context_before": "Selectable ", "context_after": " text",
        "spans": [{"page": 1, "start_item": 0, "start_offset_utf16": 11, "end_item": 0, "end_offset_utf16": 18}]}
    receipt, _ = submitted(feedback, v1, selection=hint)
    result = SelectionStore(feedback, indexes).get(v1.task_id, v1.version_id, receipt.feedback_id)
    assert result.status == "ready" and result.anchor.ranges[0].start_utf16 == 11
    assert result.anchor.ranges[0].matched_text == "preview"


@pytest.mark.parametrize("change", [
    {"selected_text": "Forged preview text"},
    {"context_before": "not in the PDF"},
    {"spans": [{"page": 1, "start_item": 900, "start_offset_utf16": 0, "end_item": 900, "end_offset_utf16": 2}]},
    {"spans": [{"page": 1, "start_item": 0, "start_offset_utf16": 0, "end_item": 0, "end_offset_utf16": 1000}]},
])
def test_client_cannot_invent_pdf_text_or_item_offsets(feedback_env, indexes, change):
    versions, feedback, _, v1, _ = feedback_env
    request = submission_for(feedback.context(v1.task_id, v1.version_id)).model_dump(mode="json")
    request["selection"].update(change)
    receipt, _ = submitted(feedback, v1, selection=request["selection"])
    with pytest.raises(FeedbackError):
        SelectionStore(feedback, indexes).get(v1.task_id, v1.version_id, receipt.feedback_id)
    assert len(versions.list_versions(v1.task_id)) == 2


def test_unknown_and_stale_target_decisions_rejected_without_publication(tmp_path, indexes):
    versions, feedback, v1 = custom_task(tmp_path, ["Selectable preview text"] * 2)
    receipt, context = submitted(feedback, v1)
    selections = SelectionStore(feedback, indexes)
    result = selections.get(v1.task_id, v1.version_id, receipt.feedback_id)
    for changes in [{"candidate_id": "f" * 64}, {"expected_resolution_id": "e" * 64}, {"expected_context_id": "d" * 64}]:
        with pytest.raises(FeedbackError):
            selections.confirm(v1.task_id, v1.version_id, receipt.feedback_id, answer_for(receipt, context, result, **changes))
    publish_fixture(versions, v1, tmp_path / "v2")
    with pytest.raises(FeedbackError) as error:
        selections.confirm(v1.task_id, v1.version_id, receipt.feedback_id, answer_for(receipt, context, result))
    assert error.value.status == "stale_base_version"


def test_concurrent_confirmation_records_one_decision(tmp_path, indexes):
    versions, feedback, v1 = custom_task(tmp_path, ["Selectable preview text"] * 2)
    receipt, context = submitted(feedback, v1)
    selections = SelectionStore(feedback, indexes)
    result = selections.get(v1.task_id, v1.version_id, receipt.feedback_id)
    answer = answer_for(receipt, context, result)
    with ThreadPoolExecutor(max_workers=2) as pool:
        results = list(pool.map(lambda _: selections.confirm(v1.task_id, v1.version_id, receipt.feedback_id, answer), range(2)))
    assert results[0] == results[1]
    with versions._transaction() as connection:
        assert connection.execute("SELECT count(*) FROM selection_target_decisions").fetchone()[0] == 1


def test_persisted_anchor_is_compared_with_source_replay(feedback_env, indexes):
    versions, feedback, _, v1, _ = feedback_env
    receipt, _ = submitted(feedback, v1)
    selections = SelectionStore(feedback, indexes)
    result = selections.get(v1.task_id, v1.version_id, receipt.feedback_id)
    forged = result.model_dump(mode="json")
    forged["anchor"]["ranges"][0]["start_utf16"] = 7
    with versions._transaction() as connection:
        connection.execute("UPDATE selection_resolutions SET payload_json=?, payload_sha256=? WHERE feedback_id=?",
            (json.dumps(forged), hashlib.sha256(b"forged").hexdigest(), receipt.feedback_id))
    with pytest.raises(FeedbackError, match="重放不一致"):
        selections.get(v1.task_id, v1.version_id, receipt.feedback_id)


def test_no_selection_does_not_start_any_indexer(feedback_env):
    versions, feedback, _, v1, _ = feedback_env
    receipt, _ = submitted(feedback, v1, selection=None, requested_scope="document")
    class NoIndexes:
        fingerprint = "0" * 64
        def __getattr__(self, name):
            raise AssertionError("No selection must not start an index process")
    assert SelectionStore(feedback, NoIndexes()).get(v1.task_id, v1.version_id, receipt.feedback_id).status == "needs_user_input"


def test_http_routes_authorize_and_never_accept_client_anchors(feedback_env, indexes):
    versions, feedback, _, v1, _ = feedback_env
    receipt, _ = submitted(feedback, v1)
    app = create_local_api(store=versions, snapshots=RuleSnapshotStore(versions), indexer=None, editor=None,
        renderer_factory=lambda _: None, session_token=TOKEN, allowed_origin=ORIGIN, font_fingerprint="test-fonts", selection_indexes=indexes)
    prefix = f"/api/tasks/{v1.task_id}/versions/{v1.version_id}/feedback/{receipt.feedback_id}/selection"
    with TestClient(app, base_url=ORIGIN) as client:
        assert client.get(prefix).status_code == 401
        assert client.get(prefix, headers={**HEADERS, "Origin": "https://evil.test"}).status_code == 403
        result = client.get(prefix, headers=HEADERS)
        assert result.status_code == 200 and result.json()["status"] == "ready"
        assert "word/document.xml" not in result.text and "xml_path" not in result.text
        assert client.post(prefix + "/confirm", headers=HEADERS, json={"anchor": {"status": "ready"}}).status_code == 422
        assert client.get(prefix.replace(receipt.feedback_id, new_id()), headers=HEADERS).status_code == 422


def test_component_fingerprint_drift_is_rejected(feedback_env, indexes, monkeypatch):
    versions, feedback, _, v1, _ = feedback_env
    receipt, _ = submitted(feedback, v1)
    monkeypatch.setattr(indexes, "digests", ["0" * 64] * len(indexes.digests))
    with pytest.raises(FeedbackError, match="定位组件已变化"):
        SelectionStore(feedback, indexes).get(v1.task_id, v1.version_id, receipt.feedback_id)


def test_wrapped_text_preserves_unicode_and_rejects_semantic_normalization(feedback_env):
    from autoformat.selection_indexes import DocxParagraph, DocxSelectionIndex
    from autoformat.selection_resolver import PdfEvidence
    _, feedback, _, v1, _ = feedback_env
    context = feedback.context(v1.task_id, v1.version_id)
    def match(pdf, source):
        index = DocxSelectionIndex(schema_version="selection-docx-index/1.0", sha256=context.base_docx_sha256,
            paragraphs=[DocxParagraph(object_id="o:" + "a" * 24, story_id="body", text=source,
                location_label="正文第1段", editable=True, semantic_role="body", ordinal=1)])
        return find_matches(index, PdfEvidence(pdf, "", ""), context)
    span = match("😀中\n文", "前😀中文后")[0].ranges[0]
    assert (span.start_utf16, span.length_utf16, span.matched_text) == (1, 4, "😀中文")
    for pdf, source in [("ab", "a b"), ("Ab", "ab"), ("ａ", "a"), ("a,b", "a.b"), ("ﬃ", "ffi")]:
        assert match(pdf, source) == []
    with pytest.raises(FeedbackError, match="范围过长"):
        match("a\nb", "a" + " " * 5000 + "b")


def test_utf16_surrogate_splitting_is_rejected(feedback_env):
    from autoformat.selection_indexes import PdfItem, PdfPage, PdfSelectionIndex
    _, feedback, _, v1, _ = feedback_env
    context = feedback.context(v1.task_id, v1.version_id)
    index = PdfSelectionIndex(schema_version="selection-pdf-index/1.0", pdfjs_version="6.4.299",
        pdf_sha256=context.preview_pdf_sha256, page_count=1,
        pages=[PdfPage(page=1, items=[PdfItem(index=0, str="😀x", has_eol=False)])])
    hint = SelectionHint(selected_text="x", spans=[{"page": 1, "start_item": 0, "start_offset_utf16": 1, "end_item": 0, "end_offset_utf16": 3}])
    with pytest.raises(FeedbackError, match="Unicode"):
        verify_hint(hint, context, index)


def test_contiguous_cross_page_selection_maps_multiple_docx_paragraphs(feedback_env):
    from autoformat.selection_indexes import DocxParagraph, DocxSelectionIndex, PdfItem, PdfPage, PdfSelectionIndex
    _, feedback, _, v1, _ = feedback_env
    context = feedback.context(v1.task_id, v1.version_id).model_copy(update={"page_count": 2})
    index = PdfSelectionIndex(schema_version="selection-pdf-index/1.0", pdfjs_version="6.4.299",
        pdf_sha256=context.preview_pdf_sha256, page_count=2,
        pages=[PdfPage(page=n, items=[PdfItem(index=0, str=text, has_eol=False)])
               for n, text in [(1, "前文第一段"), (2, "第二段后文")]])
    hint = SelectionHint(selected_text="第一段\n第二段", context_before="前文", context_after="后文",
        spans=[dict(page=1, start_item=0, end_item=0, start_offset_utf16=2, end_offset_utf16=5),
               dict(page=2, start_item=0, end_item=0, start_offset_utf16=0, end_offset_utf16=3)])
    evidence = verify_hint(hint, context, index)
    docx = DocxSelectionIndex(schema_version="selection-docx-index/1.0", sha256=context.base_docx_sha256,
        paragraphs=[DocxParagraph(object_id="o:" + str(n) * 24, story_id="body", text=text,
            location_label=f"正文第{n}段", editable=True, semantic_role="body", ordinal=n)
            for n, text in [(1, "前文第一段"), (2, ""), (3, "第二段后文")]])
    result = find_matches(docx, evidence, context)
    assert len(result) == 1 and result[0].supported
    assert [(r.start_utf16, r.length_utf16, r.matched_text) for r in result[0].ranges] == [(2, 3, "第一段"), (0, 3, "第二段")]
    disjoint = hint.model_copy(update={"spans": [hint.spans[0].model_copy(update={"end_offset_utf16": 4}), hint.spans[1]]})
    with pytest.raises(FeedbackError, match="连续选区"):
        verify_hint(disjoint, context, index)
    with pytest.raises(FeedbackError, match="不属于此预览"):
        verify_hint(hint, context, index.model_copy(update={"pdf_sha256": "0" * 64}))


def test_real_docx_index_blocks_unsafe_fields_hidden_revisions_and_inline_objects(tmp_path, indexes):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.enum.style import WD_STYLE_TYPE
    document = Document()
    document.add_paragraph("普通正文")
    hidden = document.add_paragraph("隐藏文字"); hidden.runs[0].font.hidden = True
    style = document.styles.add_style("HiddenSelectionStyle", WD_STYLE_TYPE.CHARACTER); style.font.hidden = True
    inherited = document.add_paragraph("样式隐藏"); inherited.runs[0].style = style
    field = document.add_paragraph("字段显示值")
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin"); field.runs[0]._r.insert(0, begin)
    document.add_paragraph("跨段字段显示值")
    end_p = document.add_paragraph("字段末尾")
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end"); end_p.runs[0]._r.append(end)
    document.add_paragraph("字段之后")
    revised = document.add_paragraph(); insertion = OxmlElement("w:ins")
    run = OxmlElement("w:r"); text = OxmlElement("w:t"); text.text = "修订文字"; run.append(text); insertion.append(run); revised._p.append(insertion)
    document.add_paragraph("含制表符\t文字")
    document.add_paragraph("含换行\n文字")
    hyphen = document.add_paragraph("含连字符文字"); hyphen.runs[0]._r.append(OxmlElement("w:noBreakHyphen"))
    table = document.add_table(rows=1, cols=1); table.cell(0, 0).text = "普通表格文字"
    hidden_style = document.styles.add_style("HiddenTableStyle", WD_STYLE_TYPE.TABLE)
    conditional = OxmlElement("w:tblStylePr"); conditional.set(qn("w:type"), "firstRow")
    properties = OxmlElement("w:rPr"); properties.append(OxmlElement("w:vanish")); conditional.append(properties); hidden_style.element.append(conditional)
    hidden_table = document.add_table(rows=1, cols=1); hidden_table.style = hidden_style; hidden_table.cell(0, 0).text = "表格条件样式隐藏"
    document.sections[0].header.paragraphs[0].text = "普通页眉文字"
    source = tmp_path / "unsafe.docx"; document.save(source)
    before = source.read_bytes()
    indexed = indexes.docx(source)
    by_text = {p.text: p for p in indexed.paragraphs}
    for value in ["隐藏文字", "样式隐藏", "字段显示值", "跨段字段显示值", "字段末尾", "修订文字", "含制表符文字", "含换行文字", "含连字符文字", "表格条件样式隐藏"]:
        assert not by_text[value].editable, value
    for value in ["普通正文", "字段之后", "普通表格文字", "普通页眉文字"]:
        assert by_text[value].editable, value
    assert by_text["普通正文"].story_id != by_text["普通页眉文字"].story_id
    assert source.read_bytes() == before


def test_unsafe_duplicate_is_never_silently_ignored(tmp_path, indexes):
    source = tmp_path / "hidden.docx"
    document = Document()
    document.add_paragraph("Selectable preview text")
    hidden = document.add_paragraph("Selectable preview text"); hidden.runs[0].font.hidden = True
    document.save(source)
    versions = VersionStore(tmp_path / "store"); v0 = versions.create_task(source)
    v1 = publish_fixture(versions, v0, tmp_path / "output")
    feedback = FeedbackStore(versions); receipt, _ = submitted(feedback, v1)
    result = SelectionStore(feedback, indexes).get(v1.task_id, v1.version_id, receipt.feedback_id)
    assert result.status == "unsupported_target" and result.anchor is None


def test_http_target_confirmation_is_bound_and_idempotent(tmp_path, indexes):
    versions, feedback, v1 = custom_task(tmp_path, ["Selectable preview text"] * 2)
    receipt, context = submitted(feedback, v1)
    app = create_local_api(store=versions, snapshots=RuleSnapshotStore(versions), indexer=None, editor=None,
        renderer_factory=lambda _: None, session_token=TOKEN, allowed_origin=ORIGIN, font_fingerprint="test-fonts", selection_indexes=indexes)
    prefix = f"/api/tasks/{v1.task_id}/versions/{v1.version_id}/feedback/{receipt.feedback_id}/selection"
    with TestClient(app, base_url=ORIGIN) as client:
        from autoformat.feedback_contracts import SelectionResolution
        result = SelectionResolution.model_validate(client.get(prefix, headers=HEADERS).json())
        answer = answer_for(receipt, context, result).model_dump(mode="json")
        first = client.post(prefix + "/confirm", headers=HEADERS, json=answer)
        assert first.status_code == 200 and first.json()["status"] == "ready"
        assert client.post(prefix + "/confirm", headers=HEADERS, json=answer).json() == first.json()
        assert client.get(prefix, headers=HEADERS).json() == first.json()
        assert len(versions.list_versions(v1.task_id)) == 2


def test_multiple_paragraph_target_requires_explicit_review(tmp_path, indexes):
    from autoformat.selection_resolver import PdfEvidence, resolution
    versions, feedback, v1 = custom_task(tmp_path, ["Selectable", "preview text"])
    selection = SelectionHint(selected_text="Selectable\npreview text", spans=[
        dict(page=1, start_item=0, start_offset_utf16=0, end_item=1, end_offset_utf16=12)])
    receipt, context = submitted(feedback, v1, selection=selection)
    record = feedback.get(v1.task_id, v1.version_id, receipt.feedback_id)
    matches = find_matches(indexes.docx(versions.artifact_path(v1.task_id, v1.version_id, "docx")),
        PdfEvidence(selection.selected_text, "", ""), context)
    result = resolution(record, context, matches)
    assert result.status == "ambiguous" and len(result.candidates) == 1 and result.anchor is None
    confirmed = resolution(record, context, matches, chosen=result.candidates[0].candidate_id)
    assert confirmed.status == "ready" and len(confirmed.anchor.ranges) == 2


def test_header_target_requires_review_even_when_unique(tmp_path, indexes):
    source = tmp_path / "header.docx"
    document = Document(); document.add_paragraph("Other body text")
    document.sections[0].header.paragraphs[0].text = "Selectable preview text"
    document.save(source)
    versions = VersionStore(tmp_path / "store"); v0 = versions.create_task(source)
    v1 = publish_fixture(versions, v0, tmp_path / "output")
    feedback = FeedbackStore(versions); receipt, context = submitted(feedback, v1)
    selections = SelectionStore(feedback, indexes)
    result = selections.get(v1.task_id, v1.version_id, receipt.feedback_id)
    assert result.status == "ambiguous" and len(result.candidates) == 1
    assert "页眉" in result.candidates[0].location_label
    assert selections.confirm(v1.task_id, v1.version_id, receipt.feedback_id,
        answer_for(receipt, context, result)).status == "ready"


def test_same_paragraph_repeated_targets_have_distinct_offsets_and_marked_excerpts(tmp_path, indexes):
    versions, feedback, v1 = custom_task(tmp_path, ["Selectable preview text / Selectable preview text"])
    receipt, context = submitted(feedback, v1)
    selections = SelectionStore(feedback, indexes)
    result = selections.get(v1.task_id, v1.version_id, receipt.feedback_id)
    assert result.status == "ambiguous" and len(result.candidates) == 2
    assert result.candidates[0].object_id == result.candidates[1].object_id
    assert result.candidates[0].excerpt != result.candidates[1].excerpt
    assert all("「Selectable preview text」" in c.excerpt for c in result.candidates)
    chosen = selections.confirm(v1.task_id, v1.version_id, receipt.feedback_id, answer_for(receipt, context, result))
    assert chosen.anchor.ranges[0].start_utf16 == 26
