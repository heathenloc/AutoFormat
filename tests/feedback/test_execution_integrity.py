"""Concurrency, source races, API boundaries and Unicode editing checks."""
import json
from concurrent.futures import ThreadPoolExecutor
from threading import Barrier

import pytest
from docx import Document
from fastapi.testclient import TestClient

from autoformat.feedback_contracts import content_hash
from autoformat.feedback_execution import FeedbackExecutor
from autoformat.feedback_store import FeedbackError
from autoformat.local_api import create_local_api
from autoformat.rule_snapshots import RuleSnapshotStore
from autoformat.version_store import VersionStoreError
from test_execution import prepared, editor, request_for, FixtureRenderer
from conftest import submission_for
from test_processing import saved, confirm_for
from test_rules import rules_store, basis_for
from test_selection import indexes, custom_task


def executor_for(store, renderer=FixtureRenderer):
    return FeedbackExecutor(store, editor=editor(), renderer_factory=renderer, font_fingerprint="test-fonts")


def test_rules_changed_during_render_never_publish_a_partial_version(feedback_env, indexes):
    versions, feedback, v1, store, context, _, review = prepared(feedback_env, indexes)
    class RacingRenderer(FixtureRenderer):
        def render(self, source, artifact_id):
            artifact = super().render(source, artifact_id)
            receipt, _, _ = saved(feedback, v1, "加粗")
            store.review(v1.task_id, v1.version_id, receipt.feedback_id)
            return artifact
    executor = executor_for(store, RacingRenderer)
    request = request_for(review.snapshot, context)
    with pytest.raises(FeedbackError): executor.execute(request)
    assert executor.get(v1.task_id, request.request_id).phase == "failed"
    assert len(versions.list_versions(v1.task_id)) == 2
    assert versions.working_version_id(v1.task_id) == v1.version_id
    assert not list((versions.root / ".operations").iterdir())


def test_identical_requests_across_instances_publish_only_one_version(feedback_env, indexes):
    versions, _, v1, store, context, _, review = prepared(feedback_env, indexes)
    executors = [executor_for(store), executor_for(store)]
    request = request_for(review.snapshot, context)
    barrier = Barrier(2)
    for executor in executors:
        original = executor._fresh
        calls = [0]
        def fresh(value, fn=original, counter=calls):
            result = fn(value)
            counter[0] += 1
            if counter[0] == 1: barrier.wait(timeout=45)
            return result
        executor._fresh = fresh
    with ThreadPoolExecutor(2) as pool:
        results = list(pool.map(lambda e: e.execute(request), executors))
    assert all(r.request_id == request.request_id for r in results)
    assert executors[0].get(v1.task_id, request.request_id).phase == "ready"
    assert len(versions.list_versions(v1.task_id)) == 3


def test_committed_version_recovers_even_if_receipt_update_was_interrupted(feedback_env, indexes):
    versions, _, v1, store, context, _, review = prepared(feedback_env, indexes)
    executor = executor_for(store)
    original = executor._update
    def interrupted(request, attempt, phase, **kwargs):
        if phase == "ready": raise RuntimeError("Simulated crash after atomic publication")
        return original(request, attempt, phase, **kwargs)
    executor._update = interrupted
    request = request_for(review.snapshot, context)
    with pytest.raises(RuntimeError): executor.execute(request)
    recovered = executor.get(v1.task_id, request.request_id)
    assert recovered.phase == "ready" and recovered.version_id is not None
    assert executor.execute(request) == recovered
    assert len(versions.list_versions(v1.task_id)) == 3


def test_scope_choice_race_is_rechecked_in_publication_transaction(feedback_env, indexes):
    versions, feedback, _, v1, _ = feedback_env
    store = rules_store(feedback, indexes)
    receipt, context, _ = saved(feedback, v1, "这一段居中")
    analysis = store.processor.get(v1.task_id, v1.version_id, receipt.feedback_id)
    answer = dict(action="confirm_scope", accepted_review_ids=[r.review_id for r in analysis.interpretation.scope_reviews])
    confirmed = store.processor.confirm(v1.task_id, v1.version_id, receipt.feedback_id, confirm_for(analysis, context, answer))
    review = store.review(v1.task_id, v1.version_id, receipt.feedback_id)
    original = versions.publish_version
    def raced(**kwargs):
        store.processor.confirm(v1.task_id, v1.version_id, receipt.feedback_id, confirm_for(confirmed, context, answer))
        return original(**kwargs)
    versions.publish_version = raced
    executor = executor_for(store)
    with pytest.raises((FeedbackError, VersionStoreError)): executor.execute(request_for(review.snapshot, context))
    assert len(versions.list_versions(v1.task_id)) == 2
    assert versions.working_version_id(v1.task_id) == v1.version_id


def test_unicode_character_filter_does_not_touch_latin_digits_or_emoji(tmp_path, indexes):
    versions, feedback, v1 = custom_task(tmp_path, ["Selectable preview text", "A😀中1B𠀀"])
    store = rules_store(feedback, indexes)
    receipt, context, _ = saved(feedback, v1, "中文黑体", requested_scope="document", selection=None)
    analysis = store.processor.get(v1.task_id, v1.version_id, receipt.feedback_id)
    store.processor.confirm(v1.task_id, v1.version_id, receipt.feedback_id, confirm_for(analysis, context,
        dict(action="confirm_scope", accepted_review_ids=[r.review_id for r in analysis.interpretation.scope_reviews])))
    review = store.review(v1.task_id, v1.version_id, receipt.feedback_id)
    result = executor_for(store).execute(request_for(review.snapshot, context))
    document = Document(versions.artifact_path(v1.task_id, result.version_id, "docx"))
    assert document.paragraphs[1].text == "A😀中1B𠀀"
    w = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    for run in document.paragraphs[1].runs:
        fonts = run._r.find(w + "rPr/" + w + "rFonts")
        if any(c in "中𠀀" for c in run.text): assert fonts is not None and fonts.get(w + "eastAsia") == "黑体"
        else: assert fonts is None or fonts.get(w + "eastAsia") != "黑体"


def test_already_satisfied_middle_range_keeps_original_bytes(feedback_env, indexes):
    versions, feedback, _, v1, _ = feedback_env
    store = rules_store(feedback, indexes)
    context = feedback.context(v1.task_id, v1.version_id)
    hint = submission_for(context).selection.model_dump(mode="json")
    hint["selected_text"] = "preview"
    hint["spans"][0].update(start_offset_utf16=11, end_offset_utf16=18)
    receipt, _, _ = saved(feedback, v1, "取消加粗", selection=hint)
    review = store.review(v1.task_id, v1.version_id, receipt.feedback_id)
    result = executor_for(store).execute(request_for(review.snapshot, context))
    assert result.phase == "no_document_change" and len(versions.list_versions(v1.task_id)) == 2


def test_execution_api_is_session_protected_and_cannot_accept_raw_rules(feedback_env, indexes):
    versions, _, v1, store, context, _, review = prepared(feedback_env, indexes)
    token, origin = "execution-test-session-token-1234567890", "http://127.0.0.1:57843"
    client = TestClient(create_local_api(store=versions, snapshots=RuleSnapshotStore(versions),
        indexer=None, editor=editor(), renderer_factory=FixtureRenderer, session_token=token,
        allowed_origin=origin, font_fingerprint="test-fonts", selection_indexes=indexes,
        feedback_basis_loader=lambda c: basis_for(c)), base_url=origin)
    headers = {"X-AutoFormat-Session": token, "Origin": origin}
    request = request_for(review.snapshot, context)
    path = f"/api/tasks/{v1.task_id}/versions/{v1.version_id}/feedback-executions"
    body = request.model_dump(mode="json")
    assert client.post(path, json=body).status_code == 401
    assert client.post(path, json={**body, "raw_text": "任意修改", "rules": []}, headers=headers).status_code == 422
    response = client.post(path, json=body, headers=headers)
    assert response.status_code == 200 and response.json()["phase"] == "ready"
    assert client.post(path, json=body, headers=headers).json() == response.json()
    assert client.get(f"/api/tasks/{v1.task_id}/feedback-executions/{request.request_id}", headers=headers).json() == response.json()


def test_source_replayed_history_survives_a_changed_component_tag(feedback_env, indexes):
    _, _, v1, store, _, _, review = prepared(feedback_env, indexes)
    store.fingerprint = content_hash({"new-component": "same verified source semantics"})
    assert store.load(v1.task_id, v1.version_id, review.snapshot.snapshot_id) == review.snapshot
