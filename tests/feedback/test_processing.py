"""Controlled prose, source-derived Word scopes and replayed typed answers."""
from __future__ import annotations

from concurrent.futures import ThreadPoolExecutor
from hashlib import sha256
from threading import Barrier

from docx import Document
from fastapi.testclient import TestClient
import pytest

from autoformat.feedback_contracts import FeedbackConfirmationRequest
from autoformat.feedback_parser import interpret
from autoformat.feedback_processing import FeedbackProcessor
from autoformat.feedback_store import FeedbackError, FeedbackStore
from autoformat.local_api import create_local_api
from autoformat.rule_snapshots import RuleSnapshotStore
from autoformat.version_store import VersionStore, new_id
from conftest import publish_fixture, submission_for
from test_selection import indexes, answer_for, custom_task


@pytest.mark.parametrize("raw,prop,value,unit", [
    ("这里宋体", "font.family", "宋体", None), ("字体“Body Custom Font”", "font.family", "Body Custom Font", None),
    ("小四号", "font.size", 12, "pt"), ("字号10.5pt", "font.size", 10.5, "pt"),
    ("加粗", "font.bold", True, None), ("取消加粗", "font.bold", False, None),
    ("不使用斜体", "font.italic", False, None), ("居中", "paragraph.alignment", "center", None),
    ("两端对齐", "paragraph.alignment", "both", None),
    ("行距1.5倍", "paragraph.line_spacing", 1.5, "multiple"),
    ("双倍行距", "paragraph.line_spacing", 2, "multiple"),
    ("行距固定值20磅", "paragraph.line_spacing", {"mode": "exact", "amount": 20}, "pt"),
    ("段前6磅", "paragraph.spacing_before", 6, "pt"), ("段后1行", "paragraph.spacing_after", 1, "lines"),
    ("首行缩进2字符", "paragraph.first_line_indent", 2, "characters"),
    ("左缩进1cm", "paragraph.left_indent", 1, "cm"),
    ("上页边距2.5厘米", "page.margin.top", 2.5, "cm"), ("页边距25mm", "page.margin.all", 25, "mm"),
    ("纸张A4", "page.paper_size", "A4", None), ("横向", "page.orientation", "landscape", None),
    ("三线表", "table.border_style", "three_line", None),
    ("font size 12pt", "font.size", 12, "pt"), ("not bold", "font.bold", False, None),
    ("英文Arial", "font.family", "Arial", None),
])
def test_common_controlled_feedback(raw, prop, value, unit):
    facts = interpret(raw, "selection", "0" * 64)
    assert facts.kind == "new_rule", (raw, facts.missing_fields)
    assert len(facts.candidates) == 1
    candidate = facts.candidates[0]
    assert (candidate.property_id, candidate.value.normalized, candidate.value.unit) == (prop, value, unit)


@pytest.mark.parametrize("raw", ["不要居中", "宋体或黑体", "字号12到14pt", "如果是标题就加粗", "除了表格外加粗",
    "宋体但不改英文", "这里12pt并润色句子", "加粗且添加下划线", "第2章加粗", "0pt", "字号12", "把字体换成随便一种"])
def test_unknown_negative_conditional_and_incomplete_input_never_becomes_complete(raw):
    assert interpret(raw, "selection", "0" * 64).kind in {"clarification_needed", "out_of_scope"}


@pytest.mark.parametrize("raw,kind", [("为什么是宋体？", "question"), ("字号该用多少号", "question"),
    ("润色正文", "out_of_scope"), ("摘要不少于300字", "out_of_scope"), ("图片尺寸修改", "out_of_scope"),
    ("这里还是宋体", "clarification_needed")])
def test_non_instructions_do_not_create_candidates(raw, kind):
    result = interpret(raw, "selection", "0" * 64)
    assert result.kind == kind and not result.candidates


@pytest.mark.parametrize("raw", ["宋体12pt加粗", "字号12pt并且加粗", "字号12磅加粗"])
def test_adjacent_chinese_format_tokens_do_not_hide_size_or_bold(raw):
    facts = interpret(raw, "selection", "0" * 64)
    assert facts.kind == "new_rule"
    assert {c.property_id for c in facts.candidates} >= {"font.size", "font.bold"}


@pytest.mark.parametrize("raw", ["中文和英文居中", "中文和数字行距1.5倍", "中文和英文和数字段后6磅"])
def test_multiple_script_limits_on_paragraph_geometry_are_preserved_as_clarification(raw):
    facts = interpret(raw, "selection", "0" * 64)
    assert facts.kind == "clarification_needed" and "character_scope" in facts.missing_fields
    assert facts.candidates[0].applicability_scope.axes["characters"].include != ["chinese", "complex_script", "digits", "latin", "other"]


def saved(feedback, version, raw, **updates):
    context = feedback.context(version.task_id, version.version_id)
    submission = submission_for(context, raw_text=raw, **updates)
    receipt = feedback.submit(version.task_id, version.version_id, submission)
    return receipt, context, submission


def confirm_for(analysis, context, answer, **updates):
    data = dict(schema_version="feedback-confirmation-request/1.0", request_id=new_id(), task_id=context.task_id,
        base_version_id=context.base_version_id, expected_base_docx_sha256=context.base_docx_sha256,
        expected_context_id=context.context_id, feedback_id=analysis.interpretation.feedback_id,
        expected_parse_result_id=analysis.interpretation.result_id, raw_answer="确认所显示的补充或范围", answer=answer)
    data.update(updates)
    return FeedbackConfirmationRequest.model_validate(data)


def test_draft_is_read_only_and_displays_paragraph_promotion_before_submission(feedback_env, indexes):
    versions, feedback, _, v1, _ = feedback_env
    processor = FeedbackProcessor(feedback, indexes)
    context = feedback.context(v1.task_id, v1.version_id)
    request = submission_for(context, raw_text="这里居中")
    before = versions.artifact_path(v1.task_id, v1.version_id, "docx").read_bytes()
    result = processor.draft(v1.task_id, v1.version_id, request)
    assert result.kind == "new_rule" and result.target_status == "ready"
    review = result.scope_reviews[0]
    assert (review.minimum_unit, review.effective_scope, review.estimated_affected_count) == ("paragraph", "paragraph", 1)
    assert review.requires_confirmation and "整段" in review.reason
    assert feedback.list_receipts(v1.task_id, v1.version_id) == []
    assert versions.artifact_path(v1.task_id, v1.version_id, "docx").read_bytes() == before
    assert len(versions.list_versions(v1.task_id)) == 2


def test_scope_confirmation_is_bound_idempotent_and_survives_restart(feedback_env, indexes):
    versions, feedback, _, v1, _ = feedback_env
    receipt, context, _ = saved(feedback, v1, "这里居中")
    processor = FeedbackProcessor(feedback, indexes)
    first = processor.get(v1.task_id, v1.version_id, receipt.feedback_id)
    request = confirm_for(first, context, {"action": "confirm_scope", "accepted_review_ids": [first.interpretation.scope_reviews[0].review_id]})
    confirmed = processor.confirm(v1.task_id, v1.version_id, receipt.feedback_id, request)
    assert confirmed.confirmed_review_ids == request.answer.accepted_review_ids
    assert confirmed.decision_count == 1
    assert processor.confirm(v1.task_id, v1.version_id, receipt.feedback_id, request) == confirmed
    assert FeedbackProcessor(FeedbackStore(versions), indexes).get(v1.task_id, v1.version_id, receipt.feedback_id) == confirmed
    publish_fixture(versions, v1, versions.root.parent / "v2")
    assert processor.confirm(v1.task_id, v1.version_id, receipt.feedback_id, request) == confirmed
    with pytest.raises(FeedbackError):
        processor.confirm(v1.task_id, v1.version_id, receipt.feedback_id, request.model_copy(update={"request_id": new_id()}))


def test_partial_font_size_can_be_completed_without_rewriting_raw_feedback(feedback_env, indexes):
    versions, feedback, _, v1, _ = feedback_env
    raw = "  这里的字号需要调整\n"
    receipt, context, submission = saved(feedback, v1, raw)
    processor = FeedbackProcessor(feedback, indexes)
    first = processor.get(v1.task_id, v1.version_id, receipt.feedback_id)
    assert first.interpretation.kind == "clarification_needed"
    partial = first.interpretation.understood[0]
    request = confirm_for(first, context, {"action": "clarify_format", "format": {
        "candidate_id": partial.candidate_id, "property_id": "font.size", "value": {"normalized": 12, "unit": "pt"},
        "requested_scope": "selection", "semantic_role": None}})
    done = processor.confirm(v1.task_id, v1.version_id, receipt.feedback_id, request)
    assert done.interpretation.kind == "new_rule" and done.interpretation.candidates[0].value.normalized == 12
    assert feedback.get(v1.task_id, v1.version_id, receipt.feedback_id).submission == submission
    assert feedback.get(v1.task_id, v1.version_id, receipt.feedback_id).submission.raw_text == raw
    assert processor.get(v1.task_id, v1.version_id, receipt.feedback_id) == done
    assert len(versions.list_versions(v1.task_id)) == 2


def test_structured_answer_cannot_drop_unknown_limits_or_rewrite_known_property(feedback_env, indexes):
    _, feedback, _, v1, _ = feedback_env
    processor = FeedbackProcessor(feedback, indexes)
    for raw in ["如果是标题就修改字号", "这里的字号需要调整"]:
        receipt, context, _ = saved(feedback, v1, raw)
        first = processor.get(v1.task_id, v1.version_id, receipt.feedback_id)
        partial = first.interpretation.understood[0]
        request = confirm_for(first, context, {"action": "clarify_format", "format": {
            "candidate_id": partial.candidate_id, "property_id": "font.bold", "value": {"normalized": True}, "requested_scope": "selection"}})
        with pytest.raises(FeedbackError): processor.confirm(v1.task_id, v1.version_id, receipt.feedback_id, request)


def test_scope_mismatch_and_script_limits_remain_pending(feedback_env, indexes):
    _, feedback, _, v1, _ = feedback_env
    processor = FeedbackProcessor(feedback, indexes)
    for raw, field in [("所有二级标题居中", "requested_scope"), ("中文居中", "character_scope"), ("页边距2.5cm", "requested_scope")]:
        receipt, _, _ = saved(feedback, v1, raw)
        result = processor.get(v1.task_id, v1.version_id, receipt.feedback_id).interpretation
        assert result.kind == "clarification_needed" and field in result.missing_fields


def test_same_kind_targets_come_from_actual_role_index(tmp_path, indexes):
    source = tmp_path / "roles.docx"
    document = Document(); document.add_paragraph("Selectable preview text", style="Heading 2")
    document.add_paragraph("Another heading", style="Heading 2"); document.add_paragraph("Ordinary body")
    document.save(source)
    versions = VersionStore(tmp_path / "store"); v0 = versions.create_task(source)
    v1 = publish_fixture(versions, v0, tmp_path / "output"); feedback = FeedbackStore(versions)
    receipt, _, _ = saved(feedback, v1, "所有二级标题居中", requested_scope="same_kind")
    result = FeedbackProcessor(feedback, indexes).get(v1.task_id, v1.version_id, receipt.feedback_id).interpretation
    assert result.kind == "new_rule"
    assert result.candidates[0].semantic_role == "heading_2"
    assert result.scope_reviews[0].estimated_affected_count == 2 and result.scope_reviews[0].requires_confirmation


def test_whole_document_page_settings_use_real_section_ids(feedback_env, indexes):
    _, feedback, _, v1, _ = feedback_env
    receipt, _, _ = saved(feedback, v1, "全文页边距2.5cm", requested_scope="document", selection=None)
    result = FeedbackProcessor(feedback, indexes).get(v1.task_id, v1.version_id, receipt.feedback_id).interpretation
    assert result.kind == "new_rule"
    assert result.scope_reviews[0].minimum_unit == "section" and result.scope_reviews[0].requires_confirmation


def test_questions_and_content_do_not_start_index_processes(feedback_env):
    _, feedback, _, v1, _ = feedback_env
    class NoIndexes:
        fingerprint = "0" * 64
        def __getattr__(self, name): raise AssertionError("Consultation must not start indexing")
    processor = FeedbackProcessor(feedback, NoIndexes())
    for raw, kind in [("为什么是宋体？", "question"), ("润色这段文字", "out_of_scope")]:
        receipt, _, _ = saved(feedback, v1, raw)
        result = processor.get(v1.task_id, v1.version_id, receipt.feedback_id)
        assert result.interpretation.kind == kind and result.selection is None


def test_persisted_parse_result_cannot_replace_source_replay(feedback_env, indexes):
    versions, feedback, _, v1, _ = feedback_env
    receipt, _, _ = saved(feedback, v1, "加粗")
    processor = FeedbackProcessor(feedback, indexes); result = processor.get(v1.task_id, v1.version_id, receipt.feedback_id)
    with versions._transaction() as connection:
        connection.execute("UPDATE feedback_parse_results SET payload_json=? WHERE feedback_id=?", ("{}", receipt.feedback_id))
    with pytest.raises(FeedbackError, match="重放不一致"): processor.get(v1.task_id, v1.version_id, receipt.feedback_id)


def test_analysis_routes_require_session_and_reject_client_target_ids(feedback_env, indexes):
    versions, feedback, _, v1, _ = feedback_env
    receipt, context, request = saved(feedback, v1, "加粗")
    origin, token = "http://127.0.0.1:57843", "processing-test-session-token-1234567890"
    app = create_local_api(store=versions, snapshots=RuleSnapshotStore(versions), indexer=None, editor=None,
        renderer_factory=lambda _: None, session_token=token, allowed_origin=origin, font_fingerprint="test-fonts", selection_indexes=indexes)
    prefix = f"/api/tasks/{v1.task_id}/versions/{v1.version_id}/feedback"
    with TestClient(app, base_url=origin) as client:
        headers = {"X-AutoFormat-Session": token}
        assert client.post(prefix + "/review", json=request.model_dump(mode="json")).status_code == 401
        draft = client.post(prefix + "/review", headers=headers, json=request.model_dump(mode="json"))
        assert draft.status_code == 200 and draft.json()["scope_reviews"][0]["minimum_unit"] == "character"
        payload = request.model_dump(mode="json"); payload["target_ids"] = ["o:" + "a" * 24]
        assert client.post(prefix + "/review", headers=headers, json=payload).status_code == 422
        response = client.get(prefix + f"/{receipt.feedback_id}/analysis", headers=headers)
        assert response.status_code == 200 and response.json()["interpretation"]["kind"] == "new_rule"
        assert "word/document.xml" not in response.text and "xml_path" not in response.text


def test_two_processors_retrying_one_confirmation_append_one_decision(feedback_env, indexes):
    versions, feedback, _, v1, _ = feedback_env
    receipt, context, _ = saved(feedback, v1, "这里居中")
    processors = [FeedbackProcessor(FeedbackStore(versions), indexes) for _ in range(2)]
    first = processors[0].get(v1.task_id, v1.version_id, receipt.feedback_id)
    request = confirm_for(first, context, {"action": "confirm_scope", "accepted_review_ids": [first.interpretation.scope_reviews[0].review_id]})
    gate = Barrier(2)
    for processor in processors:
        original = processor._apply
        def synchronize(*args, apply=original, seen=[]):
            apply(*args)
            if not seen:
                seen.append(True); gate.wait(timeout=30)
        processor._apply = synchronize
    with ThreadPoolExecutor(max_workers=2) as pool:
        results = list(pool.map(lambda processor: processor.confirm(v1.task_id, v1.version_id, receipt.feedback_id, request), processors))
    assert results[0] == results[1] and results[0].decision_count == 1
    with versions._transaction() as connection:
        assert connection.execute("SELECT COUNT(*) FROM feedback_processing_decisions WHERE feedback_id=?", (receipt.feedback_id,)).fetchone()[0] == 1


def test_multiple_partial_answers_replay_without_changing_any_version_artifact(feedback_env, indexes):
    versions, feedback, v0, v1, _ = feedback_env
    before = {(v.version_id, kind): versions.artifact_path(v.task_id, v.version_id, kind).read_bytes()
        for v in [v0, v1] for kind in (["docx"] if v.kind == "original" else ["docx", "preview", "report_json", "report_markdown"])}
    receipt, context, original = saved(feedback, v1, "这里字体和字号需要调整")
    processor = FeedbackProcessor(feedback, indexes)
    analysis = processor.get(v1.task_id, v1.version_id, receipt.feedback_id)
    assert len(analysis.interpretation.understood) == 2
    for prop, value in [("font.family", {"normalized": "宋体"}), ("font.size", {"normalized": 12, "unit": "pt"})]:
        partial = next(p for p in analysis.interpretation.understood if p.property_id == prop)
        request = confirm_for(analysis, context, {"action": "clarify_format", "format": {
            "candidate_id": partial.candidate_id, "property_id": prop, "value": value, "requested_scope": "selection"}})
        analysis = processor.confirm(v1.task_id, v1.version_id, receipt.feedback_id, request)
    assert analysis.interpretation.kind == "new_rule" and analysis.decision_count == 2
    assert FeedbackProcessor(FeedbackStore(versions), indexes).get(v1.task_id, v1.version_id, receipt.feedback_id) == analysis
    assert feedback.get(v1.task_id, v1.version_id, receipt.feedback_id).submission == original
    for (version, kind), content in before.items():
        assert versions.artifact_path(v1.task_id, version, kind).read_bytes() == content


def test_ambiguous_target_must_be_confirmed_before_structured_completion(tmp_path, indexes):
    _, feedback, v1 = custom_task(tmp_path, ["Selectable preview text", "Selectable preview text"])
    receipt, context, _ = saved(feedback, v1, "修改字号")
    processor = FeedbackProcessor(feedback, indexes)
    initial = processor.get(v1.task_id, v1.version_id, receipt.feedback_id)
    partial = initial.interpretation.understood[0]
    def complete(analysis):
        return confirm_for(analysis, context, {"action": "clarify_format", "format": {
            "candidate_id": partial.candidate_id, "property_id": "font.size", "value": {"normalized": 12, "unit": "pt"}, "requested_scope": "selection"}})
    with pytest.raises(FeedbackError): processor.confirm(v1.task_id, v1.version_id, receipt.feedback_id, complete(initial))
    processor.selections.confirm(v1.task_id, v1.version_id, receipt.feedback_id, answer_for(receipt, context, initial.selection))
    resolved = processor.get(v1.task_id, v1.version_id, receipt.feedback_id)
    done = processor.confirm(v1.task_id, v1.version_id, receipt.feedback_id, complete(resolved))
    assert done.interpretation.kind == "new_rule"
    assert done.interpretation.scope_reviews[0].target_ids == [done.selection.anchor.ranges[0].object_id]


def test_scope_confirmation_rejects_stale_result_unknown_review_and_request_collision(feedback_env, indexes):
    _, feedback, _, v1, _ = feedback_env
    receipt, context, _ = saved(feedback, v1, "这里居中")
    processor = FeedbackProcessor(feedback, indexes)
    first = processor.get(v1.task_id, v1.version_id, receipt.feedback_id)
    request = confirm_for(first, context, {"action": "confirm_scope", "accepted_review_ids": [first.interpretation.scope_reviews[0].review_id]})
    for updates in [{"expected_parse_result_id": "f" * 64}, {"expected_base_docx_sha256": "f" * 64}]:
        with pytest.raises(FeedbackError): processor.confirm(v1.task_id, v1.version_id, receipt.feedback_id, request.model_copy(update=updates))
    unknown = confirm_for(first, context, {"action": "confirm_scope", "accepted_review_ids": ["f" * 64]})
    with pytest.raises(FeedbackError): processor.confirm(v1.task_id, v1.version_id, receipt.feedback_id, unknown)
    done = processor.confirm(v1.task_id, v1.version_id, receipt.feedback_id, request)
    with pytest.raises(FeedbackError, match="请求 ID"):
        processor.confirm(v1.task_id, v1.version_id, receipt.feedback_id, request.model_copy(update={"raw_answer": "不同的答案"}))
    assert processor.get(v1.task_id, v1.version_id, receipt.feedback_id) == done


def historical_rule(role="body", character="all", legacy=False):
    from brief_parser.domain.academic_format import AcademicFormatRule
    from brief_parser.compilers.academic_scope import base_scope
    data = dict(schema_version="academic-format-rule/1.1", rule_id=f"rule-{role}-{character}", source_requirement_id="r1",
        source_kind="brief_explicit", source_artifact_id="brief1", source_version="brief:v1",
        evidence_refs=[{"evidence_id": "ev1", "source_id": "brief1", "block_id": "p1", "start": 0, "end": 4,
            "quote": "正文宋体", "quote_sha256": sha256("正文宋体".encode()).hexdigest(), "locator": {"paragraph": 0}, "match_status": "exact"}],
        target={"object_kind": "paragraph", "semantic_role": role, "selector": "semantic_role"}, property_id="font.family",
        value={"original_text": "宋体", "normalized": "宋体", "dictionary_version": "test"}, character_scope=character,
        requested_scope="same_kind", effective_scope="same_kind", status="verified", reason="Confirmed fixture",
        semantics={"scope": base_scope(role, "font.family", character), "constraint": {"operator": "eq"}})
    if legacy: data.update(schema_version="academic-format-rule/1.0", semantics=None)
    return AcademicFormatRule.model_validate(data)


@pytest.mark.parametrize("rule,expected", [(historical_rule(), True), (historical_rule(character="latin"), False),
    (historical_rule(role="heading_2"), False), (historical_rule(legacy=True), False)])
def test_execution_defect_only_references_proven_applicable_selected_version_rules(tmp_path, indexes, rule, expected):
    document = Document(); document.add_paragraph("Selectable preview text", style="Body Text")
    source = tmp_path / "body.docx"; document.save(source)
    versions = VersionStore(tmp_path / "store"); v0 = versions.create_task(source)
    v1 = publish_fixture(versions, v0, tmp_path / "v1"); feedback = FeedbackStore(versions)
    receipt, context, _ = saved(feedback, v1, "这里还是宋体")
    loaded = []
    def selected_rules(selected_context):
        loaded.append(selected_context); return [rule, rule]
    analysis = FeedbackProcessor(feedback, indexes, rule_loader=selected_rules).get(v1.task_id, v1.version_id, receipt.feedback_id)
    result = analysis.interpretation
    assert loaded == [context]
    assert result.kind == ("execution_defect" if expected else "clarification_needed")
    assert result.existing_rule_ids == ([rule.rule_id] if expected else [])
    assert not result.candidates and not result.scope_reviews


@pytest.mark.parametrize("scope,expected", [("selection", "clarification_needed"), ("paragraph", "new_rule"), ("document", "new_rule")])
def test_script_scope_counts_only_objects_with_matching_characters_in_actual_range(tmp_path, indexes, scope, expected):
    _, feedback, v1 = custom_task(tmp_path, ["Selectable preview text中文𠀀", "Another English paragraph"])
    receipt, _, _ = saved(feedback, v1, "中文宋体", requested_scope=scope, **({"selection": None} if scope == "document" else {}))
    result = FeedbackProcessor(feedback, indexes).get(v1.task_id, v1.version_id, receipt.feedback_id).interpretation
    assert result.kind == expected
    if expected == "new_rule":
        assert result.scope_reviews[0].estimated_affected_count == 1
        assert "符合所示字符条件" in result.scope_reviews[0].reason
    else:
        assert not result.scope_reviews and "selection_target" in result.missing_fields


@pytest.mark.parametrize("text,script", [("中", "chinese"), ("𠀀", "chinese"), ("Ａ", "chinese"), ("9", "digits"),
    ("é", "latin"), ("A", "latin"), ("ا", "complex_script"), ("😀", "other"), (" ", "other")])
def test_character_classification_matches_editor_ranges(text, script):
    from autoformat.feedback_scope import character_kind
    assert character_kind(text) == script
