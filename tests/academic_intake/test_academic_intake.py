"""Parse, confirm or explicitly skip, then publish an immutable V1."""

from __future__ import annotations

import asyncio
import hashlib
import json
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from docx.shared import Pt
from fastapi.testclient import TestClient
from jsonschema import Draft202012Validator
from pypdf import PdfReader

from brief_parser.bootstrap import build_service
from brief_parser.compilers.template_layout import build_template_groups
from brief_parser.config import Settings
from brief_parser.domain.academic_bundle import AcademicParseBundle
from brief_parser.errors import BriefParserError, ErrorCode

from autoformat.academic_intake import AcademicTaskIntake
from autoformat.initial_draft import InitialDraftPublisher
from autoformat.local_api import create_local_api
from autoformat.preview import PreviewArtifact, PreviewError
from autoformat.rule_snapshots import RuleSnapshotStore
from autoformat.version_store import VersionStore, VersionStoreError, new_id


ROOT = Path(__file__).resolve().parents[2]
ORIGIN = "http://127.0.0.1:57843"
TOKEN = "academic-intake-session-token-long-enough"
DECISION_HEADERS = {"X-AutoFormat-Session": TOKEN, "Origin": ORIGIN}
BRIEF = "正文段落必须居中。".encode("utf-8")


def _decision(client, task_id, queue, *, action="select_rule", request_id=None, question=None, **answer_fields):
    pending = question or next(item for item in queue["questions"] if item["status"] == "pending")
    data = {"request_id": request_id or new_id(), "base_version_id": queue["base_version_id"],
            "expected_snapshot_id": queue["snapshot_id"], "question_id": pending["question_id"],
            "answer": {"action": action, **answer_fields}}
    return client.post(f"/api/tasks/{task_id}/decisions", headers=DECISION_HEADERS, json=data), data


def _decision_task(setup, brief="正文：宋体。\n正文：黑体。"):
    client, store, intake, paper = setup
    created = _post(client, paper, request_id=new_id(), brief=brief.encode("utf-8")).json()
    task_id = created["task_id"]
    queue = client.get(f"/api/tasks/{task_id}/questions", headers=DECISION_HEADERS)
    assert queue.status_code == 200, queue.text
    return client, store, intake, created, queue.json()


def _finish_review(client, payload):
    task_id = payload["task_id"]
    for _ in range(256):
        state = client.get(f"/api/tasks/{task_id}/initial-review", headers=DECISION_HEADERS).json()
        if state["phase"] != "reviewing_requirements":
            break
        response, _ = _decision(client, task_id, state, action="skip")
        assert response.status_code == 200, response.text
    else:
        pytest.fail("Review failed to terminate within its decision budget.")
    response = client.post(f"/api/tasks/{task_id}/initial-draft", headers=DECISION_HEADERS, json={
        "schema_version": "initial-draft-request/1.0", "base_version_id": state["base_version_id"],
        "expected_base_docx_sha256": state["base_docx_sha256"], "expected_snapshot_id": state["snapshot_id"],
    })
    assert response.status_code == 200, response.text
    version_id = response.json()["initial_version_id"]
    return {**payload, "initial_version_id": version_id,
            "initial_docx_sha256": hashlib.sha256(client.get(
                f"/api/tasks/{task_id}/versions/{version_id}/artifacts/docx", headers=DECISION_HEADERS).content).hexdigest()}


def test_parser_decision_roundtrip_preserves_source_and_paper_history(setup):
    from autoformat.academic_decisions import AcademicDecisionStore
    client, store, intake, created, queue = _decision_task(setup)
    task_id = created["task_id"]
    original_parse = intake.load_bundle(task_id).model_dump_json()
    files = {item.version_id: store.artifact_path(task_id, item.version_id, "docx").read_bytes() for item in store.list_versions(task_id)}
    conflict = next(item for item in queue["questions"] if item["kind"] == "conflict")
    response, data = _decision(client, task_id, queue, question=conflict, selected_id=conflict["rule_ids"][0])
    assert response.status_code == 200, response.text
    result = response.json()
    assert result["revision"] == 1 and not result["replayed"]
    snapshot = client.get(f"/api/tasks/{task_id}/resolved-rules", headers=DECISION_HEADERS).json()
    assert snapshot["snapshot_id"] == result["snapshot_id"] and not snapshot["conflicts"]
    assert snapshot["rules"][0]["status"] == "verified" and snapshot["rules"][0]["decision_ids"]
    for kind, value in (("academic-question-queue", queue), ("academic-decision-snapshot", snapshot),
                        ("academic-rule-decision", snapshot["events"][0]),
                        ("academic-decision-request", snapshot["events"][0]["request"]),
                        ("academic-effective-rule", snapshot["rules"][0])):
        schema = json.loads((ROOT / "contracts" / kind / f"{kind}-1.0.schema.json").read_text(encoding="utf-8"))
        Draft202012Validator.check_schema(schema)
        Draft202012Validator(schema).validate(value)
    restarted = AcademicDecisionStore(VersionStore(store.root), intake).load(task_id)
    assert restarted.model_dump(mode="json") == snapshot
    history = client.get(f"/api/tasks/{task_id}/rule-decisions/snapshots", headers=DECISION_HEADERS).json()["snapshots"]
    assert [item["revision"] for item in history] == [0, 1]
    baseline = client.get(f"/api/tasks/{task_id}/rule-decisions/snapshots/{queue['snapshot_id']}", headers=DECISION_HEADERS).json()
    assert baseline["revision"] == 0 and baseline["conflicts"]
    report = client.get(f"/api/tasks/{task_id}/rule-decisions/report", headers=DECISION_HEADERS)
    assert report.status_code == 200 and "Brief 解析决策报告" in report.text
    assert len(store.list_decisions(task_id)) == 1 and len(store.list_versions(task_id)) == 1
    assert intake.load_bundle(task_id).model_dump_json() == original_parse
    assert files == {item.version_id: store.artifact_path(task_id, item.version_id, "docx").read_bytes() for item in store.list_versions(task_id)}
    assert store.working_version_id(task_id) == created["base_version_id"]


def test_parser_decision_request_is_idempotent_and_rejects_payload_reuse(setup):
    client, store, _, created, queue = _decision_task(setup)
    task_id = created["task_id"]
    conflict = next(item for item in queue["questions"] if item["kind"] == "conflict")
    response, data = _decision(client, task_id, queue, question=conflict, selected_id=conflict["rule_ids"][0])
    replay = client.post(f"/api/tasks/{task_id}/decisions", headers=DECISION_HEADERS, json=data)
    assert replay.status_code == 200 and replay.json()["replayed"]
    assert replay.json()["snapshot_id"] == response.json()["snapshot_id"]
    changed = {**data, "answer": {"action": "select_rule", "selected_id": conflict["rule_ids"][1]}}
    rejected = client.post(f"/api/tasks/{task_id}/decisions", headers=DECISION_HEADERS, json=changed)
    assert rejected.status_code == 409
    assert len(store.list_decisions(task_id)) == 1


def test_parser_decision_rejects_stale_snapshot_atomically(setup):
    client, store, _, created, queue = _decision_task(setup)
    task_id = created["task_id"]
    conflict = next(item for item in queue["questions"] if item["kind"] == "conflict")
    response, _ = _decision(client, task_id, queue, question=conflict, action="defer")
    assert response.status_code == 200, response.text
    expected_count = 1
    rejected, _ = _decision(client, task_id, queue, question=conflict, selected_id=conflict["rule_ids"][0])
    assert rejected.status_code == 409, rejected.text
    assert len(store.list_decisions(task_id)) == expected_count
    history = client.get(f"/api/tasks/{task_id}/rule-decisions/snapshots", headers=DECISION_HEADERS).json()["snapshots"]
    assert len(history) == expected_count + 1


def test_parser_typed_value_and_question_ownership_are_enforced(setup):
    client, store, _, created, queue = _decision_task(setup, "正文：字号12-14pt。")
    task_id = created["task_id"]
    response, data = _decision(client, task_id, queue, action="clarify_format", format={"property_id": "font.size", "normalized": 18, "unit": "pt"})
    assert response.status_code == 409 and not store.list_decisions(task_id)
    data["question_id"] = "fmtq_unknown"
    rejected = client.post(f"/api/tasks/{task_id}/decisions", headers=DECISION_HEADERS, json=data)
    assert rejected.status_code == 422 and not store.list_decisions(task_id)
    for invalid in ({"action": "confirm_condition", "applies": "true"}, {"action": "skip", "selected_id": "unexpected"}):
        data["answer"] = invalid
        assert client.post(f"/api/tasks/{task_id}/decisions", headers=DECISION_HEADERS, json=data).status_code == 422


@pytest.mark.parametrize("target", ["snapshot", "event", "chain"])
def test_parser_history_tampering_is_detected(setup, target):
    client, store, _, created, queue = _decision_task(setup)
    task_id = created["task_id"]
    conflict = next(item for item in queue["questions"] if item["kind"] == "conflict")
    response, _ = _decision(client, task_id, queue, question=conflict, selected_id=conflict["rule_ids"][0])
    assert response.status_code == 200
    with store._transaction() as connection:
        if target == "snapshot":
            connection.execute("UPDATE academic_decision_snapshots SET payload_json='{}' WHERE task_id=? AND revision=1", (task_id,))
        elif target == "event":
            connection.execute("UPDATE decisions SET payload_json='{}' WHERE task_id=?", (task_id,))
        else:
            connection.execute("UPDATE academic_decision_snapshots SET parent_snapshot_id=NULL WHERE task_id=? AND revision=1", (task_id,))
    rejected = client.get(f"/api/tasks/{task_id}/resolved-rules", headers=DECISION_HEADERS)
    assert rejected.status_code == 500 and rejected.json()["status"] == "validation_failed"


def test_parser_question_loop_opens_at_v0_with_session_protection(setup):
    client, store, intake, paper = setup
    result = asyncio.run(intake.create(request_id=new_id(), brief_filename="school.txt", brief="正文：字号。".encode(),
        paper_filename="paper.docx", paper=paper, model_processing_acknowledged=True))
    task_id = result.version.task_id
    rejected = client.get(f"/api/tasks/{task_id}/questions", headers=DECISION_HEADERS)
    assert rejected.status_code == 200 and rejected.json()["questions"]
    assert rejected.json()["base_version_id"] == result.version.version_id
    assert client.get(f"/api/tasks/{task_id}/questions").status_code == 401
    with store._transaction() as connection:
        assert connection.execute("SELECT * FROM academic_decision_heads WHERE task_id=?", (task_id,)).fetchone()


def test_parser_revision_endpoint_preserves_old_answer(setup):
    client, store, _, created, queue = _decision_task(setup)
    task_id = created["task_id"]
    conflict = next(item for item in queue["questions"] if item["kind"] == "conflict")
    first, first_request = _decision(client, task_id, queue, question=conflict, selected_id=conflict["rule_ids"][0])
    decision_id = first.json()["decision_id"]
    reopened = client.get(f"/api/tasks/{task_id}/questions?reopen_decision_id={decision_id}", headers=DECISION_HEADERS).json()
    conflict = next(item for item in reopened["questions"] if item["kind"] == "conflict" and item["status"] == "pending")
    data = {"request_id": new_id(), "base_version_id": reopened["base_version_id"], "expected_snapshot_id": reopened["snapshot_id"],
            "question_id": conflict["question_id"], "supersedes_decision_id": decision_id,
            "answer": {"action": "select_rule", "selected_id": conflict["rule_ids"][1]}}
    revised = client.post(f"/api/tasks/{task_id}/decisions", headers=DECISION_HEADERS, json=data)
    assert revised.status_code == 200, revised.text
    result = client.get(f"/api/tasks/{task_id}/resolved-rules", headers=DECISION_HEADERS).json()
    assert [effect["status"] for effect in result["decision_effects"]] == ["superseded", "applied"]
    assert len(store.list_decisions(task_id)) == 2
    replay = client.get(f"/api/tasks/{task_id}/rule-decisions/snapshots/{first.json()['snapshot_id']}", headers=DECISION_HEADERS).json()
    assert replay["decision_effects"][0]["status"] == "applied"
    retried = client.post(f"/api/tasks/{task_id}/decisions", headers=DECISION_HEADERS, json=first_request)
    assert retried.status_code == 200 and retried.json()["snapshot_id"] == first.json()["snapshot_id"]
    current = client.get(f"/api/tasks/{task_id}/resolved-rules", headers=DECISION_HEADERS).json()
    assert current["snapshot_id"] == revised.json()["snapshot_id"] and len(store.list_decisions(task_id)) == 2


def test_parser_concurrent_answers_publish_one_snapshot(setup, monkeypatch):
    from concurrent.futures import ThreadPoolExecutor
    from threading import Barrier
    from brief_parser.application.academic_decisions import AcademicDecisionEngine
    client, store, _, created, queue = _decision_task(setup)
    task_id = created["task_id"]
    conflict = next(item for item in queue["questions"] if item["kind"] == "conflict")
    barrier = Barrier(2)
    original = AcademicDecisionEngine.submit

    def simultaneously(self, *args, **kwargs):
        result = original(self, *args, **kwargs)
        barrier.wait(timeout=10)
        return result

    monkeypatch.setattr(AcademicDecisionEngine, "submit", simultaneously)
    with ThreadPoolExecutor(max_workers=2) as pool:
        responses = list(pool.map(lambda index: _decision(client, task_id, queue, question=conflict,
            selected_id=conflict["rule_ids"][index])[0], (0, 1)))
    assert sorted(response.status_code for response in responses) == [200, 409]
    assert len(store.list_decisions(task_id)) == 1
    history = client.get(f"/api/tasks/{task_id}/rule-decisions/snapshots", headers=DECISION_HEADERS).json()["snapshots"]
    assert len(history) == 2


def test_parser_answer_cannot_cross_tasks(setup):
    client, store, _, first, queue = _decision_task(setup)
    _, _, _, second, _ = _decision_task(setup)
    conflict = next(item for item in queue["questions"] if item["kind"] == "conflict")
    rejected, _ = _decision(client, second["task_id"], queue, question=conflict, selected_id=conflict["rule_ids"][0])
    assert rejected.status_code == 409 and not store.list_decisions(second["task_id"])
    assert not store.list_decisions(first["task_id"])


def _pdf(path: Path) -> None:
    content = b"BT /F1 14 Tf 72 720 Td (Selectable preview text) Tj ET"
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream",
    ]
    output = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for number, body in enumerate(objects, 1):
        offsets.append(len(output))
        output.extend(f"{number} 0 obj\n".encode() + body + b"\nendobj\n")
    startxref = len(output)
    output.extend(f"xref\n0 {len(offsets)}\n0000000000 65535 f \n".encode())
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode())
    output.extend((f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\n"
                   f"startxref\n{startxref}\n%%EOF\n").encode())
    path.write_bytes(output)


class FixtureRenderer:
    def __init__(self, root: Path) -> None:
        self.root = root

    def render(self, source_docx: Path, artifact_id: str) -> PreviewArtifact:
        relative = f"artifacts/previews/{artifact_id}.pdf"
        pdf = self.root / relative
        pdf.parent.mkdir(parents=True)
        _pdf(pdf)
        reader = PdfReader(pdf, strict=True)
        return PreviewArtifact(
            schema_version="preview-artifact/1.0", artifact_id=artifact_id,
            source_sha256=hashlib.sha256(source_docx.read_bytes()).hexdigest(),
            pdf_sha256=hashlib.sha256(pdf.read_bytes()).hexdigest(),
            pdf_relative_path=relative, page_count=len(reader.pages),
            selectable_text_chars=sum(len(page.extract_text() or "") for page in reader.pages),
            renderer_id="fixture", renderer_version="1.0",
        )


def _paper(tmp_path: Path) -> bytes:
    path = tmp_path / "paper.docx"
    document = Document()
    document.add_paragraph("Original thesis body.")
    document.save(path)
    return path.read_bytes()


@pytest.fixture
def setup(tmp_path: Path):
    store = VersionStore(tmp_path / "runtime")
    snapshots = RuleSnapshotStore(store)
    service = build_service(Settings(
        environment="test", provider="fake", parser_isolation_mode="in_process",
    ))
    intake = AcademicTaskIntake(store=store, service=service)
    app = create_local_api(
        store=store, snapshots=snapshots, indexer=None, editor=None,
        renderer_factory=FixtureRenderer, session_token=TOKEN,
        allowed_origin=ORIGIN, font_fingerprint="test-fonts", intake=intake,
    )
    with TestClient(app, base_url=ORIGIN) as client:
        yield client, store, intake, _paper(tmp_path)


def _post(client: TestClient, paper: bytes, *, request_id: str, brief: bytes = BRIEF,
          acknowledged: str = "true"):
    return client.post(
        "/api/tasks", headers={"X-AutoFormat-Session": TOKEN, "Origin": ORIGIN},
        data={"request_id": request_id,
              "model_processing_acknowledged": acknowledged},
        files={
            "brief": ("school.txt", brief, "text/plain"),
            "paper": ("paper.docx", paper,
                      "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
        },
    )


def test_academic_intake_persists_parse_and_v0_without_executable_rules(setup):
    client, store, intake, paper = setup
    request_id = new_id()
    response = _post(client, paper, request_id=request_id)
    assert response.status_code == 200, response.text
    payload = response.json()
    schema = json.loads((ROOT / "contracts" / "task-intake" /
                         "task-intake-1.2.schema.json").read_text(encoding="utf-8"))
    Draft202012Validator.check_schema(schema)
    Draft202012Validator(schema).validate(payload)
    assert payload["replayed"] is False
    assert payload["phase"] == "initial_draft_ready"
    assert payload["initial_format_status"] == "unassessed"
    assert "未作格式修改" in payload["next_action"]
    assert payload["base_docx_sha256"] == hashlib.sha256(paper).hexdigest()
    assert [item.kind for item in store.list_versions(payload["task_id"])] == [
        "original", "output"]
    assert payload["initial_docx_sha256"] == payload["base_docx_sha256"]
    assert store.get_academic_intake(payload["task_id"]).brief_path.read_bytes() == BRIEF
    assert not (store.root / "rule_snapshots" / f"{payload['task_id']}.json").exists()

    headers = {"X-AutoFormat-Session": TOKEN, "Origin": ORIGIN}
    task = client.get(f"/api/tasks/{payload['task_id']}", headers=headers)
    assert task.status_code == 200
    assert task.json()["phase"] == payload["phase"]
    assert task.json()["working_version_id"] == payload["initial_version_id"]
    assert task.json()["initial_format_status"] == "unassessed"
    downloaded = client.get(
        f"/api/tasks/{payload['task_id']}/versions/{payload['initial_version_id']}/artifacts/docx",
        headers=headers)
    assert downloaded.status_code == 200 and downloaded.content == paper
    report = json.loads(store.artifact_path(
        payload["task_id"], payload["initial_version_id"], "report_json").read_text(encoding="utf-8"))
    assert report["summary"]["this_round_change_count"] == 0
    assert "未修改任何格式" in report["summary"]["note"]
    parsed = client.get(f"/api/tasks/{payload['task_id']}/academic-parse",
                        headers=headers)
    assert parsed.status_code == 200
    assert parsed.json()["bundle_version"] == "academic-parse-bundle/1.6"
    assert parsed.json()["source_review"]["policy_version"] == "academic-format-scope/1.2"
    ruleset = parsed.json()["format_ruleset"]
    assert ruleset["rules"][0]["property_id"] == "paragraph.alignment"
    assert ruleset["rules"][0]["value"]["normalized"] == "center"
    assert parsed.json()["manifest"]["sources"][0]["sha256"] == hashlib.sha256(BRIEF).hexdigest()
    parse_schema = json.loads((ROOT / "contracts" / "academic-parse-bundle" /
                                   "academic-parse-bundle-1.6.schema.json").read_text(encoding="utf-8"))
    Draft202012Validator(parse_schema).validate(parsed.json())
    assert intake.load_bundle(payload["task_id"]).manifest.requested_task_types == [
        "format_academic_paper"]

    replay = _post(client, paper, request_id=request_id)
    assert replay.status_code == 200
    assert replay.json()["task_id"] == payload["task_id"]
    assert replay.json()["replayed"] is True
    assert replay.json()["initial_version_id"] == payload["initial_version_id"]
    assert len(store.list_versions(payload["task_id"])) == 2
    reopened = VersionStore(store.root)
    replay_after_restart = asyncio.run(AcademicTaskIntake(
        store=reopened, service=object(),
    ).create(
        request_id=request_id, brief_filename="school.txt", brief=BRIEF,
        paper_filename="paper.docx", paper=paper,
        model_processing_acknowledged=True,
    ))
    assert replay_after_restart.replayed is True
    assert replay_after_restart.version.task_id == payload["task_id"]
    assert [item.kind for item in reopened.list_versions(payload["task_id"])] == [
        "original", "output"]
    changed = _post(client, paper + b"x", request_id=request_id)
    assert changed.status_code == 409


def test_count_and_content_only_brief_does_not_create_initial_questions(setup):
    client, store, _, paper = setup
    response = _post(client, paper, request_id=new_id(),
                     brief="论文不少于10000字。摘要应包括研究目的和结论。".encode("utf-8"))
    assert response.status_code == 200, response.text
    payload = response.json()
    report = json.loads(store.artifact_path(
        payload["task_id"], payload["initial_version_id"], "report_json").read_text(encoding="utf-8"))
    assert report["summary"]["pending_count"] == 0
    assert not report["issues"]
    assert payload["initial_docx_sha256"] == payload["base_docx_sha256"]


def test_initial_report_contains_format_gap_and_its_source_location(setup):
    client, store, _, paper = setup
    response = _post(client, paper, request_id=new_id(), brief="正文：宋体，段前自动间距。".encode("utf-8"))
    assert response.status_code == 200, response.text
    payload = _finish_review(client, response.json())
    report = json.loads(store.artifact_path(
        payload["task_id"], payload["initial_version_id"], "report_json").read_text(encoding="utf-8"))
    assert any("用户暂时忽略" in item["reason"] and "段前自动间距" in item["reason"]
               and "paragraph" in item["reason"] for item in report["issues"])


@pytest.mark.parametrize("brief,has_conflict", [
    ("正文：纸质版宋体；电子版黑体。", False),
    ("正文：宋体。英文字母：Times New Roman。正文：字号12-14pt。", True),
])
def test_initial_report_preserves_condition_or_intersection_and_range(setup, brief, has_conflict):
    client, store, _, paper = setup
    response = _post(client, paper, request_id=new_id(), brief=brief.encode("utf-8"))
    assert response.status_code == 200, response.text
    payload = response.json()
    assert payload["phase"] == "reviewing_requirements" and payload["initial_version_id"] is None
    questions = client.get(f"/api/tasks/{payload['task_id']}/questions", headers=DECISION_HEADERS).json()["questions"]
    conflicts = [question for question in questions if question["kind"] == "conflict"]
    assert bool(conflicts) == has_conflict
    if has_conflict:
        assert all("latin" in json.dumps(question["scope"]) for question in conflicts)
        assert any("12" in json.dumps(question["candidate_values"]) and "14" in json.dumps(question["candidate_values"]) for question in questions)
    else:
        assert any(question["kind"] == "submission" for question in questions)
    payload = _finish_review(client, payload)
    assert payload["initial_docx_sha256"] == payload["base_docx_sha256"]


def test_intake_rejects_unacknowledged_invalid_and_oversized_inputs(setup):
    client, store, _, paper = setup
    assert _post(client, paper, request_id=new_id(), acknowledged="false").status_code == 422
    assert _post(client, b"not-a-docx", request_id=new_id()).status_code == 422
    assert _post(client, paper, request_id=new_id(), brief=b"\x00bad").status_code == 422
    assert client.post("/api/tasks").status_code == 401
    with store._transaction() as connection:
        assert connection.execute("SELECT COUNT(*) FROM tasks").fetchone()[0] == 0


def test_intake_detects_stored_brief_tampering(setup):
    client, store, intake, paper = setup
    response = _post(client, paper, request_id=new_id())
    assert response.status_code == 200
    task_id = response.json()["task_id"]
    record = store.get_academic_intake(task_id)
    record.brief_path.write_bytes(b"changed")
    headers = {"X-AutoFormat-Session": TOKEN, "Origin": ORIGIN}
    result = client.get(f"/api/tasks/{task_id}/academic-parse", headers=headers)
    assert result.status_code == 500
    assert client.get(f"/api/tasks/{task_id}", headers=headers).status_code == 500
    download = client.get(
        f"/api/tasks/{task_id}/versions/{response.json()['base_version_id']}/artifacts/docx",
        headers=headers)
    assert download.status_code == 500
    store.recover()
    assert store.get_version(task_id, response.json()["base_version_id"]).status == "validation_failed"
    assert store.get_version(task_id, response.json()["initial_version_id"]).status == "validation_failed"


def test_parse_failure_does_not_create_v0(tmp_path: Path):
    class FailingParser:
        async def parse_academic_with_timeout(self, request, *, timeout_seconds):
            raise BriefParserError(ErrorCode.PROVIDER_TRANSIENT, "Temporary parse failure.")

    store = VersionStore(tmp_path / "runtime")
    intake = AcademicTaskIntake(store=store, service=FailingParser())
    with pytest.raises(BriefParserError):
        asyncio.run(intake.create(
            request_id=new_id(), brief_filename="school.txt", brief=BRIEF,
            paper_filename="paper.docx", paper=_paper(tmp_path),
            model_processing_acknowledged=True,
        ))
    with store._transaction() as connection:
        assert connection.execute("SELECT COUNT(*) FROM tasks").fetchone()[0] == 0


def test_initial_preview_failure_keeps_v0_and_retry_publishes_v1(tmp_path: Path):
    class FailingRenderer:
        def render(self, source_docx, artifact_id):
            raise PreviewError("render_failed", "Synthetic preview failure.")

    store = VersionStore(tmp_path / "runtime")
    service = build_service(Settings(
        environment="test", provider="fake", parser_isolation_mode="in_process",
    ))
    intake = AcademicTaskIntake(store=store, service=service)
    result = asyncio.run(intake.create(
        request_id=new_id(), brief_filename="school.txt", brief=BRIEF,
        paper_filename="paper.docx", paper=_paper(tmp_path),
        model_processing_acknowledged=True,
    ))
    failing = InitialDraftPublisher(
        store=store, renderer_factory=lambda _: FailingRenderer(),
        font_fingerprint="test-fonts")
    with pytest.raises(PreviewError):
        failing.publish(result.version, result.bundle)
    assert [item.kind for item in store.list_versions(result.version.task_id)] == ["original"]
    assert store.working_version_id(result.version.task_id) == result.version.version_id

    working = InitialDraftPublisher(
        store=store, renderer_factory=FixtureRenderer, font_fingerprint="test-fonts")
    first = working.publish(result.version, result.bundle)
    assert first.kind == "output"
    assert working.publish(result.version, result.bundle).version_id == first.version_id


def _template_brief():
    document = Document()
    document.add_paragraph("摘要")
    document.add_paragraph("摘要示例").runs[0].font.size = Pt(12)
    document.add_paragraph("目录")
    document.add_heading("1 绪论", level=1)
    document.add_paragraph("正文示例").runs[0].font.size = Pt(12)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_docx_template_intake_reopens_and_reports_without_editing_paper(setup):
    client, store, intake, paper = setup
    brief = _template_brief()
    response = client.post("/api/tasks", headers={"X-AutoFormat-Session": TOKEN, "Origin": ORIGIN},
                           data={"request_id": new_id(), "model_processing_acknowledged": "true"},
                           files={"brief": ("学校模板.docx", brief), "paper": ("paper.docx", paper)})
    assert response.status_code == 200, response.text
    payload = response.json()
    bundle = intake.load_bundle(payload["task_id"])
    assert bundle.template_layout.observations and bundle.template_layout.groups
    assert all(item.source_sha256 == hashlib.sha256(brief).hexdigest() for item in bundle.template_layout.observations)
    assert not any(rule.source_kind == "template_observation" for rule in bundle.format_ruleset.rules)
    assert payload["phase"] == "reviewing_requirements" and payload["initial_version_id"] is None
    payload = _finish_review(client, payload)
    assert payload["initial_docx_sha256"] == payload["base_docx_sha256"] == hashlib.sha256(paper).hexdigest()
    report = json.loads(store.artifact_path(payload["task_id"], payload["initial_version_id"], "report_json").read_text(encoding="utf-8"))
    assert any("模板" in item["reason"] and "用户暂时忽略" in item["reason"] for item in report["issues"])
    reopened = AcademicTaskIntake(store=VersionStore(store.root), service=object())
    assert reopened.load_bundle(payload["task_id"]) == bundle
    assert store.get_academic_intake(payload["task_id"]).brief_path.read_bytes() == brief


def test_intake_rejects_forged_observation_even_when_groups_match(tmp_path):
    service = build_service(Settings(environment="test", provider="fake", parser_isolation_mode="in_process"))

    class ForgingParser:
        async def parse_academic_with_timeout(self, request, *, timeout_seconds):
            bundle = await service.parse_academic_with_timeout(request, timeout_seconds=timeout_seconds)
            data = bundle.model_dump(mode="python")
            observation = next(item for item in data["template_layout"]["observations"] if item["property_id"] == "font.size")
            observation["value"] = 42
            layout = bundle.template_layout.model_copy(update={"observations": [
                item.model_copy(update={"value": 42}) if item.observation_id == observation["observation_id"] else item
                for item in bundle.template_layout.observations]})
            data["template_layout"]["groups"] = [item.model_dump(mode="python") for item in build_template_groups(layout.observations, bundle.format_ruleset)]
            return AcademicParseBundle.model_validate(data)

    store = VersionStore(tmp_path / "runtime")
    intake = AcademicTaskIntake(store=store, service=ForgingParser())
    with pytest.raises(VersionStoreError, match="Template observations differ"):
        asyncio.run(intake.create(request_id=new_id(), brief_filename="school.docx", brief=_template_brief(),
                                  paper_filename="paper.docx", paper=_paper(tmp_path), model_processing_acknowledged=True))
    with store._transaction() as connection:
        assert connection.execute("SELECT COUNT(*) FROM tasks").fetchone()[0] == 0


def test_pending_review_blocks_editing_and_survives_restart(setup):
    from autoformat.academic_decisions import AcademicDecisionStore
    from autoformat.initial_review import review_state
    client, store, intake, payload, queue = _decision_task(setup)
    task_id = payload["task_id"]
    schema = json.loads((ROOT / "contracts/task-intake/task-intake-1.2.schema.json").read_text(encoding="utf-8"))
    Draft202012Validator(schema).validate(payload)
    assert payload["status"] == "needs_user_input" and payload["initial_version_id"] is None
    assert [version.kind for version in store.list_versions(task_id)] == ["original"]
    state = client.get(f"/api/tasks/{task_id}/initial-review", headers=DECISION_HEADERS).json()
    for kind in ("initial-review",):
        Draft202012Validator(json.loads((ROOT / f"contracts/{kind}/{kind}-1.0.schema.json").read_text(encoding="utf-8"))).validate(state)
    body = {"base_version_id": state["base_version_id"], "expected_base_docx_sha256": state["base_docx_sha256"],
            "expected_snapshot_id": state["snapshot_id"]}
    rejected = client.post(f"/api/tasks/{task_id}/initial-draft", headers=DECISION_HEADERS, json=body)
    assert rejected.status_code == 422 and len(store.list_versions(task_id)) == 1
    assert client.post(f"/api/tasks/{task_id}/initial-draft", json=body).status_code == 401
    reopened = VersionStore(store.root)
    assert review_state(reopened, AcademicDecisionStore(reopened, intake), task_id) == state
    assert client.get(f"/api/tasks/{task_id}", headers=DECISION_HEADERS).json()["phase"] == "reviewing_requirements"


@pytest.mark.parametrize("operation", ["review", "confirm", "execute"])
def test_legacy_local_edit_cannot_bypass_initial_review(setup, operation):
    client, store, intake, payload, queue = _decision_task(setup)
    url = f"/api/tasks/{payload['task_id']}/versions/{payload['base_version_id']}/body-alignment/{operation}"
    if operation == "review":
        response = client.get(url, headers=DECISION_HEADERS)
    else:
        body = {"schema_version": "body-scope-confirmation/1.0" if operation == "confirm" else "body-plan-execution/1.0",
                "request_id": new_id(), "expected_base_docx_sha256": payload["base_docx_sha256"]}
        if operation == "confirm":
            body.update(review_sha256="a" * 64, body_object_ids=["p:000001"],
                        excluded_object_ids=[], raw_answer="确认正文")
        else:
            body.update(decision_id=new_id(), plan_id="a" * 64)
        response = client.post(url, headers=DECISION_HEADERS, json=body)
    assert response.status_code == 422
    assert response.json()["status"] == "needs_user_input"
    assert "编辑前确认" in response.json()["reason"]
    assert len(store.list_versions(payload["task_id"])) == 1


def test_skip_then_publish_uses_decision_snapshot_and_replays_once(setup):
    client, store, intake, payload, queue = _decision_task(setup)
    original_bytes = store.artifact_path(payload["task_id"], payload["base_version_id"], "docx").read_bytes()
    final = _finish_review(client, payload)
    assert len(store.list_versions(payload["task_id"])) == 2
    snapshot = client.get(f"/api/tasks/{payload['task_id']}/resolved-rules", headers=DECISION_HEADERS).json()
    assert snapshot["initial_version_id"] == payload["base_version_id"]
    assert all(question["status"] in {"answered", "skipped"} for question in snapshot["questions"])
    record = store.get_version(payload["task_id"], final["initial_version_id"])
    assert record.references.rules_snapshot_id == snapshot["snapshot_id"]
    report = json.loads(store.artifact_path(payload["task_id"], record.version_id, "report_json").read_text(encoding="utf-8"))
    assert report["decisions"] and any("暂时忽略" in issue["reason"] for issue in report["issues"])
    assert _finish_review(client, payload)["initial_version_id"] == record.version_id
    assert store.artifact_path(payload["task_id"], payload["base_version_id"], "docx").read_bytes() == original_bytes
    assert len(store.list_versions(payload["task_id"])) == 2


def test_initial_publication_rejects_stale_snapshot_and_hash(setup):
    client, store, _, payload, queue = _decision_task(setup)
    task_id = payload["task_id"]
    conflict = next(question for question in queue["questions"] if question["kind"] == "conflict")
    response, _ = _decision(client, task_id, queue, question=conflict, selected_id=conflict["rule_ids"][0])
    assert response.status_code == 200
    body = {"base_version_id": queue["base_version_id"], "expected_base_docx_sha256": payload["base_docx_sha256"],
            "expected_snapshot_id": queue["snapshot_id"]}
    assert client.post(f"/api/tasks/{task_id}/initial-draft", headers=DECISION_HEADERS, json=body).status_code == 409
    body["expected_snapshot_id"] = response.json()["snapshot_id"]
    body["expected_base_docx_sha256"] = "0" * 64
    assert client.post(f"/api/tasks/{task_id}/initial-draft", headers=DECISION_HEADERS, json=body).status_code == 409
    assert len(store.list_versions(task_id)) == 1


def test_v0_answer_cannot_write_after_v1_becomes_working_version(setup):
    client, store, _, payload, _ = _decision_task(setup)
    final = _finish_review(client, payload)
    task_id = payload["task_id"]
    snapshot = client.get(f"/api/tasks/{task_id}/resolved-rules", headers=DECISION_HEADERS).json()
    event = next(event for event in snapshot["events"] if event["question_kind"] == "conflict")
    reopened = client.get(f"/api/tasks/{task_id}/questions?reopen_decision_id={event['decision_id']}", headers=DECISION_HEADERS).json()
    question = next(question for question in reopened["questions"] if question["kind"] == "conflict" and question["status"] == "pending")
    body = {"request_id": new_id(), "base_version_id": payload["base_version_id"],
            "expected_snapshot_id": snapshot["snapshot_id"], "question_id": question["question_id"],
            "supersedes_decision_id": event["decision_id"],
            "answer": {"action": "select_rule", "selected_id": question["rule_ids"][0]}}
    response = client.post(f"/api/tasks/{task_id}/decisions", headers=DECISION_HEADERS, json=body)
    assert response.status_code == 409 and response.json()["status"] == "stale_base_version"
    assert len(store.list_decisions(task_id)) == len(snapshot["events"])
    assert store.working_version_id(task_id) == final["initial_version_id"]


def test_decisions_changed_while_rendering_cannot_publish_stale_v1(setup):
    from autoformat.academic_decisions import AcademicDecisionStore
    from brief_parser.domain.academic_decisions import AcademicDecisionRequest
    client, store, intake, payload, queue = _decision_task(setup)
    task_id = payload["task_id"]
    decisions = AcademicDecisionStore(store, intake)
    while True:
        snapshot = decisions.load(task_id)
        question = next((question for question in snapshot.questions if question.status == "pending"), None)
        if question is None:
            break
        decisions.submit(task_id, AcademicDecisionRequest(request_id=new_id(), base_version_id=payload["base_version_id"],
            expected_snapshot_id=snapshot.snapshot_id, question_id=question.question_id, answer={"action": "skip"}))
    old = snapshot
    conflict_event = next(event for event in old.events if event.question_kind == "conflict")

    class ChangingRenderer(FixtureRenderer):
        def render(self, docx_path, artifact_id):
            reopened = decisions.questions(task_id, conflict_event.decision_id)
            question = next(question for question in reopened["questions"] if question["kind"] == "conflict" and question["status"] == "pending")
            decisions.submit(task_id, AcademicDecisionRequest(request_id=new_id(), base_version_id=payload["base_version_id"],
                expected_snapshot_id=old.snapshot_id, question_id=question["question_id"],
                supersedes_decision_id=conflict_event.decision_id,
                answer={"action": "select_rule", "selected_id": question["rule_ids"][0]}))
            return super().render(docx_path, artifact_id)

    publisher = InitialDraftPublisher(store=store, renderer_factory=ChangingRenderer, font_fingerprint="test-fonts")
    with pytest.raises(VersionStoreError, match="编辑期间"):
        publisher.publish(store.get_version(task_id, payload["base_version_id"]), intake.load_bundle(task_id), old.snapshot_id)
    assert len(store.list_versions(task_id)) == 1
    assert store.working_version_id(task_id) == payload["base_version_id"]
