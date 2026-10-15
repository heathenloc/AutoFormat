"""Immutable publication, branching, recovery, and fault checks."""

from __future__ import annotations

import hashlib
import json
import os
import shutil
import sqlite3
from dataclasses import asdict, replace
from pathlib import Path
from uuid import uuid4

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pypdf import PdfReader
from jsonschema import Draft202012Validator

from autoformat.preview import LibreOfficeRenderer, PreviewArtifact
from autoformat.reports import ReportContext, build_report, render_markdown, write_report
from autoformat.version_store import (
    VersionReferences,
    VersionStore,
    VersionStoreError,
    new_id,
)


def _hash(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _pdf(path: Path) -> None:
    """Make a tiny text PDF without a second rendering dependency."""
    content = b"BT /F1 14 Tf 72 720 Td (Selectable preview text) Tj ET"
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream",
    ]
    output = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for number, body in enumerate(objects, 1):
        offsets.append(len(output))
        output.extend(f"{number} 0 obj\n".encode() + body + b"\nendobj\n")
    cross_reference = len(output)
    output.extend(f"xref\n0 {len(offsets)}\n0000000000 65535 f \n".encode())
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode())
    output.extend(
        f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\n"
        f"startxref\n{cross_reference}\n%%EOF\n".encode()
    )
    path.write_bytes(output)


def _source(path: Path) -> None:
    document = Document()
    document.add_heading("A thesis heading", level=1)
    document.add_paragraph("A paragraph that remains semantically unchanged.")
    document.save(path)


def _report_files(
    directory: Path, *, store: VersionStore, parent, version_id: str,
    docx: Path, artifact: PreviewArtifact,
) -> tuple[Path, Path]:
    report = directory / "report.json"
    markdown = directory / "report.md"
    parent_report = (json.loads(store.artifact_path(parent.task_id, parent.version_id,
                                                    "report_json").read_text(encoding="utf-8"))
                     if parent.kind == "output" else None)
    rules = list(parent_report["rules"]) if parent_report else []
    previous_changes = list(parent_report["changes_cumulative"]) if parent_report else []
    decisions = list(parent_report["decisions"]) if parent_report else []
    changed = _hash(docx) != parent.docx_sha256
    changes = []
    if changed:
        rule_id = f"rule-{version_id}"
        verification_ref = f"fixture-verification-{version_id}"
        rules.append({
            "rule_id": rule_id, "source_kind": "brief_explicit",
            "source_artifact_id": "fixture-brief", "source_version": "brief-1",
            "original_requirement": "测试用例要求正文段落居中",
            "evidence_refs": [{
                "source_artifact_id": "fixture-brief", "source_version": "brief-1",
                "locator": "fixture:1", "quote": "正文段落居中",
            }],
            "object_id": "p:000002", "semantic_role": "body",
            "property_id": "paragraph.alignment",
            "requested_scope": "paragraph", "effective_scope": "paragraph",
            "before_value": None, "after_value": "center",
            "status": "modified_verified", "reason": "Fixture result verified by caller.",
            "next_action": None, "decision_id": None,
            "verification_ref": verification_ref,
        })
        changes.append({
            "change_id": f"change-{version_id}", "rule_id": rule_id,
            "object_id": "p:000002", "property_id": "paragraph.alignment",
            "before_value": None, "after_value": "center",
            "effective_scope": "paragraph", "validation_status": "verified",
            "verification_ref": verification_ref,
        })
    context = ReportContext(
        task_id=parent.task_id, version_id=version_id,
        parent_version_id=parent.version_id,
        source_docx_sha256=parent.docx_sha256,
        output_docx_sha256=_hash(docx), brief_version_id="brief-1",
        rules_snapshot_id="rules-1", edit_plan_id="plan-1",
        preview_artifact_id=artifact.artifact_id,
        renderer_id=artifact.renderer_id,
        renderer_version=artifact.renderer_version, font_fingerprint="font-1",
    )
    payload = build_report(
        context, rules=rules, changes_this_round=changes,
        changes_cumulative=previous_changes + changes,
        decisions=decisions, issues=[],
        note=("Fixture candidate has a verified format change." if changed else
              "No confirmed executable changes; draft remains unassessed."),
    )
    write_report(report, markdown, payload)
    return report, markdown


def _candidate_set(
    directory: Path, *, store: VersionStore, parent, version_id: str,
    docx: Path,
) -> tuple[Path, Path, PreviewArtifact]:
    directory.mkdir()
    pdf = directory / "preview.pdf"
    docx_hash = _hash(docx)
    _pdf(pdf)
    with pdf.open("rb") as stream:
        reader = PdfReader(stream, strict=True)
        page_count = len(reader.pages)
        text_chars = sum(len(page.extract_text() or "") for page in reader.pages)
    artifact_id = new_id()
    artifact = PreviewArtifact(
        schema_version="preview-artifact/1.0", artifact_id=artifact_id,
        source_sha256=docx_hash, pdf_sha256=_hash(pdf),
        pdf_relative_path=f"artifacts/previews/{artifact_id}.pdf",
        page_count=page_count, selectable_text_chars=text_chars,
        renderer_id="fixture", renderer_version="test-1",
    )
    report, markdown = _report_files(
        directory, store=store, parent=parent, version_id=version_id,
        docx=docx, artifact=artifact,
    )
    return report, markdown, artifact


def _publish(
    store: VersionStore, *, parent, candidate: Path, directory: Path,
    version_id: str | None = None, request_id: str | None = None,
):
    version_id = version_id or new_id()
    request_id = request_id or new_id()
    report, markdown, artifact = _candidate_set(
        directory, store=store, parent=parent, version_id=version_id,
        docx=candidate,
    )
    result = store.publish_version(
        task_id=parent.task_id, version_id=version_id,
        parent_version_id=parent.version_id,
        expected_parent_sha256=parent.docx_sha256, request_id=request_id,
        candidate_docx=candidate, report_json=report, report_markdown=markdown,
        preview_pdf=directory / "preview.pdf", preview_artifact=artifact,
        references=VersionReferences("rules-1", "brief-1", "plan-1", "font-1"),
    )
    return result, (report, markdown, artifact, request_id)


def test_v0_v1_branch_and_restart_keep_every_artifact(tmp_path: Path) -> None:
    source = tmp_path / "source.docx"
    _source(source)
    original_hash = _hash(source)
    store = VersionStore(tmp_path / "store")
    v0 = store.create_task(source)
    assert v0.sequence == 0 and v0.parent_version_id is None
    assert store.artifact_path(v0.task_id, v0.version_id, "docx") != source
    assert _hash(source) == original_hash == v0.docx_sha256

    v1, (report, markdown, preview, request_id) = _publish(
        store, parent=v0, candidate=source, directory=tmp_path / "v1"
    )
    schema = json.loads((Path(__file__).resolve().parents[2] / "contracts" /
                         "version-node" / "version-node-1.0.schema.json").read_text(encoding="utf-8"))
    Draft202012Validator.check_schema(schema)
    Draft202012Validator(schema).validate(asdict(v0))
    Draft202012Validator(schema).validate(asdict(v1))
    assert v1.sequence == 1 and v1.docx_sha256 == v0.docx_sha256
    assert v1.parent_version_id == v0.version_id
    for kind in ("docx", "report_json", "report_markdown", "preview", "preview_metadata"):
        assert store.artifact_path(v1.task_id, v1.version_id, kind).is_file()
    assert store.working_version_id(v0.task_id) == v1.version_id

    duplicate = store.publish_version(
        task_id=v0.task_id, version_id=v1.version_id,
        parent_version_id=v0.version_id, expected_parent_sha256=v0.docx_sha256,
        request_id=request_id, candidate_docx=source,
        report_json=report, report_markdown=markdown,
        preview_pdf=tmp_path / "v1" / "preview.pdf", preview_artifact=preview,
        references=VersionReferences("rules-1", "brief-1", "plan-1", "font-1"),
    )
    assert duplicate.version_id == v1.version_id and len(store.list_versions(v0.task_id)) == 2
    with pytest.raises(VersionStoreError) as request_conflict:
        store.publish_version(
            task_id=v0.task_id, version_id=v1.version_id,
            parent_version_id=v0.version_id, expected_parent_sha256=v0.docx_sha256,
            request_id=request_id, candidate_docx=source,
            report_json=report, report_markdown=markdown,
            preview_pdf=tmp_path / "v1" / "preview.pdf", preview_artifact=preview,
            references=VersionReferences("rules-1", "brief-1", "plan-1", "other-font"),
        )
    assert request_conflict.value.status == "blocked_conflict"

    store.select_working_version(v0.task_id, v0.version_id)
    formatted = tmp_path / "formatted.docx"
    shutil.copyfile(source, formatted)
    document = Document(formatted)
    document.paragraphs[1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.save(formatted)
    v2, _ = _publish(store, parent=v0, candidate=formatted, directory=tmp_path / "v2")
    assert v2.parent_version_id == v0.version_id and v2.docx_sha256 != v0.docx_sha256
    assert store.working_version_id(v0.task_id) == v2.version_id
    assert [v.version_id for v in store.list_versions(v0.task_id)] == [
        v0.version_id, v1.version_id, v2.version_id,
    ]
    restarted = VersionStore(tmp_path / "store")
    assert restarted.working_version_id(v0.task_id) == v2.version_id
    assert restarted.artifact_path(v1.task_id, v1.version_id, "preview").is_file()
    assert restarted.artifact_path(v0.task_id, v0.version_id, "docx").is_file()
    assert _hash(source) == original_hash


def test_no_change_after_v1_and_stale_parent_do_not_publish(tmp_path: Path) -> None:
    source = tmp_path / "source.docx"
    _source(source)
    store = VersionStore(tmp_path / "store")
    v0 = store.create_task(source)
    v1, _ = _publish(store, parent=v0, candidate=source, directory=tmp_path / "v1")
    with pytest.raises(VersionStoreError) as stale:
        _publish(store, parent=v0, candidate=source, directory=tmp_path / "stale")
    assert stale.value.status == "stale_base_version"
    with pytest.raises(VersionStoreError) as unchanged:
        _publish(store, parent=v1, candidate=source, directory=tmp_path / "unchanged")
    assert unchanged.value.status == "no_document_change"
    assert len(store.list_versions(v0.task_id)) == 2
    assert store.working_version_id(v0.task_id) == v1.version_id


def test_database_failure_and_restart_recovery_hide_partial_versions(tmp_path: Path) -> None:
    source = tmp_path / "source.docx"
    _source(source)
    store = VersionStore(tmp_path / "store")
    v0 = store.create_task(source)
    with sqlite3.connect(store.database) as connection:
        connection.execute("""
            CREATE TRIGGER fail_output BEFORE INSERT ON versions
            WHEN NEW.kind = 'output' BEGIN SELECT RAISE(ABORT, 'fault injection'); END
        """)
    with pytest.raises(sqlite3.IntegrityError):
        _publish(store, parent=v0, candidate=source, directory=tmp_path / "failed")
    assert [v.version_id for v in store.list_versions(v0.task_id)] == [v0.version_id]
    assert store.working_version_id(v0.task_id) == v0.version_id
    assert not list((store.root / ".staging").iterdir())
    assert len(list((store.root / "tasks" / v0.task_id / "versions").iterdir())) == 1

    orphan_id = uuid4().hex
    (store.root / ".staging" / uuid4().hex).mkdir()
    (store.root / "tasks" / v0.task_id / "versions" / orphan_id).mkdir()
    restarted = VersionStore(store.root)
    assert not list((store.root / ".staging").iterdir())
    assert not (store.root / "tasks" / v0.task_id / "versions" / orphan_id).exists()
    assert restarted.artifact_path(v0.task_id, v0.version_id, "docx").is_file()


def test_tampering_marks_version_invalid_after_restart(tmp_path: Path) -> None:
    source = tmp_path / "source.docx"
    _source(source)
    store = VersionStore(tmp_path / "store")
    v0 = store.create_task(source)
    v1, _ = _publish(store, parent=v0, candidate=source, directory=tmp_path / "v1")
    report = store.artifact_path(v1.task_id, v1.version_id, "report_markdown")
    report.write_text("tampered", encoding="utf-8")
    with pytest.raises(VersionStoreError) as invalid:
        store.artifact_path(v1.task_id, v1.version_id, "report_markdown")
    assert invalid.value.status == "validation_failed"
    restarted = VersionStore(store.root)
    assert restarted.get_version(v1.task_id, v1.version_id).status == "validation_failed"
    assert restarted.artifact_path(v0.task_id, v0.version_id, "docx").is_file()


def test_preview_mismatch_never_publishes_a_version(tmp_path: Path) -> None:
    source = tmp_path / "source.docx"
    _source(source)
    store = VersionStore(tmp_path / "store")
    v0 = store.create_task(source)
    version_id = new_id()
    directory = tmp_path / "candidate"
    report, markdown, artifact = _candidate_set(
        directory, store=store, parent=v0, version_id=version_id,
        docx=source,
    )
    with pytest.raises(VersionStoreError) as error:
        store.publish_version(
            task_id=v0.task_id, version_id=version_id,
            parent_version_id=v0.version_id,
            expected_parent_sha256=v0.docx_sha256, request_id=new_id(),
            candidate_docx=source, report_json=report, report_markdown=markdown,
            preview_pdf=directory / "preview.pdf",
            preview_artifact=replace(artifact, pdf_sha256="0" * 64),
            references=VersionReferences("rules-1", "brief-1", "plan-1", "font-1"),
        )
    assert error.value.status == "render_failed"
    assert [v.version_id for v in store.list_versions(v0.task_id)] == [v0.version_id]
    assert not list((store.root / ".staging").iterdir())


def test_decision_log_is_append_only_without_redundant_docx(tmp_path: Path) -> None:
    source = tmp_path / "source.docx"
    _source(source)
    store = VersionStore(tmp_path / "store")
    v0 = store.create_task(source)
    request_id = new_id()
    payload = {"question_id": "q-1", "raw_answer": "暂缓这项格式要求"}
    decision = store.record_decision(
        task_id=v0.task_id, base_version_id=v0.version_id,
        request_id=request_id, kind="defer", payload=payload,
    )
    schema = json.loads((Path(__file__).resolve().parents[2] / "contracts" /
                         "decision-event" / "decision-event-1.0.schema.json").read_text(encoding="utf-8"))
    Draft202012Validator.check_schema(schema)
    Draft202012Validator(schema).validate(asdict(decision))
    same = store.record_decision(
        task_id=v0.task_id, base_version_id=v0.version_id,
        request_id=request_id, kind="defer", payload=payload,
    )
    assert same.decision_id == decision.decision_id
    with pytest.raises(VersionStoreError) as conflict:
        store.record_decision(
            task_id=v0.task_id, base_version_id=v0.version_id,
            request_id=request_id, kind="skip", payload=payload,
        )
    assert conflict.value.status == "blocked_conflict"
    assert len(store.list_versions(v0.task_id)) == 1
    restarted = VersionStore(store.root)
    assert restarted.list_decisions(v0.task_id) == [decision]
    assert restarted.artifact_path(v0.task_id, v0.version_id, "docx").is_file()


def test_real_libreoffice_preview_can_publish_v1() -> None:
    executable = os.environ.get("AUTOFORMAT_TEST_SOFFICE")
    expected_hash = os.environ.get("AUTOFORMAT_TEST_SOFFICE_SHA256")
    if not executable or not expected_hash:
        pytest.skip("Set the verified LibreOffice path and SHA-256 for integration.")
    root = Path(__file__).resolve().parents[2] / "runtime" / "tests" / "version-store-lo" / uuid4().hex
    root.mkdir(parents=True)
    source = root / "source.docx"
    _source(source)
    store = VersionStore(root / "store")
    v0 = store.create_task(source)
    render_root = root / "render"
    render_root.mkdir()
    render_input = render_root / "source.docx"
    shutil.copyfile(source, render_input)
    renderer = LibreOfficeRenderer(
        soffice_path=Path(executable), expected_soffice_sha256=expected_hash,
        data_root=render_root, timeout_seconds=60,
    )
    artifact = renderer.render(render_input, new_id())
    version_id = new_id()
    report, markdown = _report_files(
        root, store=store, parent=v0, version_id=version_id,
        docx=source, artifact=artifact,
    )
    v1 = store.publish_version(
        task_id=v0.task_id, version_id=version_id,
        parent_version_id=v0.version_id,
        expected_parent_sha256=v0.docx_sha256, request_id=new_id(),
        candidate_docx=source, report_json=report, report_markdown=markdown,
        preview_pdf=render_root / artifact.pdf_relative_path,
        preview_artifact=artifact,
        references=VersionReferences("rules-1", "brief-1", "plan-1", "font-1"),
    )
    assert v1.sequence == 1 and v1.preview_artifact_id == artifact.artifact_id
    assert _hash(store.artifact_path(v1.task_id, v1.version_id, "preview")) == artifact.pdf_sha256
    assert _hash(store.artifact_path(v0.task_id, v0.version_id, "docx")) == v0.docx_sha256
    assert not list((render_root / "preview-work").iterdir())


def test_markdown_tamper_and_lost_parent_history_block_publication(tmp_path: Path) -> None:
    source = tmp_path / "source.docx"
    _source(source)
    store = VersionStore(tmp_path / "store")
    v0 = store.create_task(source)
    first_id = new_id()
    first_dir = tmp_path / "bad-markdown"
    report, markdown, artifact = _candidate_set(
        first_dir, store=store, parent=v0, version_id=first_id, docx=source,
    )
    markdown.write_text("# Unrelated text\n", encoding="utf-8")
    with pytest.raises(VersionStoreError) as mismatch:
        store.publish_version(
            task_id=v0.task_id, version_id=first_id,
            parent_version_id=v0.version_id,
            expected_parent_sha256=v0.docx_sha256, request_id=new_id(),
            candidate_docx=source, report_json=report, report_markdown=markdown,
            preview_pdf=first_dir / "preview.pdf", preview_artifact=artifact,
            references=VersionReferences("rules-1", "brief-1", "plan-1", "font-1"),
        )
    assert mismatch.value.status == "validation_failed"
    assert len(store.list_versions(v0.task_id)) == 1

    formatted = tmp_path / "formatted.docx"
    shutil.copyfile(source, formatted)
    document = Document(formatted)
    document.paragraphs[1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.save(formatted)
    v1, _ = _publish(store, parent=v0, candidate=formatted, directory=tmp_path / "v1")
    second_id = new_id()
    second_dir = tmp_path / "lost-history"
    report, markdown, artifact = _candidate_set(
        second_dir, store=store, parent=v1, version_id=second_id, docx=formatted,
    )
    original = json.loads(report.read_text(encoding="utf-8"))
    context = ReportContext(
        task_id=original["task_id"], version_id=original["version_id"],
        parent_version_id=original["parent_version_id"],
        source_docx_sha256=original["input"]["source_docx_sha256"],
        output_docx_sha256=original["output_docx_sha256"],
        brief_version_id=original["input"]["brief_version_id"],
        rules_snapshot_id=original["input"]["rules_snapshot_id"],
        edit_plan_id=original["input"]["edit_plan_id"],
        preview_artifact_id=original["preview"]["artifact_id"],
        renderer_id=original["preview"]["renderer_id"],
        renderer_version=original["preview"]["renderer_version"],
        font_fingerprint=original["preview"]["font_fingerprint"],
        created_at=original["created_at"],
    )
    dropped = build_report(
        context, rules=[], changes_this_round=[], changes_cumulative=[],
        decisions=[], issues=[], note="Incorrectly dropped parent history.",
    )
    report.write_text(json.dumps(dropped, ensure_ascii=False), encoding="utf-8")
    markdown.write_text(render_markdown(dropped), encoding="utf-8")
    with pytest.raises(VersionStoreError) as history:
        store.publish_version(
            task_id=v1.task_id, version_id=second_id,
            parent_version_id=v1.version_id,
            expected_parent_sha256=v1.docx_sha256, request_id=new_id(),
            candidate_docx=formatted, report_json=report, report_markdown=markdown,
            preview_pdf=second_dir / "preview.pdf", preview_artifact=artifact,
            references=VersionReferences("rules-1", "brief-1", "plan-1", "font-1"),
        )
    assert history.value.status == "validation_failed"
    assert len(store.list_versions(v0.task_id)) == 2
