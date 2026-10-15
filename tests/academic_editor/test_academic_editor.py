"""Real parser -> replayed snapshot -> .NET -> immutable V1 integration checks."""
from __future__ import annotations

import asyncio
import hashlib
import json
import shutil
import subprocess
from pathlib import Path

import pytest
from docx import Document
from docx.shared import Pt
from jsonschema import Draft202012Validator

from brief_parser.application.academic_decisions import AcademicDecisionEngine
from brief_parser.bootstrap import build_service
from brief_parser.config import Settings
from brief_parser.domain.academic_decisions import AcademicDecisionRequest, content_hash
from brief_parser.validators.academic_decisions import ENUMS, BOOLEANS, LENGTHS, INTEGERS
from autoformat.academic_editor import AcademicEditor, report_facts
from autoformat.academic_intake import AcademicTaskIntake
from autoformat.initial_draft import InitialDraftPublisher
from autoformat.reports import ReportContext, build_report
from autoformat.version_store import VersionStore, new_id
from autoformat.workflow import WorkflowError

ROOT = Path(__file__).resolve().parents[2]
DLL = ROOT / "services/docx_engine/AutoFormat.DocxEditor/bin/Debug/net10.0/AutoFormat.DocxEditor.dll"


@pytest.fixture(scope="module")
def editor():
    dotnet = shutil.which("dotnet")
    assert dotnet, "The .NET editor must be built for the required integration checks."
    subprocess.run([dotnet, "build", str(ROOT / "services/docx_engine/AutoFormat.DocxEditor"), "--no-restore"], cwd=ROOT, capture_output=True, check=True)
    return AcademicEditor(dotnet_path=Path(dotnet), editor_dll=DLL,
                          expected_editor_sha256=hashlib.sha256(DLL.read_bytes()).hexdigest())


def task(tmp_path, brief):
    store = VersionStore(tmp_path / "versions")
    intake = AcademicTaskIntake(store=store, service=build_service(Settings(environment="test", provider="fake", parser_isolation_mode="in_process")))
    paper = tmp_path / "paper.docx"
    doc = Document()
    doc.add_heading("第1章 绪论", level=1)
    doc.add_paragraph("正文中文Ab12", style="Body Text")
    doc.add_heading("第2章 结果", level=1)
    doc.add_paragraph("第二章正文Ab12", style="Body Text")
    doc.add_paragraph("表1 实验数据")
    table = doc.add_table(rows=2, cols=2)
    for index, cell in enumerate(table._cells):
        cell.text = f"数据{index}Ab"
    doc.add_paragraph("参考文献")
    doc.add_paragraph("[1] Existing reference.")
    doc.save(paper)
    result = asyncio.run(intake.create(request_id=new_id(), brief_filename="school.txt", brief=brief.encode("utf-8"),
                                      paper_filename="paper.docx", paper=paper.read_bytes(), model_processing_acknowledged=True))
    engine = AcademicDecisionEngine(result.bundle, task_id=result.version.task_id, initial_version_id=new_id())
    return store, intake, result, engine


def execute(editor, tmp_path, result, engine, snapshot=None, plan_only=False):
    work = tmp_path / new_id()
    work.mkdir()
    (work / "parent.docx").write_bytes((tmp_path / "paper.docx").read_bytes())
    output = editor.execute(work, bundle=result.bundle, snapshot=snapshot or engine.compile(),
                            base_version_id=result.version.version_id, plan_only=plan_only)
    return work, output


def test_current_registry_and_both_languages_use_the_same_contract(editor):
    raw = subprocess.run([str(editor.dotnet_path), str(DLL), "capabilities"], capture_output=True, text=True, check=True)
    capabilities = json.loads(raw.stdout)
    typed = set(ENUMS) | BOOLEANS | LENGTHS | INTEGERS | {"font.family", "font.size", "paragraph.line_spacing", "page.dimensions"}
    assert {item["property_id"] for item in capabilities} == typed
    assert len(typed) == 35
    for name in ("academic-edit-manifest", "academic-edit-result"):
        checked = (ROOT / "contracts" / name / f"{name}-1.0.schema.json").read_bytes()
        bundled = (ROOT / "apps/api/src/autoformat/schemas" / f"{name}-1.0.schema.json").read_bytes()
        assert checked == bundled
        Draft202012Validator.check_schema(json.loads(checked))


def test_parser_to_editor_and_report_preserve_content(editor, tmp_path):
    store, _, result, engine = task(tmp_path, "正文：宋体，12pt，两端对齐，首行缩进2字符。\n一级标题：黑体，15pt，加粗，居中。\n页面：A4，上页边距2.5cm。\n表格使用三线表。")
    snapshot = engine.compile()
    assert {r.property_id for r in snapshot.rules} >= {"font.family", "font.size", "font.bold", "paragraph.alignment", "paragraph.first_line_indent", "table.border_style"}
    original = (tmp_path / "paper.docx").read_bytes()
    work, output = execute(editor, tmp_path, result, engine)
    assert output["status"] == "ready", output
    assert all(r["status"] in {"modified_verified", "already_satisfied"} for r in output["rules"]), output["rules"]
    before, after = Document(tmp_path / "paper.docx"), Document(work / "edited.docx")
    assert [p.text for p in before.paragraphs] == [p.text for p in after.paragraphs]
    assert [[c.text for c in t._cells] for t in before.tables] == [[c.text for c in t._cells] for t in after.tables]
    assert (tmp_path / "paper.docx").read_bytes() == original
    assert store.artifact_path(result.version.task_id, result.version.version_id, "docx").read_bytes() == original
    rules, changes, decisions = report_facts(result.bundle, snapshot, output)
    report = build_report(ReportContext(task_id=result.version.task_id, version_id=new_id(), parent_version_id=result.version.version_id,
                          source_docx_sha256=output["base_docx_sha256"], output_docx_sha256=output["output_docx_sha256"],
                          brief_version_id="brief:v1", rules_snapshot_id=snapshot.snapshot_id, edit_plan_id="test-plan",
                          preview_artifact_id="test-preview", renderer_id="test", renderer_version="1", font_fingerprint="test"),
                          rules=rules, changes_this_round=changes, changes_cumulative=changes, decisions=decisions, issues=[], note="verified integration")
    assert report["summary"]["format_status"] == "verified"
    assert any(r["after_value"]["operations"] for r in rules)


def test_partial_conflict_keeps_clear_fragments(editor, tmp_path):
    _, _, result, engine = task(tmp_path, "正文：宋体。\n第2章正文：黑体。")
    snapshot = engine.compile()
    assert snapshot.conflicts and any(f.status == "verified" for f in snapshot.fragments)
    work, output = execute(editor, tmp_path, result, engine)
    assert output["status"] == "ready", output
    assert output["operations"] and all(r["status"] == "pending_conflict" for r in output["rules"])
    doc = Document(work / "edited.docx")
    assert any(op["character_type"] == "chinese" and op["status"] == "already_satisfied" for op in output["operations"])
    assert doc.paragraphs[1].runs[1]._element.rPr.rFonts.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}ascii") == "宋体"
    assert doc.paragraphs[3].runs[0]._element.rPr is None
    facts, changes, _ = report_facts(result.bundle, snapshot, output)
    assert any(r["status"] == "modified_verified" for r in facts) and any(r["status"] == "pending_conflict" for r in facts)
    assert changes


def test_decided_value_is_replayed_before_editing(editor, tmp_path):
    _, _, result, engine = task(tmp_path, "正文：字号12-14pt。")
    baseline = engine.compile()
    question = next(q for q in baseline.questions if q.kind == "value")
    request = AcademicDecisionRequest(request_id=new_id(), base_version_id=engine.initial_version_id,
                                      expected_snapshot_id=baseline.snapshot_id, question_id=question.question_id,
                                      answer={"action": "clarify_format", "format": {"property_id": "font.size", "normalized": 13, "unit": "pt"}})
    snapshot = engine.submit(baseline, request, decision_id=new_id(), created_at="2026-10-04T00:00:00+00:00")
    work, output = execute(editor, tmp_path, result, engine, snapshot)
    assert output["status"] == "ready" and all(r["status"] == "modified_verified" for r in output["rules"])
    assert Document(work / "edited.docx").paragraphs[1].runs[0].font.size == Pt(13)
    forged_rule = snapshot.rules[0].model_copy(update={"value": snapshot.rules[0].value.model_copy(update={"normalized": 14})})
    forged = snapshot.model_copy(update={"rules": [forged_rule]})
    forged = forged.model_copy(update={"snapshot_id": content_hash(forged.model_dump(mode="json", exclude={"snapshot_id"}))})
    with pytest.raises(WorkflowError, match="source replay"):
        execute(editor, tmp_path, result, engine, forged)


def test_plan_is_read_only_and_keeps_scope(editor, tmp_path):
    _, _, result, engine = task(tmp_path, "第2章正文：右对齐。")
    work, output = execute(editor, tmp_path, result, engine, plan_only=True)
    assert output["status"] == "planned" and len(output["operations"]) == 1
    assert not (work / "edited.docx").exists()
    assert (work / "parent.docx").read_bytes() == (tmp_path / "paper.docx").read_bytes()


def test_first_v1_waits_for_pre_edit_confirmation(editor, tmp_path):
    # Use the existing deterministic fixture renderer; this checks publication, not visual fidelity.
    from tests.academic_intake.test_academic_intake import FixtureRenderer
    from autoformat.academic_decisions import AcademicDecisionStore
    store, intake, result, _ = task(tmp_path, "正文：12pt，居中。\n正文：12-14pt。")
    publisher = InitialDraftPublisher(store=store, renderer_factory=FixtureRenderer, font_fingerprint="test-fonts", academic_editor=editor)
    from autoformat.version_store import VersionStoreError
    with pytest.raises(VersionStoreError, match="先选择"):
        publisher.publish(result.version, result.bundle)
    assert len(store.list_versions(result.version.task_id)) == 1
    decisions = AcademicDecisionStore(store, intake)
    for _ in range(256):
        snapshot = decisions.load(result.version.task_id)
        question = next((q for q in snapshot.questions if q.status in {"pending", "deferred"}), None)
        if question is None:
            break
        decisions.submit(result.version.task_id, AcademicDecisionRequest(
            request_id=new_id(), base_version_id=result.version.version_id, expected_snapshot_id=snapshot.snapshot_id,
            question_id=question.question_id, answer={"action": "skip"}))
    published = publisher.publish(result.version, result.bundle)
    assert published.sequence == 1
    assert publisher.publish(result.version, result.bundle).version_id == published.version_id
    doc = Document(store.artifact_path(published.task_id, published.version_id, "docx"))
    assert doc.paragraphs[1].alignment == 1
    assert result.version.docx_sha256 != published.docx_sha256
    report = json.loads(store.artifact_path(published.task_id, published.version_id, "report_json").read_text(encoding="utf-8"))
    assert report["summary"]["this_round_change_count"] > 0
    snapshot = AcademicDecisionStore(store, intake).load(published.task_id)
    assert report["input"]["rules_snapshot_id"] == snapshot.snapshot_id
    assert snapshot.questions
    assert len(store.list_versions(published.task_id)) == 2


def test_pre_edit_conflict_choice_is_used_by_real_editor_in_v1(editor, tmp_path):
    from tests.academic_intake.test_academic_intake import FixtureRenderer
    from autoformat.academic_decisions import AcademicDecisionStore
    store, intake, result, _ = task(tmp_path, "正文：左对齐。\n正文：居中。")
    decisions = AcademicDecisionStore(store, intake)
    snapshot = decisions.load(result.version.task_id)
    question = next(q for q in snapshot.questions if q.kind == "conflict")
    rules = {rule.rule_id: rule for rule in snapshot.rules}
    chosen = next(identifier for identifier in question.rule_ids if rules[identifier].value.normalized == "center")
    decisions.submit(result.version.task_id, AcademicDecisionRequest(request_id=new_id(),
        base_version_id=result.version.version_id, expected_snapshot_id=snapshot.snapshot_id,
        question_id=question.question_id, answer={"action": "select_rule", "selected_id": chosen}))
    publisher = InitialDraftPublisher(store=store, renderer_factory=FixtureRenderer,
        font_fingerprint="test-fonts", academic_editor=editor, decision_store=decisions)
    published = publisher.publish(result.version, result.bundle)
    document = Document(store.artifact_path(published.task_id, published.version_id, "docx"))
    assert document.paragraphs[1].alignment == 1
    report = json.loads(store.artifact_path(published.task_id, published.version_id, "report_json").read_text(encoding="utf-8"))
    assert report["decisions"][0]["normalized_choice"] == "select_rule"
    assert not any("明文规则冲突" in issue["reason"] for issue in report["issues"])
