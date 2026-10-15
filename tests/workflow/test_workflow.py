"""Real C# editor to LibreOffice to immutable version publication."""

from __future__ import annotations

import hashlib
import json
import os
import shutil
import subprocess
from dataclasses import replace
from pathlib import Path

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pypdf import PdfReader
from brief_parser.domain.academic_format import (
    AcademicFormatIssue, AcademicFormatRule, AcademicFormatRuleSet,
    AcademicTarget, AcademicValue,
)
from brief_parser.domain.requirement_ir import EvidenceReference

from autoformat.preview import LibreOfficeRenderer
from autoformat.target_planning import (
    DotnetTargetIndexer, TargetPlanningError, build_confirmed_body_alignment_plan,
)
from autoformat.version_store import VersionReferences, VersionStore, new_id
from autoformat.workflow import (
    ConfirmedBodyAlignment, ConfirmedBriefAlignment, DotnetEditor, WorkflowError,
    apply_confirmed_alignment, apply_confirmed_body_alignment_plan,
    confirm_body_scope,
)


ROOT = Path(__file__).resolve().parents[2]
EDITOR_PROJECT = ROOT / "services" / "docx_engine" / "AutoFormat.DocxEditor"
INDEX_PROJECT = ROOT / "services" / "docx_engine" / "AutoFormat.DocxIndex"


@pytest.fixture(scope="module")
def editor() -> DotnetEditor:
    dotnet = shutil.which("dotnet")
    if dotnet is None:
        pytest.skip(".NET 10 SDK is required for the editor integration test.")
    completed = subprocess.run(
        [dotnet, "build", str(EDITOR_PROJECT), "--no-restore", "-c", "Debug"],
        cwd=ROOT, capture_output=True, text=True, timeout=120, check=False,
        encoding="utf-8", errors="replace",
    )
    assert completed.returncode == 0, completed.stderr[-2000:]
    dll = EDITOR_PROJECT / "bin" / "Debug" / "net10.0" / "AutoFormat.DocxEditor.dll"
    return DotnetEditor(
        dotnet_path=Path(dotnet), editor_dll=dll,
        expected_editor_sha256=hashlib.sha256(dll.read_bytes()).hexdigest(),
    )


@pytest.fixture(scope="module")
def indexer() -> DotnetTargetIndexer:
    dotnet = shutil.which("dotnet")
    if dotnet is None:
        pytest.skip(".NET 10 SDK is required for target indexing.")
    completed = subprocess.run(
        [dotnet, "build", str(INDEX_PROJECT), "--no-restore", "-c", "Debug"],
        cwd=ROOT, capture_output=True, text=True, timeout=120, check=False,
        encoding="utf-8", errors="replace",
    )
    assert completed.returncode == 0, completed.stderr[-2000:]
    dll = INDEX_PROJECT / "bin" / "Debug" / "net10.0" / "AutoFormat.DocxIndex.dll"
    return DotnetTargetIndexer(
        dotnet_path=Path(dotnet), index_dll=dll,
        expected_index_sha256=hashlib.sha256(dll.read_bytes()).hexdigest(),
    )


def _body_ruleset(*, include_issue: bool = False) -> AcademicFormatRuleSet:
    rule = AcademicFormatRule.model_construct(
        rule_id="brief-body-center", source_requirement_id="req-body",
        source_kind="brief_explicit", source_artifact_id="fixture-brief",
        source_version="brief-v1", evidence_refs=[EvidenceReference.model_construct(
            block_id="fixture:block:1", quote="正文居中", match_status="exact")],
        target=AcademicTarget(
            object_kind="paragraph", semantic_role="body", selector="semantic_role"),
        property_id="paragraph.alignment",
        value=AcademicValue(
            original_text="居中", normalized="center",
            dictionary_version="academic-alignment/1.0"),
        character_scope="all", condition=None, requested_scope="same_kind",
        effective_scope="same_kind", status="verified", reason="Fixture rule.",
    )
    issues = ([AcademicFormatIssue(
        source_requirement_id="req-unclassified",
        status="needs_source_classification", reason="Source type is unknown.",
        next_action="Classify the source.", evidence_ids=[],
    )] if include_issue else [])
    return AcademicFormatRuleSet.model_construct(rules=[rule], issues=issues)


def test_confirmed_body_scope_batch_publishes_one_version(
    tmp_path: Path, editor: DotnetEditor, indexer: DotnetTargetIndexer,
) -> None:
    executable = os.environ.get("AUTOFORMAT_TEST_SOFFICE")
    expected_hash = os.environ.get("AUTOFORMAT_TEST_SOFFICE_SHA256")
    if not executable or not expected_hash:
        pytest.skip("Set verified LibreOffice path and SHA-256 for integration.")
    source = tmp_path / "original.docx"
    _document(source)
    store = VersionStore(tmp_path / "store")
    v0 = store.create_task(source)
    ruleset = _body_ruleset(include_issue=True)
    confirmation, plan = confirm_body_scope(
        store=store, parent=v0, request_id=new_id(),
        raw_answer="前两段是正文，第三段不是正文。",
        ruleset=ruleset, rules_snapshot_id="rules-body",
        indexer=indexer,
        body_object_ids=("p:000001", "p:000002"),
        excluded_object_ids=("p:000003",),
    )
    assert plan["estimated_affected_count"] == 2
    references = VersionReferences(
        rules_snapshot_id="rules-body", brief_version_id="brief-v1",
        edit_plan_id=plan["plan_id"], decision_version_id=confirmation.decision_id,
        font_fingerprint="fixture-fonts",
    )
    source_rule = ConfirmedBodyAlignment(
        rule_id="brief-body-center", source_artifact_id="fixture-brief",
        source_version="brief-v1", original_requirement="正文居中。",
        evidence_locator="fixture:block:1", evidence_quote="正文居中",
        value="center",
    )

    def renderer(work_dir: Path) -> LibreOfficeRenderer:
        return LibreOfficeRenderer(
            soffice_path=Path(executable), expected_soffice_sha256=expected_hash,
            data_root=work_dir, timeout_seconds=60,
        )

    request_id = new_id()
    result = apply_confirmed_body_alignment_plan(
        store=store, parent=v0, request_id=request_id, source_rule=source_rule,
        ruleset=ruleset, confirmation=confirmation, plan=plan,
        references=references, indexer=indexer, editor=editor,
        renderer_factory=renderer,
    )
    assert result.status == "ready" and result.version is not None
    v1 = result.version
    assert len(store.list_versions(v0.task_id)) == 2
    output = Document(store.artifact_path(v1.task_id, v1.version_id, "docx"))
    assert output.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert output.paragraphs[1].alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert output.paragraphs[2].alignment is None
    report = json.loads(store.artifact_path(
        v1.task_id, v1.version_id, "report_json").read_text(encoding="utf-8"))
    assert report["summary"]["this_round_change_count"] == 2
    assert {item["object_id"] for item in report["changes_this_round"]} == {
        "p:000001", "p:000002"}
    assert report["decisions"][0]["decision_id"] == confirmation.decision_id
    assert report["input"]["edit_plan_id"] == plan["plan_id"]
    assert report["summary"]["pending_count"] == 1
    assert report["issues"][0]["status"] == "needs_user_input"
    assert not list((store.root / ".operations").iterdir())

    replay = apply_confirmed_body_alignment_plan(
        store=store, parent=v0, request_id=request_id, source_rule=source_rule,
        ruleset=ruleset, confirmation=confirmation, plan=plan,
        references=references, indexer=indexer, editor=editor,
        renderer_factory=renderer,
    )
    assert replay.version == v1


def test_body_scope_requires_complete_classification_before_recording(
    tmp_path: Path, indexer: DotnetTargetIndexer,
) -> None:
    source = tmp_path / "original.docx"
    _document(source)
    store = VersionStore(tmp_path / "store")
    v0 = store.create_task(source)
    with pytest.raises(TargetPlanningError) as incomplete:
        confirm_body_scope(
            store=store, parent=v0, request_id=new_id(),
            raw_answer="只确认第一段。", ruleset=_body_ruleset(),
            rules_snapshot_id="rules-body", indexer=indexer,
            body_object_ids=("p:000001",),
            excluded_object_ids=("p:000002",),
        )
    assert incomplete.value.status == "needs_user_input"
    assert store.list_decisions(v0.task_id) == []


def test_body_plan_requires_recorded_decision_and_unchanged_plan(
    tmp_path: Path, editor: DotnetEditor, indexer: DotnetTargetIndexer,
) -> None:
    source = tmp_path / "original.docx"
    _document(source)
    store = VersionStore(tmp_path / "store")
    v0 = store.create_task(source)
    ruleset = _body_ruleset()
    confirmation, plan = confirm_body_scope(
        store=store, parent=v0, request_id=new_id(),
        raw_answer="第一段为正文。", ruleset=ruleset,
        rules_snapshot_id="rules-body", indexer=indexer,
        body_object_ids=("p:000001",),
        excluded_object_ids=("p:000002", "p:000003"),
    )
    source_rule = ConfirmedBodyAlignment(
        rule_id="brief-body-center", source_artifact_id="fixture-brief",
        source_version="brief-v1", original_requirement="正文居中。",
        evidence_locator="fixture:block:1", evidence_quote="正文居中", value="center",
    )

    def unexpected_renderer(_: Path) -> LibreOfficeRenderer:
        raise AssertionError("Rejected plan must not render a preview.")

    def references(decision_id: str, plan_id: str) -> VersionReferences:
        return VersionReferences(
            rules_snapshot_id="rules-body", brief_version_id="brief-v1",
            edit_plan_id=plan_id, decision_version_id=decision_id,
            font_fingerprint="fixture-fonts",
        )

    forged = replace(confirmation, decision_id="not-recorded")
    forged_plan = build_confirmed_body_alignment_plan(
        ruleset=ruleset,
        index=indexer.inspect(store.artifact_path(v0.task_id, v0.version_id, "docx"),
                              expected_docx_sha256=v0.docx_sha256),
        confirmation=forged,
    )
    with pytest.raises(WorkflowError) as missing:
        apply_confirmed_body_alignment_plan(
            store=store, parent=v0, request_id=new_id(),
            source_rule=source_rule, ruleset=ruleset,
            confirmation=forged,
            plan=forged_plan,
            references=references("not-recorded", forged_plan["plan_id"]), indexer=indexer,
            editor=editor, renderer_factory=unexpected_renderer,
        )
    assert missing.value.status == "needs_user_input"

    altered = {**plan, "operations": [{**plan["operations"][0], "after_value": "right"}]}
    with pytest.raises(WorkflowError) as invalid:
        apply_confirmed_body_alignment_plan(
            store=store, parent=v0, request_id=new_id(),
            source_rule=source_rule, ruleset=ruleset,
            confirmation=confirmation, plan=altered,
            references=references(confirmation.decision_id, plan["plan_id"]), indexer=indexer,
            editor=editor, renderer_factory=unexpected_renderer,
        )
    assert invalid.value.status == "validation_failed"
    with pytest.raises(WorkflowError) as evidence:
        apply_confirmed_body_alignment_plan(
            store=store, parent=v0, request_id=new_id(),
            source_rule=replace(source_rule, evidence_quote="invented"),
            ruleset=ruleset, confirmation=confirmation, plan=plan,
            references=references(confirmation.decision_id, plan["plan_id"]),
            indexer=indexer, editor=editor, renderer_factory=unexpected_renderer,
        )
    assert evidence.value.status == "validation_failed"
    assert len(store.list_versions(v0.task_id)) == 1


def _document(path: Path) -> None:
    document = Document()
    document.add_paragraph("First body paragraph remains intact.")
    document.add_paragraph("Second body paragraph remains intact.")
    document.add_paragraph("Third body paragraph remains intact.")
    document.save(path)


def _rule(number: int, value: str) -> ConfirmedBriefAlignment:
    text = ("First body paragraph remains intact." if number == 1 else
            "Second body paragraph remains intact.")
    return ConfirmedBriefAlignment(
        rule_id=f"brief-alignment-{number}-{value}",
        source_artifact_id="fixture-brief", source_version="brief-v1",
        original_requirement=f"Fixture requires paragraph {number} to be {value}.",
        evidence_locator=f"fixture:block:{number}",
        evidence_quote=f"paragraph {number}: {value}",
        target_object_id=f"p:{number:06d}",
        expected_paragraph_text_sha256=hashlib.sha256(text.encode("utf-8")).hexdigest(),
        value=value, semantic_role="body",
    )


def _references(plan: str) -> VersionReferences:
    return VersionReferences(
        rules_snapshot_id=f"rules-{plan}", brief_version_id="brief-v1",
        edit_plan_id=f"plan-{plan}", font_fingerprint="fixture-fonts",
    )


def test_editor_renderer_report_and_version_lineage(tmp_path: Path, editor: DotnetEditor) -> None:
    executable = os.environ.get("AUTOFORMAT_TEST_SOFFICE")
    expected_hash = os.environ.get("AUTOFORMAT_TEST_SOFFICE_SHA256")
    if not executable or not expected_hash:
        pytest.skip("Set the verified LibreOffice path and SHA-256 for integration.")
    source = tmp_path / "original.docx"
    _document(source)
    original_hash = hashlib.sha256(source.read_bytes()).hexdigest()
    store = VersionStore(tmp_path / "store")
    v0 = store.create_task(source)

    def renderer(work_dir: Path) -> LibreOfficeRenderer:
        return LibreOfficeRenderer(
            soffice_path=Path(executable), expected_soffice_sha256=expected_hash,
            data_root=work_dir, timeout_seconds=60,
        )

    first_request = new_id()
    first = apply_confirmed_alignment(
        store=store, parent=v0, expected_parent_sha256=v0.docx_sha256,
        request_id=first_request, rule=_rule(1, "center"),
        references=_references("one"), editor=editor, renderer_factory=renderer,
    )
    assert first.status == "ready" and first.version is not None
    v1 = first.version
    second = apply_confirmed_alignment(
        store=store, parent=v1, expected_parent_sha256=v1.docx_sha256,
        request_id=new_id(), rule=_rule(2, "right"),
        references=_references("two"), editor=editor, renderer_factory=renderer,
    )
    assert second.status == "ready" and second.version is not None
    v2 = second.version
    assert v1.parent_version_id == v0.version_id
    assert v2.parent_version_id == v1.version_id
    assert [v.sequence for v in store.list_versions(v0.task_id)] == [0, 1, 2]
    assert hashlib.sha256(store.artifact_path(v0.task_id, v0.version_id, "docx").read_bytes()).hexdigest() == original_hash
    assert hashlib.sha256(source.read_bytes()).hexdigest() == original_hash

    output = Document(store.artifact_path(v2.task_id, v2.version_id, "docx"))
    assert [p.text for p in output.paragraphs] == [p.text for p in Document(source).paragraphs]
    assert output.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert output.paragraphs[1].alignment == WD_ALIGN_PARAGRAPH.RIGHT
    assert output.paragraphs[2].alignment is None
    report = json.loads(store.artifact_path(v2.task_id, v2.version_id, "report_json").read_text(encoding="utf-8"))
    assert report["summary"]["this_round_change_count"] == 1
    assert report["summary"]["cumulative_change_count"] == 2
    assert [change["after_value"] for change in report["changes_cumulative"]] == ["center", "right"]
    assert all(change["verification_ref"].startswith("docx:") for change in report["changes_cumulative"])
    preview = store.artifact_path(v2.task_id, v2.version_id, "preview")
    assert "Second body paragraph" in (PdfReader(preview).pages[0].extract_text() or "")
    assert not list((store.root / ".operations").iterdir())

    def unexpected_renderer(_: Path) -> LibreOfficeRenderer:
        raise AssertionError("Already satisfied edit must not render a preview.")

    replay = apply_confirmed_alignment(
        store=store, parent=v0, expected_parent_sha256=v0.docx_sha256,
        request_id=first_request, rule=_rule(1, "center"),
        references=_references("one"), editor=editor,
        renderer_factory=unexpected_renderer,
    )
    assert replay.version == v1
    with pytest.raises(WorkflowError) as conflicting_request:
        apply_confirmed_alignment(
            store=store, parent=v0, expected_parent_sha256=v0.docx_sha256,
            request_id=first_request, rule=_rule(2, "right"),
            references=_references("one"), editor=editor,
            renderer_factory=unexpected_renderer,
        )
    assert conflicting_request.value.status == "blocked_conflict"

    no_change = apply_confirmed_alignment(
        store=store, parent=v2, expected_parent_sha256=v2.docx_sha256,
        request_id=new_id(),
        rule=replace(_rule(1, "center"), rule_id="another-confirmed-rule"),
        references=_references("three"), editor=editor,
        renderer_factory=unexpected_renderer,
    )
    assert no_change.status == "already_satisfied" and no_change.version is None
    assert len(store.list_versions(v0.task_id)) == 3


def test_stale_parent_and_unconfirmed_rule_do_not_run_editor(tmp_path: Path, editor: DotnetEditor) -> None:
    source = tmp_path / "source.docx"
    _document(source)
    store = VersionStore(tmp_path / "store")
    v0 = store.create_task(source)

    def unexpected_renderer(_: Path) -> LibreOfficeRenderer:
        raise AssertionError("Rejected edit must not render a preview.")

    with pytest.raises(WorkflowError) as stale:
        apply_confirmed_alignment(
            store=store, parent=v0, expected_parent_sha256="0" * 64,
            request_id=new_id(), rule=_rule(1, "center"),
            references=_references("stale"), editor=editor,
            renderer_factory=unexpected_renderer,
        )
    assert stale.value.status == "stale_base_version"
    with pytest.raises(WorkflowError) as invalid:
        apply_confirmed_alignment(
            store=store, parent=v0, expected_parent_sha256=v0.docx_sha256,
            request_id=new_id(), rule=ConfirmedBriefAlignment(
                rule_id="bad", source_artifact_id="brief", source_version="brief-v1",
                original_requirement="unsafe", evidence_locator="block:1",
                evidence_quote="unsafe", target_object_id="p:000001",
                expected_paragraph_text_sha256="0" * 64,
                value="made-up", semantic_role="body",
            ),
            references=_references("invalid"), editor=editor,
            renderer_factory=unexpected_renderer,
        )
    assert invalid.value.status == "needs_user_input"
    with pytest.raises(WorkflowError) as ambiguous:
        apply_confirmed_alignment(
            store=store, parent=v0, expected_parent_sha256=v0.docx_sha256,
            request_id=new_id(),
            rule=replace(_rule(1, "center"), expected_paragraph_text_sha256="0" * 64),
            references=_references("ambiguous"), editor=editor,
            renderer_factory=unexpected_renderer,
        )
    assert ambiguous.value.status == "ambiguous_target"
    assert len(store.list_versions(v0.task_id)) == 1
    assert not list((store.root / ".operations").iterdir())


def test_runtime_result_schema_matches_contract() -> None:
    assert (ROOT / "contracts" / "edit-result" / "edit-result-prototype-1.0.schema.json").read_bytes() == (
        ROOT / "apps" / "api" / "src" / "autoformat" / "schemas" /
        "edit-result-prototype-1.0.schema.json"
    ).read_bytes()
    assert (ROOT / "contracts" / "edit-batch-result" / "edit-batch-result-1.0.schema.json").read_bytes() == (
        ROOT / "apps" / "api" / "src" / "autoformat" / "schemas" /
        "edit-batch-result-1.0.schema.json"
    ).read_bytes()
    assert (ROOT / "contracts" / "edit-batch-manifest" / "edit-batch-manifest-1.0.schema.json").read_bytes() == (
        ROOT / "apps" / "api" / "src" / "autoformat" / "schemas" /
        "edit-batch-manifest-1.0.schema.json"
    ).read_bytes()
