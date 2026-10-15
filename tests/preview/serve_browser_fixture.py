"""Browser fixture: real SQLite + DOCX + report + selectable two-page PDF.

Set AUTOFORMAT_TEST_SOFFICE and AUTOFORMAT_TEST_SOFFICE_SHA256 to exercise the
real pinned LibreOffice renderer. Without them, a minimal text PDF is generated.
All outputs are under runtime/; this is not a product/demo task generator.
"""
from __future__ import annotations

import json
import os
import shutil
import sys
from dataclasses import replace
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "apps/api/src"))
sys.path.insert(0, str(ROOT / "modules/brief_parsing/src"))

from docx import Document
from docx.shared import Pt
from pypdf import PdfReader
import uvicorn

from autoformat.local_api import create_local_api
from autoformat.academic_intake import AcademicTaskIntake
from brief_parser.bootstrap import build_service
from brief_parser.config import Settings
from autoformat.preview import LibreOfficeRenderer, PreviewArtifact, _sha256
from autoformat.workflow import DotnetEditor
from autoformat.reports import ReportContext, build_report, write_report
from autoformat.rule_snapshots import RuleSnapshotStore
from autoformat.selection_indexes import SelectionIndexes
from autoformat.version_store import VersionStore, VersionReferences, new_id
from autoformat.feedback_rules_contracts import FeedbackRuleBasis
from autoformat.feedback_rule_basis import source_loader
from brief_parser.domain.academic_decisions import AcademicEffectiveRule
from brief_parser.compilers.academic_scope import analyze_rules, base_scope
from hashlib import sha256

TOKEN = "preview-browser-test-session-token-1234567890"
PORT = 58371


def minimal_pdf(path: Path):
    objects = [b"<< /Type /Catalog /Pages 2 0 R >>", b"<< /Type /Pages /Kids [3 0 R 4 0 R] /Count 2 >>"]
    for number in [6, 7]:
        objects.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Resources << /Font << /F1 5 0 R >> >> /Contents {number} 0 R >>".encode())
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    for page in [1, 2]:
        text = (b"BT /F1 16 Tf 60 770 Td (" + (b"Selectable preview text" if page == 1 else b"Second page content") +
                b") Tj 0 -35 Td (Another line of text) Tj ET")
        objects.append(b"<< /Length " + str(len(text)).encode() + b" >>\nstream\n" + text + b"\nendstream")
    data = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for i, item in enumerate(objects, 1):
        offsets.append(len(data)); data.extend(f"{i} 0 obj\n".encode() + item + b"\nendobj\n")
    xref = len(data)
    data.extend(f"xref\n0 {len(offsets)}\n0000000000 65535 f \n".encode())
    for offset in offsets[1:]:
        data.extend(f"{offset:010d} 00000 n \n".encode())
    data.extend(f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    path.write_bytes(data)


def publish(store, parent, work, *, synthetic=False, preview_template=None, pending=False):
    work.mkdir()
    candidate = work / "document.docx"
    source = store.artifact_path(parent.task_id, parent.version_id, "docx")
    rules, changes = [], []
    if parent.kind == "original":
        shutil.copyfile(source, candidate)
    else:
        document = Document(source)
        before = document.paragraphs[0].runs[0].bold
        document.paragraphs[0].runs[0].bold = True
        document.save(candidate)
        assert Document(candidate).paragraphs[0].runs[0].bold is True
        rule = {"rule_id": "browser-fixture-bold", "source_kind": "brief_explicit", "source_artifact_id": "test-brief", "source_version": "brief:fixture",
            "original_requirement": "测试首行加粗", "evidence_refs": [{"source_artifact_id": "test-brief", "source_version": "brief:fixture", "locator": "fixture:1", "quote": "测试首行加粗"}],
            "object_id": "p:000001", "semantic_role": "body", "property_id": "font.bold", "requested_scope": "paragraph", "effective_scope": "paragraph",
            "before_value": before, "after_value": True, "status": "modified_verified", "reason": "Fixture reopened and verified.", "next_action": None, "decision_id": None, "verification_ref": "fixture-verify-bold"}
        rules = [rule]
        changes = [{"change_id": "fixture-change-bold", "rule_id": rule["rule_id"], "object_id": rule["object_id"], "property_id": rule["property_id"],
            "before_value": before, "after_value": True, "effective_scope": "paragraph", "validation_status": "verified", "verification_ref": rule["verification_ref"]}]
    if pending:
        rules.append({"rule_id": "fixture-pending-font", "source_kind": "brief_explicit", "source_artifact_id": "test-brief", "source_version": "brief:fixture",
            "original_requirement": "正文使用较大字号（测试待确认原文）", "evidence_refs": [{"source_artifact_id": "test-brief", "source_version": "brief:fixture", "locator": "fixture:2", "quote": "正文使用较大字号（测试待确认原文）"}],
            "object_id": None, "semantic_role": "body", "property_id": "font.size", "requested_scope": "same_kind", "effective_scope": "same_kind",
            "before_value": None, "after_value": None, "status": "needs_user_input", "reason": "字号数值尚未明确。", "next_action": "请核对学校要求的具体字号。", "decision_id": None, "verification_ref": None})
    version_id, artifact_id = new_id(), new_id()
    if preview_template is not None:
        template_path, template = preview_template
        assert _sha256(candidate) == template.source_sha256
        relative = f"artifacts/previews/{artifact_id}.pdf"
        pdf = work / relative
        pdf.parent.mkdir(parents=True)
        shutil.copyfile(template_path, pdf)
        artifact = replace(template, artifact_id=artifact_id, pdf_relative_path=relative)
    elif os.environ.get("AUTOFORMAT_TEST_SOFFICE") and not synthetic:
        artifact = LibreOfficeRenderer(soffice_path=Path(os.environ["AUTOFORMAT_TEST_SOFFICE"]),
            expected_soffice_sha256=os.environ["AUTOFORMAT_TEST_SOFFICE_SHA256"], data_root=work).render(candidate, artifact_id)
        pdf = work / artifact.pdf_relative_path
    else:
        pdf = work / "preview.pdf"; minimal_pdf(pdf)
        artifact = PreviewArtifact("preview-artifact/1.0", artifact_id, _sha256(candidate), _sha256(pdf), f"artifacts/previews/{artifact_id}.pdf", 2,
            sum(len(page.extract_text()) for page in PdfReader(pdf).pages), "test-fixture", "1.0")
    report = build_report(ReportContext(task_id=parent.task_id, version_id=version_id, parent_version_id=parent.version_id,
        source_docx_sha256=parent.docx_sha256, output_docx_sha256=_sha256(candidate), brief_version_id="brief:fixture",
        rules_snapshot_id="rules:fixture", edit_plan_id="plan:fixture", preview_artifact_id=artifact_id,
        renderer_id=artifact.renderer_id, renderer_version=artifact.renderer_version, font_fingerprint="test-fonts"),
        rules=rules, changes_this_round=changes, changes_cumulative=changes, decisions=[], issues=[],
        note="浏览器测试样本。真实格式变化已在保存后重新读取校验。" if changes else "本测试初稿未作格式修改；请核对报告中的完成状态。")
    report_path, markdown = work / "report.json", work / "report.md"
    write_report(report_path, markdown, report)
    return store.publish_version(task_id=parent.task_id, version_id=version_id, parent_version_id=parent.version_id,
        expected_parent_sha256=parent.docx_sha256, request_id=new_id(), candidate_docx=candidate,
        report_json=report_path, report_markdown=markdown, preview_pdf=pdf, preview_artifact=artifact,
        references=VersionReferences("rules:fixture", "brief:fixture", "plan:fixture", "test-fonts"))


def main():
    work = ROOT / "runtime/browser-preview" / new_id()
    work.mkdir(parents=True)
    document = Document()
    document.styles["Normal"].font.name = "Arial"
    document.styles["Normal"].font.size = Pt(12)
    document.add_paragraph("Selectable preview text", style="Body Text")
    document.add_paragraph("材料学院毕业论文格式预览", style="Body Text")
    document.add_paragraph("Another line of text", style="Body Text")
    document.add_paragraph("选中文字后，可以记录具体的格式反馈。", style="Body Text")
    document.add_page_break()
    document.add_paragraph("Second page content", style="Body Text")
    document.add_paragraph("第二页用于验证翻页、缩放和选区清理。", style="Body Text")
    source = work / "original.docx"; document.save(source)
    store = VersionStore(work / "store")
    v0 = store.create_task(source)
    v1 = publish(store, v0, work / "v1")
    v2 = publish(store, v1, work / "v2")
    duplicate = Document()
    for _ in range(2):
        duplicate.add_paragraph("Selectable preview text")
    duplicate_source = work / "duplicate.docx"; duplicate.save(duplicate_source)
    duplicate_v0 = store.create_task(duplicate_source)
    duplicate_v1 = publish(store, duplicate_v0, work / "duplicate-v1", synthetic=True)
    rules_tasks = {}
    preview_template = (store.artifact_path(v0.task_id, v1.version_id, "preview"),
        PreviewArtifact(**json.loads(store.artifact_path(v0.task_id, v1.version_id, "preview_metadata").read_text(encoding="utf-8"))))
    for number, name in enumerate(["local", "global", "retry", "execution_auto", "execution_exception", "execution_retry", "execution_paragraph",
            "workspace_history", "workspace_retry", "workspace_queue", "workspace_branch", "workspace_report"]):
        original = store.create_task(source)
        output = publish(store, original, work / f"r{number}", preview_template=preview_template, pending=name == "workspace_report")
        rules_tasks[name] = {"task": original.task_id, "version": output.version_id}
        if name in {"workspace_history", "workspace_retry"}:
            child = publish(store, output, work / f"r{number}-child", synthetic=True)
            rules_tasks[name]["child"] = child.version_id
    manifest = {"task": v0.task_id, "v0": v0.version_id, "v1": v1.version_id, "v2": v2.version_id,
        "store": str(store.root), "pdf_sha256": _sha256(store.artifact_path(v0.task_id, v2.version_id, "preview")),
        "duplicate_task": duplicate_v0.task_id, "duplicate_v1": duplicate_v1.version_id, "paper": str(source), "rules_tasks": rules_tasks}
    (ROOT / "runtime/browser-preview/manifest.json").write_text(json.dumps(manifest), encoding="utf-8")
    # Freeze the built UI for this run. Another build in the shared workspace
    # can replace hashed assets while a test is navigating between pages.
    web_dist = work / "web-dist"
    shutil.copytree(ROOT / "apps/web/dist", web_dist)
    class ReviewRenderer:
        def __init__(self, root):
            self.root = root

        def render(self, docx_path, artifact_id):
            relative = f"artifacts/previews/{artifact_id}.pdf"
            pdf = self.root / relative
            pdf.parent.mkdir(parents=True)
            minimal_pdf(pdf)
            return PreviewArtifact("preview-artifact/1.0", artifact_id, _sha256(docx_path), _sha256(pdf), relative, 2,
                sum(len(page.extract_text() or "") for page in PdfReader(pdf).pages), "test-fixture", "1.0")

    intake = AcademicTaskIntake(store=store, service=build_service(Settings(
        environment="test", provider="fake", parser_isolation_mode="in_process")))
    snapshots = RuleSnapshotStore(store)
    real_basis = source_loader(store, snapshots, intake)
    quote = "正文采用12pt字号"
    rule = AcademicEffectiveRule.model_validate(dict(rule_id="fixture-body-size", source_kind="brief_explicit", source_rule_ids=["fixture-body-size"],
        evidence_refs=[dict(evidence_id="fixture-evidence", source_id="fixture-brief", block_id="fixture-block", start=0, end=len(quote),
            quote=quote, quote_sha256=sha256(quote.encode()).hexdigest(), locator={}, match_status="exact")],
        observation_ids=[], decision_ids=[], target=dict(object_kind="paragraph", semantic_role="body", selector="semantic_role"),
        property_id="font.size", value=dict(original_text=quote, normalized=12, unit="pt", dictionary_version="fixture/1.0"),
        character_scope="all", semantics=dict(scope=base_scope("body", "font.size").model_dump(), constraint={"operator":"eq"}),
        status="verified", reason="Trusted test fixture requirement."))
    def fixture_basis(context):
        empty_names = {"execution_auto", "execution_retry", "execution_paragraph", "workspace_history", "workspace_retry", "workspace_queue", "workspace_report"}
        if context.task_id in {item["task"] for name, item in rules_tasks.items() if name not in empty_names}:
            return FeedbackRuleBasis(academic_rules_snapshot_id=context.rules_snapshot_id, rules=[rule], fragments=analyze_rules([rule])[2])
        if context.task_id in {v0.task_id, duplicate_v0.task_id} | {item["task"] for name, item in rules_tasks.items() if name in empty_names}:
            return FeedbackRuleBasis(academic_rules_snapshot_id=context.rules_snapshot_id, rules=[], fragments=[])
        return real_basis(context)
    # Each browser run pins a private copy; backend checks may rebuild the
    # shared editor output while this session is still using its hashes.
    editor_copy = work / "editor"
    shutil.copytree(ROOT / "services/docx_engine/AutoFormat.DocxEditor/bin/Debug/net10.0", editor_copy)
    dll = editor_copy / "AutoFormat.DocxEditor.dll"
    feedback_editor, renderer_factory = None, ReviewRenderer
    if os.environ.get("AUTOFORMAT_TEST_FEEDBACK_EDITING") == "1":
        feedback_editor = DotnetEditor(dotnet_path=Path(shutil.which("dotnet")), editor_dll=dll, expected_editor_sha256=_sha256(dll))
        if os.environ.get("AUTOFORMAT_TEST_SOFFICE"):
            def renderer_factory(root):
                return LibreOfficeRenderer(soffice_path=Path(os.environ["AUTOFORMAT_TEST_SOFFICE"]),
                    expected_soffice_sha256=os.environ["AUTOFORMAT_TEST_SOFFICE_SHA256"], data_root=root)
    app = create_local_api(store=store, snapshots=snapshots, indexer=None, editor=feedback_editor,
        renderer_factory=renderer_factory, session_token=TOKEN, allowed_origin=f"http://127.0.0.1:{PORT}",
        font_fingerprint="test-fonts", web_dist_dir=web_dist, selection_indexes=SelectionIndexes.from_project(ROOT, editor_dll=dll), intake=intake,
        feedback_basis_loader=fixture_basis, application_mode="demo")
    uvicorn.run(app, host="127.0.0.1", port=PORT, access_log=False, proxy_headers=False, log_level="warning")


if __name__ == "__main__":
    main()
