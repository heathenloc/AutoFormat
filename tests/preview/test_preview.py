"""Safety and real headless conversion checks for the replaceable preview boundary."""

from __future__ import annotations

import hashlib
import json
import os
import shutil
import subprocess
from dataclasses import asdict
from pathlib import Path
from uuid import uuid4
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document
from jsonschema import Draft202012Validator
from pypdf import PdfReader

from autoformat.preview import LibreOfficeRenderer, PreviewError, inspect_docx_for_render
from autoformat.selection import verify_pdf_selection


@pytest.fixture
def safe_docx(tmp_path: Path) -> Path:
    path = tmp_path / "source.docx"
    document = Document()
    document.add_heading("预览定位样本", 0)
    document.add_paragraph("这是可选择的中文正文。")
    document.add_paragraph("Selectable English text 123.")
    document.save(path)
    return path


def _hash(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def test_safe_docx_preflight_returns_original_hash(safe_docx: Path) -> None:
    assert inspect_docx_for_render(safe_docx) == _hash(safe_docx)


def test_embedded_object_is_rejected_before_launch(safe_docx: Path) -> None:
    with ZipFile(safe_docx, "a", compression=ZIP_DEFLATED) as archive:
        archive.writestr("word/embeddings/embedded.bin", b"not executed")
    with pytest.raises(PreviewError, match="embedded") as error:
        inspect_docx_for_render(safe_docx)
    assert error.value.status == "unsafe_input"


def test_external_relationship_in_any_part_is_rejected(safe_docx: Path) -> None:
    relationships = (
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId999" Type="http://schemas.openxmlformats.org/officeDocument/'
        '2006/relationships/image" Target="https://example.invalid/x" TargetMode="External"/>'
        "</Relationships>"
    )
    with ZipFile(safe_docx, "a", compression=ZIP_DEFLATED) as archive:
        archive.writestr("custom/_rels/external.rels", relationships)
    with pytest.raises(PreviewError, match="external relationship") as error:
        inspect_docx_for_render(safe_docx)
    assert error.value.status == "unsafe_input"


def test_external_content_field_in_header_is_rejected(safe_docx: Path) -> None:
    header = (
        '<w:hdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        '<w:p><w:r><w:instrText>INCLUDEPICTURE https://example.invalid/x</w:instrText>'
        "</w:r></w:p></w:hdr>"
    )
    with ZipFile(safe_docx, "a", compression=ZIP_DEFLATED) as archive:
        archive.writestr("word/header99.xml", header)
    with pytest.raises(PreviewError, match="external-content field") as error:
        inspect_docx_for_render(safe_docx)
    assert error.value.status == "unsafe_input"


def test_binary_hash_mismatch_prevents_launch(tmp_path: Path) -> None:
    binary = tmp_path / "soffice.com"
    binary.write_bytes(b"not a real executable")
    with pytest.raises(PreviewError, match="hash changed"):
        LibreOfficeRenderer(
            soffice_path=binary, expected_soffice_sha256="0" * 64, data_root=tmp_path
        )


def test_school_template_with_embedded_object_is_not_rendered() -> None:
    template = Path(__file__).resolve().parents[2] / "TestDocuments" / "材料学院毕业论文模板2026.docx"
    with pytest.raises(PreviewError, match="embedded") as error:
        inspect_docx_for_render(template)
    assert error.value.status == "unsafe_input"


def test_real_headless_pdf_has_selectable_text(safe_docx: Path) -> None:
    executable = os.environ.get("AUTOFORMAT_TEST_SOFFICE")
    expected_hash = os.environ.get("AUTOFORMAT_TEST_SOFFICE_SHA256")
    if not executable or not expected_hash:
        pytest.skip("Set an explicit verified LibreOffice path and SHA-256 for integration.")

    repo_root = Path(__file__).resolve().parents[2]
    data_root = repo_root / "runtime" / "tests" / "preview" / uuid4().hex
    data_root.mkdir(parents=True)
    source = data_root / "source.docx"
    shutil.copyfile(safe_docx, source)
    renderer = LibreOfficeRenderer(
        soffice_path=Path(executable), expected_soffice_sha256=expected_hash,
        data_root=data_root, timeout_seconds=60,
    )
    source_hash = _hash(source)
    result = renderer.render(source, "preview-one")
    assert result.renderer_id == "libreoffice"
    assert result.renderer_version.startswith("LibreOffice 26.8.0")
    assert result.page_count >= 1
    assert result.selectable_text_chars > 20
    assert result.source_sha256 == source_hash == _hash(source)
    schema = json.loads((repo_root / "contracts" / "preview-artifact" /
                         "preview-artifact-1.0.schema.json").read_text(encoding="utf-8"))
    Draft202012Validator(schema).validate(asdict(result))
    pdf = data_root / result.pdf_relative_path
    assert pdf.is_file() and result.pdf_sha256 == _hash(pdf)
    text = "".join(page.extract_text() or "" for page in PdfReader(pdf).pages)
    assert "中文正文" in text and "Selectable English" in text
    verify_pdf_selection(
        artifact=result, data_root=data_root, source_docx=source,
        selected_text="中文正文", pdf_page=1,
    )
    with pytest.raises(PreviewError) as missing_selection:
        verify_pdf_selection(
            artifact=result, data_root=data_root, source_docx=source,
            selected_text="不存在的选区", pdf_page=1,
        )
    assert missing_selection.value.status == "ambiguous_target"
    with pytest.raises(PreviewError) as wrong_page:
        verify_pdf_selection(
            artifact=result, data_root=data_root, source_docx=source,
            selected_text="中文正文", pdf_page=result.page_count + 1,
        )
    assert wrong_page.value.status == "stale_preview"
    with pytest.raises(PreviewError) as wrong_artifact:
        verify_pdf_selection(
            artifact=result.__class__(**{**asdict(result), "pdf_sha256": "0" * 64}),
            data_root=data_root, source_docx=source,
            selected_text="中文正文", pdf_page=1,
        )
    assert wrong_artifact.value.status == "stale_preview"
    hint = {
        "schema_version": "selection-hint-prototype/1.0",
        "base_version_id": "version-1",
        "base_docx_sha256": source_hash,
        "preview_artifact_id": result.artifact_id,
        "renderer_id": result.renderer_id,
        "renderer_version": result.renderer_version,
        "selected_text": "中文正文",
        "context_before": "这是可选择的",
        "context_after": "。",
        "pdf_page": 1,
        "text_layer_start": 0,
        "text_layer_end": 4,
    }
    binding = {
        "base_version_id": "version-1",
        "base_docx_sha256": source_hash,
        "preview_artifact_id": result.artifact_id,
        "renderer_id": result.renderer_id,
        "renderer_version": result.renderer_version,
        "pdf_page_count": result.page_count,
    }
    hint_path = data_root / "selection-hint.json"
    binding_path = data_root / "trusted-preview-binding.json"
    hint_path.write_text(json.dumps(hint), encoding="utf-8")
    binding_path.write_text(json.dumps(binding), encoding="utf-8")
    for name, payload in (
        ("selection-hint/selection-hint-prototype-1.0.schema.json", hint),
        ("selection-anchor/preview-binding-1.0.schema.json", binding),
    ):
        schema = json.loads((repo_root / "contracts" / name).read_text(encoding="utf-8"))
        Draft202012Validator(schema).validate(payload)
    command = [
        "dotnet", "run", "--project",
        str(repo_root / "services" / "docx_engine" / "AutoFormat.DocxIndex"),
        "--", "locate", str(source), str(hint_path), str(binding_path),
    ]
    located = subprocess.run(command, cwd=repo_root, capture_output=True,
                             text=True, timeout=60, check=False)
    assert located.returncode == 0, located.stderr
    anchor_result = json.loads(located.stdout)
    assert anchor_result["status"] == "ready"
    anchor_schema = json.loads((repo_root / "contracts" / "selection-anchor" /
                                "selection-anchor-1.0.schema.json").read_text(encoding="utf-8"))
    Draft202012Validator(anchor_schema).validate(anchor_result["anchor"])
    assert anchor_result["anchor"]["docx_object_id"] == "p:000002"
    assert not list((data_root / "preview-work").iterdir()), "Isolated profile should be cleaned."
    with pytest.raises(PreviewError, match="already exists"):
        renderer.render(source, "preview-one")
